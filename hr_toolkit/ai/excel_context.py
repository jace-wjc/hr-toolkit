"""Convert Excel workbooks into LLM-friendly markdown analysis context.

The parser runs entirely on the local machine: rows are rendered into a
compact pipe-table plus per-column statistics so the cloud model answers
"compare these tables / who grew faster" questions from a bounded context.
``.xls`` inputs reuse the existing three-tier compatibility engine.
"""

from __future__ import annotations

import datetime as _datetime
import math
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Tuple, Union

from openpyxl import load_workbook

from hr_toolkit.common.excel_compat import ensure_xlsx_workbook

MAX_CONTEXT_ROWS = 500  # rows rendered into the prompt
MAX_STATS_ROWS = 20000  # rows scanned for statistics
MAX_SHEETS = 8
MAX_CELL_TEXT_LENGTH = 80
_HEADER_SCAN_ROWS = 10
_NUMERIC_COLUMN_RATIO = 0.8

SUPPORTED_SUFFIXES = frozenset({".xlsx", ".xls"})


@dataclass
class ColumnStats:
    header: str
    count: int
    total: float
    minimum: float
    maximum: float
    average: float


@dataclass
class SheetContext:
    name: str
    header_row: int  # 1-based row number of the detected header
    headers: List[str]
    rows: List[List[str]]  # rendered rows actually included in the prompt
    stats: List[ColumnStats]
    total_data_rows: int
    truncated: bool


@dataclass
class WorkbookContext:
    source: Path
    sheets: List[SheetContext] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, _datetime.datetime):
        if value.hour == value.minute == value.second == 0 and value.microsecond == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, _datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, _datetime.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        if math.isfinite(value) and value == int(value) and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    if len(text) > MAX_CELL_TEXT_LENGTH:
        return text[: MAX_CELL_TEXT_LENGTH - 1] + "…"
    return text


def _to_number(value) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            parsed = float(text)
        except ValueError:
            return None
        return parsed if math.isfinite(parsed) else None
    return None


def _detect_header_row(preview: List[List[str]]) -> int:
    """在非空的前几行里挑出表头行。

    人事表几乎都是「合并单元格的分组表头」：表头那行有不少空格（列分组），
    而它下面的第一行数据反而填得更满。老实现只数「非空格数最多」，于是把
    **数据行当成了表头**，真正表头之后的行全被当成数据丢掉 —— 结果整张表
    解析出 0 行数据，模型看不到任何内容，用户就以为「文件没发出去」。

    打分口径：表头几乎全是文字（数据行 numeric 密度高）、**第一格不空**
    （开头就空的是分组表头的续行，不是表头本身）、越靠前越像、太稀疏的不算。
    """
    best_index = 0
    best_score = -1.0
    fallback_index = 0
    fallback_score = -1.0
    for index, row in enumerate(preview):
        filled = [cell for cell in row if cell != ""]
        if len(filled) < 2:
            continue  # 标题、公司名这种单格行，不是表头
        textish = sum(1 for cell in filled if _to_number(cell) is None)
        score = textish / len(filled)
        score *= 1.0 / (1.0 + index * 0.15)  # 表头几乎总在数据前面
        # 只有两三个格的行不像表头；平方是故意的：数据行常常只有三四格有值，
        # 线性衰减压不住它们（身份证尾号是 X 还会被当成文字，进一步抬高分值）
        score *= min(1.0, len(filled) / 6.0) ** 2
        if score > fallback_score:
            fallback_index, fallback_score = index, score
        if row[0] == "":
            continue  # 开头就空：上一行的分组标签合并覆盖了这几列
        if score > best_score:
            best_index, best_score = index, score
    if best_score < 0:
        return fallback_index if fallback_score >= 0 else 0
    return best_index


def _is_header_continuation(row: Sequence[str], header_row: Sequence[str]) -> bool:
    """像不像紧挨着表头的那行分组子标签（「基数 / 个人(8%)」「应发工资」）。

    特征：开头几列是空的（被上一行的标签合并覆盖）、几乎没有数字、
    填的列与上一行重叠（还是同一张表）。
    """
    filled = [index for index, cell in enumerate(row) if cell != ""]
    if len(filled) < 2 or filled[0] == 0:
        return False
    textish = sum(1 for index in filled if _to_number(row[index]) is None)
    if textish < len(filled) * 0.8:
        return False
    return any(index < len(header_row) and header_row[index] for index in filled)


def _merge_header_rows(top: Sequence[str], sub: Sequence[str]) -> List[str]:
    """两行表头按列拼起来：「养老」+「基数」→「养老 基数」。"""
    width = max(len(top), len(sub))
    merged: List[str] = []
    for index in range(width):
        upper = top[index] if index < len(top) else ""
        lower = sub[index] if index < len(sub) else ""
        merged.append(f"{upper} {lower}" if upper and lower else (upper or lower))
    return merged


def _merge_header_block(preview: List[List[str]], header_index: int) -> Tuple[List[str], int]:
    """从表头行往后吸收分组续行，返回 (合并后的表头, 数据起始下标)。"""
    merged = list(preview[header_index]) if header_index < len(preview) else []
    cursor = header_index + 1
    absorbed = 0
    while cursor < len(preview) and absorbed < 3:
        candidate = preview[cursor]
        if not _is_header_continuation(candidate, merged):
            break
        merged = _merge_header_rows(merged, candidate)
        cursor += 1
        absorbed += 1
    return merged, cursor


def _unique_headers(header_row: Sequence[str], sheet_name: str) -> List[str]:
    headers: List[str] = []
    seen = {}
    for position, cell in enumerate(header_row):
        name = cell if cell else f"列{position + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 1
        headers.append(name)
    if not headers:
        headers.append(f"{sheet_name}数据")
    return headers


def _compute_stats(
    headers: Sequence[str], rows: Sequence[Sequence[str]]
) -> List[ColumnStats]:
    stats: List[ColumnStats] = []
    width = len(headers)
    for position, header in enumerate(headers):
        numbers: List[float] = []
        non_empty = 0
        for row in rows:
            if position >= len(row):
                continue
            value = row[position]
            if value == "":
                continue
            non_empty += 1
            parsed = _to_number(value)
            if parsed is not None:
                numbers.append(parsed)
        if non_empty < 3:
            continue
        if len(numbers) < non_empty * _NUMERIC_COLUMN_RATIO:
            continue
        stats.append(
            ColumnStats(
                header=header,
                count=len(numbers),
                total=round(sum(numbers), 2),
                minimum=round(min(numbers), 2),
                maximum=round(max(numbers), 2),
                average=round(sum(numbers) / len(numbers), 2),
            )
        )
    return stats


def _render_prompt_rows(rows: List[List[str]], max_rows: int) -> List[List[str]]:
    """Keep head and tail rows so period totals stay visible when truncated."""
    if len(rows) <= max_rows:
        return rows
    half = max_rows // 2
    head = rows[: max_rows - half]
    tail = rows[-half:] if half else []
    return head + tail


def build_workbook_context(path: Union[str, Path], *, max_rows: int = MAX_CONTEXT_ROWS) -> WorkbookContext:
    source = Path(path).expanduser()
    if source.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"仅支持 .xlsx 或 .xls 文件：{source.name}")
    if not source.is_file():
        raise ValueError(f"文件不存在：{source}")

    with tempfile.TemporaryDirectory(prefix="hr-ai-excel-") as temp_dir:
        converted = ensure_xlsx_workbook(source, Path(temp_dir))
        workbook = load_workbook(converted, read_only=True, data_only=True)
        try:
            sheet_names = [sheet.title for sheet in workbook.worksheets]
            contexts: List[SheetContext] = []
            notes: List[str] = []
            for sheet in workbook.worksheets[:MAX_SHEETS]:
                contexts.append(_read_sheet(sheet, max_rows=max_rows))
            skipped = sheet_names[MAX_SHEETS:]
            if skipped:
                notes.append(
                    "工作簿包含更多工作表未全部展示：" + "、".join(skipped[:5])
                    + ("等" if len(skipped) > 5 else "")
                )
            return WorkbookContext(source=source, sheets=contexts, notes=notes)
        finally:
            workbook.close()


def _read_sheet(sheet, *, max_rows: int) -> SheetContext:
    collected: List[Tuple[int, List[str]]] = []  # (真实行号, 渲染后的行)
    sheet_row_number = 0
    for row in sheet.iter_rows(values_only=True):
        sheet_row_number += 1
        rendered = [_cell_text(value) for value in row]
        while rendered and rendered[-1] == "":
            rendered.pop()
        if not any(cell != "" for cell in rendered):
            continue
        collected.append((sheet_row_number, rendered))
        if len(collected) >= MAX_STATS_ROWS + _HEADER_SCAN_ROWS:
            break

    preview = [item[1] for item in collected[:_HEADER_SCAN_ROWS]]
    header_index = _detect_header_row(preview)
    merged_header, data_start = _merge_header_block(preview, header_index)
    headers = _unique_headers(merged_header, sheet.title)
    data_rows = [item[1] for item in collected[data_start:]]
    width = len(headers)
    normalized_rows = [
        (row + [""] * (width - len(row)))[:width] for row in data_rows
    ]
    stats = _compute_stats(headers, normalized_rows)
    prompt_rows = _render_prompt_rows(normalized_rows, max_rows)
    return SheetContext(
        name=sheet.title,
        # 报真实的表头行号：行号会写进提示词，模型好对照「第几行」回答
        header_row=collected[header_index][0] if header_index < len(collected) else 1,
        headers=headers,
        rows=prompt_rows,
        stats=stats,
        total_data_rows=len(normalized_rows),
        truncated=len(normalized_rows) > len(prompt_rows),
    )


def _markdown_table(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> List[str]:
    lines = ["| " + " | ".join(headers) + " |"]
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in rows:
        cells = list(row)
        cells.extend([""] * (len(headers) - len(cells)))
        cells = cells[: len(headers)]
        lines.append("| " + " | ".join(cell for cell in cells) + " |")
    return lines


def shrink_workbook_context(context: WorkbookContext, budget: int, *, floor: int = 8) -> bool:
    """把整本工作簿的正文收紧到 ``budget`` 个字符以内；塞得下返回 True。

    真实的人事表动辄几千行，整本贴给模型会超预算。这里砍的是「贴给模型的行」，
    **数值列统计仍然基于全量数据**（统计在读取时就算完了），所以模型拿到的合计、
    平均值还是准的，只是看不到中间每一行。逐档减半而不是一步到位，能多保住多少行
    就保住多少行。
    """
    per_sheet = MAX_CONTEXT_ROWS
    while True:
        for sheet in context.sheets:
            if len(sheet.rows) > per_sheet:
                sheet.rows = _render_prompt_rows(sheet.rows, per_sheet)
                sheet.truncated = True
        if len(render_workbook_markdown(context)) <= budget:
            return True
        if per_sheet <= floor:
            return False
        per_sheet = max(floor, per_sheet // 2)


def render_workbook_markdown(context: WorkbookContext) -> str:
    lines: List[str] = [f"### 文件：{context.source.name}"]
    for note in context.notes:
        lines.append(f"（提示：{note}）")
    for sheet in context.sheets:
        lines.append("")
        lines.append(f"#### 工作表「{sheet.name}」")
        lines.append(
            f"表头位于第 {sheet.header_row} 行，数据共 {sheet.total_data_rows} 行"
            f"（本次提供 {len(sheet.rows)} 行）。"
        )
        if not sheet.rows and not sheet.stats:
            lines.append("（该表没有可识别的数据行）")
            continue
        lines.extend(_markdown_table(sheet.headers, sheet.rows))
        if sheet.stats:
            lines.append("")
            lines.append("数值列统计（本地计算，可信）：")
            stats_headers = ["列", "数值个数", "合计", "最小值", "最大值", "平均值"]
            stats_rows = [
                [
                    stat.header,
                    str(stat.count),
                    f"{stat.total:g}",
                    f"{stat.minimum:g}",
                    f"{stat.maximum:g}",
                    f"{stat.average:g}",
                ]
                for stat in sheet.stats
            ]
            lines.extend(_markdown_table(stats_headers, stats_rows))
        if sheet.truncated:
            lines.append("")
            lines.append(
                "（该表数据较多，已截断为首尾部分行；合计等统计基于全部数据，"
                "请优先依据统计结果回答，避免根据截断行数下结论。）"
            )
    return "\n".join(lines)


ANALYSIS_PREAMBLE = (
    "你是 HR Toolkit 内置的表格分析助手。下面提供用户上传的表格内容，"
    "均为本地解析后的结果；「数值列统计」由程序本地计算，数值可信，请优先依据统计结果回答。\n"
    "回答要求：\n"
    "1. 用简体中文，直接给结论，再给依据；涉及金额和人数时明确写出数字。\n"
    "2. 对比多张表时，先总结相同点和差异点，再回答用户的具体问题。\n"
    "3. 数据中可能存在身份证号等敏感信息，回答时只引用必要的业务字段。\n"
    "4. 如果数据不足以回答，明确说明缺少什么，不要编造。\n"
)


def render_analysis_context(
    workbooks: Sequence[WorkbookContext], question: str = ""
) -> str:
    parts: List[str] = [ANALYSIS_PREAMBLE]
    for context in workbooks:
        parts.append("")
        parts.append(render_workbook_markdown(context))
    if workbooks:
        parts.append("")
        parts.append("---")
        parts.append("以上共 " + str(len(workbooks)) + " 个文件。")
    if question.strip():
        parts.append("用户的问题：" + question.strip())
    return "\n".join(parts)


def estimate_tokens(text: str) -> int:
    """Rough token estimate for mixed Chinese/English prompt text."""
    return max(1, (len(text or "") + 1) // 2)
