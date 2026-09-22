"""Read-only source reconciliation; write reports and a separately filled summary.

Exports without workflow status are compared as provided, with a runtime log
explaining that status filtering was unavailable. Completion time and personnel
status are never substitutes. Ambiguous identities/events are never auto-selected.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re
from typing import Any, Callable
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from hr_toolkit.common.excel_compat import ensure_xlsx_workbook
from hr_toolkit.common.run_temp import temporary_directory
from hr_toolkit.common.template_mapping import (
    active, assigned_role, file_template_source, map_sheet, request_selection,
    template_tool,
)
from hr_toolkit.tools.personnel_change_merge import _iter_input_files


COMMON_FIELDS = {
    "姓名": ("姓名", "员工姓名"),
    "身份证号码": ("身份证号码", "身份证号", "身份证"),
    "公司": ("公司", "入职公司"),
    "工号": ("工号", "员工工号"),
    "入职日期": ("入职日期", "入职时间"),
    "离职日期": ("离职日期", "离职时间"),
    "预计离职日期": ("预计离职日期",),
    "出生日期": ("出生日期",), "年龄": ("年龄",),
    "毕业时间": ("毕业时间",), "备注": ("备注",),
    "人员分类": ("人员分类",), "用工状态": ("用工状态",),
    "试用期工资": ("试用期工资",), "薪资结算日期": ("薪资结算日期",),
}
SUMMARY_FIELDS = {
    **COMMON_FIELDS,
    "性别": ("性别",), "岗位": ("岗位", "职务"),
    "地市": ("地市",), "所属专业": ("所属专业",),
    "联系方式": ("联系方式", "工作联系电话"),
    "学历": ("学历",), "毕业学校": ("毕业学校", "毕业院校"),
    "专业": ("专业",), "婚否": ("婚否", "婚姻状况"),
    "家庭住址": ("家庭住址",),
}
FLOW_FIELDS = {
    **COMMON_FIELDS,
    "流程状态": ("流程状态", "审批状态"),
    "性别": ("性别",), "岗位": ("岗位",), "地市": ("所属市",),
    "所属专业": ("所属专业",), "联系方式": ("手机号码",),
    "学历": ("最高学历文化程度",), "毕业学校": ("最高学历毕业学校",),
    "专业": ("最高学历所学专业",), "婚否": ("婚姻状况",),
    "家庭住址": ("户籍地址",),
}
ROLES = {
    "summary_join": ("汇总表增员", "入职", SUMMARY_FIELDS),
    "summary_leave": ("汇总表减员", "离职", SUMMARY_FIELDS),
    "flow_join": ("入职流程", "入职", FLOW_FIELDS),
    "flow_leave": ("离职流程", "离职", FLOW_FIELDS),
}
EXCLUDED_STATUSES = frozenset({"未发起", "退回"})
FILL_FIELDS = ("性别", "岗位", "地市", "所属专业", "联系方式", "学历",
               "毕业学校", "专业", "婚否", "家庭住址", "入职日期",
               "出生日期", "年龄", "工号", "毕业时间", "备注",
               "人员分类", "用工状态", "试用期工资", "薪资结算日期")
DATE_FIELDS = frozenset({"入职日期", "出生日期", "毕业时间", "薪资结算日期"})
NOTICE_HEADERS = ("类型", "异动", "姓名", "身份证号码", "字段", "汇总表内容",
                  "流程内容", "汇总表位置", "流程位置", "说明",
                  "汇总表公司", "流程公司", "汇总表工号", "流程工号", "汇总表事件日期", "流程事件日期")
FILL_HEADERS = ("异动", "姓名", "身份证号码", "补入字段", "补入内容", "汇总表位置", "流程位置")
FILLED = PatternFill("solid", fgColor="E4EFEA")
WARNING = PatternFill("solid", fgColor="FFF2CC")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value: Any) -> str:
    return re.sub(r"\s+", "", text(value)).upper()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not 1 <= value <= 100000:
            return None
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else None
    raw = text(value)
    match = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?", raw)
    if not match:
        return None
    try:
        return date(*(int(part) for part in match.groups()))
    except ValueError:
        return None


def validate_month(value: str) -> str:
    raw = text(value)
    if raw and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", raw):
        raise ValueError("核对月份请填写 YYYY-MM，例如 2026-07；留空按汇总表中各类实际日期核对。")
    if raw and int(raw[:4]) < 1900:
        raise ValueError("核对年份不能早于 1900 年。")
    return raw


def parse_company_aliases(value: str) -> dict[str, str]:
    """Explicit overrides only; reject conflicting/cyclic alias definitions."""
    result = {}
    for item in re.split(r"[;；\n]+", text(value)):
        if not item.strip():
            continue
        parts = re.split(r"[=＝]", item)
        if len(parts) != 2 or not all(part.strip() for part in parts):
            raise ValueError("公司对应关系请按“简称=完整公司名称”填写，多组用分号分隔。")
        key, target = map(company_text, parts)
        if key in result and result[key] != target:
            raise ValueError("同一个公司简称不能对应两个不同名称。")
        result[key] = target
    for key, target in result.items():
        if target in result and result[target] != target:
            raise ValueError("公司对应关系请直接指向最终全称，不要连续转换。")
    return result


def company_text(value: Any) -> str:
    return re.sub(r"[\s（）()·]", "", text(value)).upper()


def _keyword_match(short: str, full: str) -> bool:
    if not 2 <= len(short) <= 12 or len(short) >= len(full):
        return False
    if short in {"公司", "集团", "中支", "分公司", "有限公司", "人力资源", "信息技术"}:
        return False
    if short in full:
        return True
    # 北京春苗 -> 春苗人力资源（北京）有限公司. Keep both pieces;
    # do not erase geography/branches or accept a vague similarity score.
    return any(short[:i] in full and short[i:] in full
               for i in range(2, len(short) - 1))


def company_mapping(summary_names, flow_names, overrides):
    names = {company_text(v) for v in flow_names if text(v)}
    mapping, ambiguous = {}, set()
    for key in {company_text(value) for value in summary_names}:
        if not key:
            continue
        if key in overrides:
            mapping[key] = overrides[key]
        elif key in names:
            mapping[key] = key
        else:
            candidates = {name for name in names if _keyword_match(key, name)}
            if len(candidates) == 1:
                mapping[key] = next(iter(candidates))
            elif candidates:
                ambiguous.add(key)
    return mapping, ambiguous


@dataclass
class Record:
    kind: str
    values: dict[str, Any]
    source: str
    sheet: str
    row: int
    columns: dict[str, int]
    day: date | None

    @property
    def identity(self):
        return norm(self.values.get("身份证号码"))

    @property
    def location(self):
        source = Path(self.source).name if Path(self.source).is_absolute() else self.source
        return f"{source} / {self.sheet} / 第{self.row}行"


@dataclass
class ReconcileResult:
    output_file: Path | None = None
    filled_output_file: Path | None = None
    warnings: list[str] = field(default_factory=list)
    matched_count: int = 0
    filled_count: int = 0
    issue_count: int = 0

    def to_dict(self):
        return {"output_file": str(self.output_file or ""),
                "filled_output_file": str(self.filled_output_file or ""),
                "matched_count": self.matched_count, "filled_count": self.filled_count,
                "issue_count": self.issue_count, "warnings": self.warnings}


def _check(cancelled):
    if cancelled is not None and cancelled():
        raise RuntimeError("本次处理已停止。")


def _headers(ws, role, file):
    view = map_sheet(ws, role, file=file)
    aliases = ROLES[role][2]
    header_row = getattr(view, "header_row", 0)
    if not header_row:
        header_row = next((r for r in range(1, min(ws.max_row, 20) + 1)
                           if any(norm(ws.cell(r, c).value) in {"姓名", "员工姓名"}
                                  for c in range(1, min(ws.max_column, 512) + 1))), 0)
    if not header_row:
        raise ValueError(f"{file} / {ws.title}：找不到姓名表头，请使用模板适配。")
    header = {c: norm(view.cell(header_row, c).value)
              for c in range(1, min(ws.max_column, 512) + 1)}
    columns = {}
    for name, choices in aliases.items():
        for alias in (name, *choices):
            found = [c for c, value in header.items() if value and value == norm(alias)]
            if len(found) > 1:
                raise ValueError(f"{file} / {ws.title}：{name}有多列，请在模板适配中明确选择。")
            if found:
                columns[name] = found[0]
                break
    for required in ("姓名", "身份证号码", "公司"):
        if required not in columns:
            raise ValueError(f"{file} / {ws.title}：缺少{required}，请使用模板适配。")
    return header_row, columns


def _records(ws, role, source, leave_date_field, cancelled):
    row, columns = _headers(ws, role, source)
    kind = ROLES[role][1]
    date_field = "入职日期" if kind == "入职" else (
        leave_date_field if role.startswith("flow_") else "离职日期")
    records = []
    for r in range(row + 1, ws.max_row + 1):
        if r % 128 == 0:
            _check(cancelled)
        values = {name: ws.cell(r, c).value for name, c in columns.items()}
        if not any(text(values.get(key)) for key in ("姓名", "身份证号码", "公司", date_field)):
            continue
        if any(marker in text(values.get("姓名")) for marker in ("审批人：", "审核人：", "合计")):
            continue
        records.append(Record(kind, values, source, ws.title, r, columns,
                              parse_date(values.get(date_field))))
    return records


class _SparseSheet:
    """Compact value-only view for template mapping; skip empty format rows."""
    def __init__(self, ws, cancelled):
        self.title = ws.title
        self.rows = {}
        self.max_row = self.max_column = 0
        ws.reset_dimensions()
        for r, values in enumerate(ws.iter_rows(max_col=512, values_only=True), 1):
            if r % 128 == 0:
                _check(cancelled)
            cells = {c: v for c, v in enumerate(values, 1) if v is not None}
            if cells:
                self.rows[r] = cells
                self.max_row = r
                self.max_column = max(self.max_column, max(cells))

    def cell(self, row, column):
        from types import SimpleNamespace
        return SimpleNamespace(value=self.rows.get(row, {}).get(column))

    def iter_rows(self, min_row=1, max_row=None, max_col=None, values_only=False):
        for r in range(min_row, (max_row or self.max_row) + 1):
            yield tuple(self.cell(r, c).value if values_only else self.cell(r, c)
                        for c in range(1, (max_col or self.max_column) + 1))


@file_template_source
def _read_flows(path, source, leave_date_field, cancelled):
    book = load_workbook(path, read_only=True, data_only=True)
    found = defaultdict(list)
    try:
        for original in book:
            ws = _SparseSheet(original, cancelled)
            if not ws.max_row:
                continue
            roles = ["flow_join", "flow_leave"]
            role = assigned_role(ws, roles, file=source, confirmed_only=True) if active() else None
            if not role:
                headers = {norm(ws.cell(r, c).value) for r in range(1, min(ws.max_row, 20) + 1)
                           for c in range(1, ws.max_column + 1)}
                if headers & {"预计离职日期", "离职日期", "离职时间", "减员类型"}:
                    role = "flow_leave"
                elif headers & {"入职日期", "入职时间"}:
                    role = "flow_join"
            if not role:
                if active():
                    request_selection([ws], roles, file=source, message="请选择这张表是入职流程还是离职流程。")
                raise ValueError(f"{source} / {ws.title}：无法识别流程类型。")
            found[ROLES[role][1]].extend(_records(ws, role, source, leave_date_field, cancelled))
    finally:
        book.close()
    return found


def _notice(kind, summary=None, flow=None, field_name="", detail=""):
    record = summary or flow
    return [kind, record.kind if record else "", text(record.values.get("姓名")) if record else "",
            record.identity if record else "", field_name,
            summary.values.get(field_name) if summary else "",
            flow.values.get(field_name) if flow else "",
            summary.location if summary else "", flow.location if flow else "", detail,
            summary.values.get("公司") if summary else "", flow.values.get("公司") if flow else "",
            text(summary.values.get("工号")) if summary else "", text(flow.values.get("工号")) if flow else "",
            summary.day if summary else None, flow.day if flow else None]


def _safe_value(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Imported text is never an executable Excel formula.


def _table(book, name, headers, rows, cancelled):
    ws = book.create_sheet(name)
    for r, values in enumerate([headers, *rows], 1):
        if r % 128 == 0:
            _check(cancelled)
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c)
            _safe_value(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            if r == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17715B")
    ws.freeze_panes = "E2" if name == "预警明细" else "A2"
    ws.auto_filter.ref = ws.dimensions
    for c, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = (100 if header == "内容" else
            40 if header in {"说明", "汇总表位置", "流程位置", "汇总表公司", "流程公司"} else 22)


@template_tool("personnel_reconcile")
def reconcile_personnel_changes(
    input_dir: str | Path | list[str | Path],
    template_path: str | Path,
    output_dir: str | Path,
    *, month: str = "", leave_date_field: str = "离职日期",
    company_aliases: str = "", highlight: bool = True,
    cancelled: Callable[[], bool] | None = None,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> ReconcileResult:
    """Compare one generated summary against department workflow exports.

    No new/removed people, no edits to source files, no updates to nonblank cells.
    Missing dates, blank status values and masked IDs stay visible in the report.
    """
    _check(cancelled)
    month = validate_month(month)
    if leave_date_field not in {"离职日期", "预计离职日期"}:
        raise ValueError("离职日期列只能选择离职日期或预计离职日期。")
    overrides = parse_company_aliases(company_aliases)
    inputs = [Path(p).expanduser().resolve() for p in (input_dir if isinstance(input_dir, list) else [input_dir])]
    source = Path(template_path).expanduser().resolve()
    if not source.is_file() or not inputs:
        raise ValueError("请选择系统流程文件及工具生成的异动汇总表。")
    result = ReconcileResult()
    notices, fills, info = [], [], []
    output_dir = Path(output_dir).expanduser().resolve()
    with temporary_directory(prefix="hr_reconcile_") as temp:
        temp = Path(temp)
        flows = defaultdict(list)
        seen = set()
        for selected in inputs:
            if not selected.exists():
                raise ValueError(f"流程资料不存在：{selected.name}")
            for candidate in _iter_input_files(selected, temp, result.warnings):
                _check(cancelled)
                if candidate.resolve() in seen:
                    continue
                if candidate.resolve() == source:
                    raise ValueError("汇总表不能同时作为系统流程表上传。")
                seen.add(candidate.resolve())
                path = ensure_xlsx_workbook(candidate, temp, warning_callback=result.warnings.append)
                try:
                    label = str(candidate.relative_to(selected)) if selected.is_dir() else candidate.name
                except ValueError:
                    label = candidate.name  # Archive member extracted outside the selected folder.
                label = selected.name + " / " + label if selected != candidate else str(selected)
                for kind, rows in _read_flows(path, label, leave_date_field, cancelled).items():
                    flows[kind].extend(rows)
        if not flows:
            raise ValueError("没有找到可识别的入职或离职流程表。")
        unfiltered_kinds = [kind for kind, rows in flows.items()
                            if any("流程状态" not in row.columns for row in rows)]
        if unfiltered_kinds and progress_callback is not None:
            progress_callback(0, 0, "提示：" + "、".join(unfiltered_kinds)
                              + "流程资料中存在无流程状态列的记录，这些记录按导入内容参与核对和空白补齐，"
                              "未进行状态筛选，无法判断是否包含未发起、退回；办结时间不作为状态依据。")
        summary_path = ensure_xlsx_workbook(source, temp, preserve_formatting=True, warning_callback=result.warnings.append)
        book = load_workbook(summary_path, data_only=False)
        try:
            _reconcile(book, source, flows, month, leave_date_field,
                       overrides, highlight, result, notices, fills, info, cancelled)
            _check(cancelled)
            report = Workbook()
            report.remove(report.active)
            try:
                _table(report, "核对说明", ("项目", "内容"), info, cancelled)
                _table(report, "预警明细", NOTICE_HEADERS, notices, cancelled)
                _table(report, "补入记录", FILL_HEADERS, fills, cancelled)
                output_dir.mkdir(parents=True, exist_ok=True)
                # Exclusive run outputs. Never overwrite an existing result/source.
                suffix = uuid4().hex[:8]
                report_target = output_dir / f"异动核对结果_{suffix}.xlsx"
                filled_target = output_dir / f"{source.stem}_核对补全_{suffix}.xlsx"
                with temporary_directory(prefix=".reconcile_", dir=str(output_dir)) as staged:
                    staged = Path(staged)
                    report.save(staged / "report.xlsx")
                    book.save(staged / "filled.xlsx")
                    _check(cancelled)
                    (staged / "report.xlsx").rename(report_target)
                    (staged / "filled.xlsx").rename(filled_target)
                    result.filled_output_file = filled_target
                result.output_file = report_target
            finally:
                report.close()
        finally:
            book.close()
    result.filled_count = len(fills)
    result.issue_count = len(notices)
    if notices:
        result.warnings.append(f"有 {len(notices)} 条预警或待确认记录，请查看核对结果中的预警明细。")
    return result


@file_template_source
def _summary_rows(path, book, cancelled):
    result = {}
    for sheet, role in (("增员", "summary_join"), ("减员", "summary_leave")):
        if sheet not in book.sheetnames:
            raise ValueError(f"请使用工具生成的异动汇总表，当前文件缺少“{sheet}”工作表。")
        result[ROLES[role][1]] = _records(book[sheet], role, str(path), "离职日期", cancelled)
    return result


def _reconcile(book, source, flows, month, leave_date_field,
               overrides, highlight, result, notices, fills, info, cancelled):
    summaries = _summary_rows(source, book, cancelled)
    info.extend([
        ("核对范围", "按本次导入资料核对增员、减员，其他工作表不处理；核对范围取决于提供的资料。"),
        ("流程筛选规则", "有流程状态列时排除未发起、退回；无状态列的记录按导入内容核对和补齐，不按办结时间或人员状态推测。"),
        ("离职流程日期列", leave_date_field),
        ("补齐规则", "仅明确匹配且身份证未脱敏时补空白；最高学历、岗位、户籍地址；已有值和公式不覆盖。"),
        ("空白保留规则", "流程未提供的字段保留空白。人员分类、用工状态等字段只读取同名列或用户在模板适配中明确指定的列，不推测含义。"),
        ("重复流程", "不自动选第一条或最新一条；逐条列出，由人工核实。"),
    ])
    for kind in ("入职", "离职"):
        _check(cancelled)
        summary_all = summaries[kind]
        periods = {month} if month else {r.day.strftime("%Y-%m") for r in summary_all if r.day}
        info.append((kind + "核对月份", "、".join(sorted(periods)) or "无法确定，请填写核对月份"))
        if not periods:
            notices.append(_notice("未核对", detail=kind + "无法确定月份，请填写核对月份后重试。"))
            continue
        target = [r for r in summary_all if r.day and r.day.strftime("%Y-%m") in periods]
        for row in summary_all:
            if row.day is None:
                notices.append(_notice("日期缺失或无效", summary=row, detail="无法确定事件月份，保留原记录且不补齐。"))
        if kind not in flows:
            notices.append(_notice("未核对", detail=kind + "流程未提供，不能判定漏提交流程。"))
            info.append((kind + "核对状态", "未完成：流程未提供"))
            continue
        eligible, valid_flows, excluded = [], [], 0
        incomplete = any(r.day is None or not r.identity or not text(r.values.get("公司")) for r in summary_all)
        uncertain_by_id = defaultdict(list)
        any_unknown_flow_identity = False
        missing_status = False
        for row in flows[kind]:
            raw_status = row.values.get("流程状态")
            has_status = "流程状态" in row.columns
            if has_status and norm(raw_status) in EXCLUDED_STATUSES:
                excluded += 1
                continue
            if not has_status:
                missing_status = True
            problem = ("流程状态为空" if has_status and not text(raw_status) else
                       "所选事件日期缺失或无效" if row.day is None else
                       "身份证或公司为空" if not row.identity or not text(row.values.get("公司")) else "")
            if problem:
                notices.append(_notice("流程待确认", flow=row, detail=problem + "；日期列为“" + ("入职日期" if kind == "入职" else leave_date_field) + "”；不自动补齐，也不直接认定对方漏登记。"))
                uncertain_by_id[row.identity].append(row)
                any_unknown_flow_identity |= not bool(row.identity)
                continue
            valid_flows.append(row)
            if row.day.strftime("%Y-%m") in periods:
                eligible.append(row)
        info.append((kind + "流程状态来源", "部分或全部记录无状态列，相关记录未进行状态筛选；其余记录读取流程状态列" if missing_status else "读取流程状态列"))
        info.append((kind + "范围内记录", f"汇总表 {len(target)} 条；参与核对流程 {len(eligible)} 条；按状态列排除未发起/退回 {excluded} 条"))
        mapping, ambiguous = company_mapping(
            (r.values.get("公司") for r in summary_all),
            (r.values.get("公司") for r in flows[kind]), overrides)
        for short, full in sorted(mapping.items()):
            if short != full:
                info.append((kind + "公司对应", f"{short} → {full}"))
        by_id = defaultdict(list)
        for row in valid_flows:
            by_id[row.identity].append(row)
        target_counts = Counter((r.identity, mapping.get(company_text(r.values.get("公司")), company_text(r.values.get("公司"))), r.day) for r in target)
        used = set()
        # Potential duplicate admissions in this scope are not auto-selected,
        # even when submission/event dates differ. Other months stay separate.
        event_groups = defaultdict(list)
        for row in eligible:
            event_groups[(row.identity, company_text(row.values.get("公司")),
                          None if kind == "入职" else row.day,
                          "" if kind == "入职" else norm(row.values.get("工号")))].append(row)
        duplicate_rows = {id(r) for group in event_groups.values() if len(group) > 1 for r in group}
        for row in eligible:
            if id(row) in duplicate_rows:
                notices.append(_notice("重复流程待确认", flow=row, detail="核对范围内同一人员、公司存在多条候选流程，不自动选择或补齐；请核实是否重复提交或再次入职。"))
        for row in target:
            _check(cancelled)
            company = company_text(row.values.get("公司"))
            canonical = mapping.get(company, company)
            candidates = by_id.get(row.identity, [])
            same_company = [p for p in candidates if company_text(p.values.get("公司")) == canonical]
            employee = norm(row.values.get("工号"))
            if employee:
                same_employee = [p for p in same_company if norm(p.values.get("工号")) == employee]
                if same_employee:
                    same_company = same_employee
            exact = [p for p in same_company if p.day == row.day]
            if not row.identity or not company:
                notices.append(_notice("身份信息不完整", summary=row, detail="身份证或公司为空，不自动关联。"))
                continue
            if company in ambiguous:
                used.update(id(p) for p in candidates)
                for candidate in candidates or [None]:
                    notices.append(_notice("公司对应不唯一", summary=row, flow=candidate, field_name="公司",
                                           detail="公司关键词匹配到多个全称，请在公司对应关系中明确填写。"))
                continue
            uncertain = [p for p in uncertain_by_id.get(row.identity, [])
                         if (not company_text(p.values.get("公司")) or company_text(p.values.get("公司")) == canonical)
                         and (not employee or not norm(p.values.get("工号")) or norm(p.values.get("工号")) == employee)]
            if uncertain:
                for candidate in uncertain:
                    notices.append(_notice("对应流程待确认", summary=row, flow=candidate,
                                           detail="同一人员、公司和工号的候选流程存在不完整字段，暂不自动补齐。"))
                continue
            if any(id(p) in duplicate_rows for p in same_company):
                used.update(id(p) for p in same_company)
                notices.append(_notice("重复记录待确认", summary=row, detail="存在多条候选流程，逐条列在预警明细中；人工确认前不补齐。"))
                continue
            if target_counts[(row.identity, canonical, row.day)] > 1 or len(exact) > 1:
                used.update(id(p) for p in exact)
                notices.append(_notice("重复记录待确认", summary=row, detail="同一事件不是一对一关系，不自动补齐。"))
                continue
            if len(exact) == 1:
                flow = exact[0]
                used.add(id(flow))
                if norm(row.values.get("姓名")) != norm(flow.values.get("姓名")):
                    notices.append(_notice("姓名不一致", row, flow, "姓名", "身份证相同但姓名不同，不自动补齐。"))
                    continue
                if text(row.values.get("工号")) and text(flow.values.get("工号")) and norm(row.values["工号"]) != norm(flow.values["工号"]):
                    notices.append(_notice("工号不一致", row, flow, "工号", "不自动补齐。"))
                    continue
                result.matched_count += 1
                identity_safe = isinstance(row.values.get("身份证号码"), str) and isinstance(flow.values.get("身份证号码"), str) and bool(re.fullmatch(r"(?:\d{17}[0-9X]|\d{15})", row.identity))
                if not identity_safe:
                    notices.append(_notice("身份证待确认", row, flow, "身份证号码", "身份证脱敏、格式异常或以数字存储；仅列出核对结果，不自动补齐。"))
                for name in FILL_FIELDS:
                    if name not in row.columns or name not in flow.columns:
                        continue
                    value, original = flow.values.get(name), row.values.get(name)
                    if not text(value):
                        continue
                    if name in DATE_FIELDS:
                        value = parse_date(value)
                        if value is None:
                            continue
                    cell = book[row.sheet].cell(row.row, row.columns[name])
                    if cell.data_type == "f":
                        continue
                    if not text(original):
                        if identity_safe:
                            _safe_value(cell, value)
                            if isinstance(value, date):
                                cell.number_format = "yyyy/m/d"
                            if highlight:
                                cell.fill = FILLED
                            fills.append([kind, row.values.get("姓名"), row.identity, name, value, row.location, flow.location])
                    elif (parse_date(original) != value if name in DATE_FIELDS else norm(original) != norm(value)):
                        notices.append(_notice("字段差异", row, flow, name, "保留汇总表已有内容。"))
                        if highlight:
                            cell.fill = WARNING
            elif candidates:
                compared = same_company or candidates
                used.update(id(p) for p in compared)
                for flow in compared:
                    field_name = "入职日期" if kind == "入职" else "离职日期"
                    notice = _notice("日期差异" if same_company else "公司差异", row, flow,
                                     field_name if same_company else "公司", "存在同一身份证的候选流程，但事件未明确匹配，不自动补齐。")
                    if same_company:
                        notice[6] = flow.day
                    notices.append(notice)
            else:
                notices.append(_notice("待确认" if any_unknown_flow_identity else "异动表有、流程无", summary=row,
                                       detail="流程中有身份不完整记录，暂不能判定漏提交。" if any_unknown_flow_identity else "本次导入流程经适用的状态及日期筛选后未找到对应人员，请核实。"))
        target_ids = {r.identity for r in summary_all}
        for flow in eligible:
            if id(flow) not in used and id(flow) not in duplicate_rows:
                uncertain = incomplete or flow.identity in target_ids
                notices.append(_notice("待确认" if uncertain else "流程有、异动表无", flow=flow,
                                       detail="汇总表存在身份/日期缺失或其他日期的同一人员，请人工确认事件。" if uncertain else "本次汇总表及日期范围内未找到对应人员，请核实是否漏登记。"))
        info.append((kind + "核对状态", "存在待确认数据，详见预警明细" if uncertain_by_id or incomplete else "已完成范围内比对，异常见预警明细"))
    info.extend([("明确关联记录数", result.matched_count), ("补入字段数", len(fills)), ("预警明细数", len(notices))])
