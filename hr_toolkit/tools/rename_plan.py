"""Editable rename plans. Legacy CLI rename modes remain in folder_rename.

A confirmed plan contains explicit relative names and portable content fingerprints;
execution never rereads an Excel roster or regenerates a naming rule.
"""
from __future__ import annotations

import copy
import csv
import hashlib
import json
import os
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from . import folder_rename as legacy


def _fingerprint(path: Path, cancelled=None) -> str:
    from hr_toolkit.project_store import _is_ignored_import_file

    digest = hashlib.sha256()
    def visit(item, relative):
        legacy._check_cancelled(cancelled)
        legacy._validate_text_source(item)
        if item.is_dir():
            digest.update(('D:' + relative + '\0').encode('utf-8'))
            for child in sorted(item.iterdir()):
                legacy._validate_text_source(child)
                if child.is_file() and _is_ignored_import_file(child):
                    continue
                visit(child, relative + '/' + child.name)
        else:
            digest.update(('F:' + relative + '\0').encode('utf-8'))
            before = item.stat()
            with item.open('rb') as stream:
                while True:
                    legacy._check_cancelled(cancelled)
                    chunk = stream.read(1024 * 1024)
                    if not chunk:
                        break
                    digest.update(chunk)
            after = item.stat()
            if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                raise ValueError(f'{item.name} 在检查期间变化，请重新预览')
            digest.update(b'\0')
    visit(path, '')
    return digest.hexdigest()


def validate_plan(plan: dict) -> dict:
    """Pure O(n) validation for background UI edits, using the inspected directory."""
    result = copy.deepcopy(plan)
    rows = result['rows']
    if any(row.get('row_id') != str(index) for index, row in enumerate(rows)):
        raise ValueError('改名项目编号无效，请重新预览')
    existing = result['existing_names']
    sources = Counter(row['source_name'].casefold() for row in rows)
    selected = [row for row in rows if row['included']]
    targets = Counter(row['target_name'].casefold() for row in selected)
    vacated = {row['source_name'].casefold() for row in selected if row['target_name'] != row['source_name']}
    for row in rows:
        issues = []
        if row['included']:
            try:
                legacy._validate_folder_name(row['source_name'])
                legacy._validate_folder_name(row['target_name'])
                if len(row['target_name'].encode('utf-16-le')) > 510:
                    raise ValueError('目标名称超过 Windows 单个名称长度限制，请缩短')
            except ValueError as exc:
                issues.append(str(exc))
            if row.get('source_error'):
                issues.append(row['source_error'])
            if sources[row['source_name'].casefold()] > 1:
                issues.append('原名称存在 Windows 大小写冲突')
            if not row['is_dir'] and (Path(row['target_name']).suffix != row['suffix'] or
                                     row['target_name'] == row['suffix']):
                issues.append('文件扩展名必须保持不变')
            target_key = row['target_name'].casefold()
            if targets[target_key] > 1:
                issues.append('目标名称重复（不区分大小写）')
            occupants = existing.get(target_key, [])
            for occupant in occupants:
                if occupant != row['source_name'] and occupant.casefold() not in vacated:
                    issues.append('目标已存在，且该项目未参与改名')
                    break
            if len(occupants) > 1:
                issues.append('目标位置存在 Windows 大小写冲突')
        row['issue'] = '；'.join(issues)
        row['status'] = ('excluded' if not row['included'] else 'conflict' if issues else
                         'unchanged' if row['source_name'] == row['target_name'] else 'ready')
        row['status_text'] = {'excluded': '已排除', 'conflict': '需处理', 'unchanged': '不变', 'ready': '待改名'}[row['status']]
    result['conflict_count'] = sum(row['status'] == 'conflict' for row in rows)
    result['change_count'] = sum(row['status'] == 'ready' for row in rows)
    return result


def _existing(root):
    names = {}
    for path in root.iterdir():
        names.setdefault(path.name.casefold(), []).append(path.name)
    return names


def _read_explicit_mapping(excel_path, source_column, target_column):
    from openpyxl import load_workbook
    from hr_toolkit.common.excel_compat import ensure_xlsx_workbook
    from hr_toolkit.common.run_temp import temporary_directory

    with temporary_directory() as tmp:
        workbook = load_workbook(ensure_xlsx_workbook(Path(excel_path), Path(tmp)), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            reset = getattr(sheet, 'reset_dimensions', None)
            if callable(reset):
                reset()
            source_index = target_index = None
            mapping = {}
            for index, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if source_index is None:
                    headers = [str(value).strip() if value is not None else '' for value in values]
                    if source_column in headers and target_column in headers:
                        if headers.count(source_column) != 1 or headers.count(target_column) != 1 or source_column == target_column:
                            raise ValueError('Excel 映射列必须各有一个且不能相同')
                        source_index, target_index = headers.index(source_column), headers.index(target_column)
                    elif index >= 20:
                        break
                    continue
                source = str(values[source_index] or '').strip() if source_index < len(values) else ''
                target = str(values[target_index] or '').strip() if target_index < len(values) else ''
                if not source and not target:
                    continue
                if not source:
                    raise ValueError(f'Excel 第 {index} 行缺少原文件名')
                legacy._validate_folder_name(source)
                if source.casefold() in mapping:
                    raise ValueError(f'Excel 原文件名重复：{source}')
                mapping[source.casefold()] = target
            if source_index is None:
                raise ValueError(f'Excel 前 20 行未找到唯一的“{source_column}”和“{target_column}”列')
            return mapping
        finally:
            workbook.close()


def build_rename_plan(root_dir, *, mode, text='', target_name='', replacement_name='',
                      file_type='folder', excel_path=None, name_column='姓名', header_row=1,
                      source_column='原文件名', target_column='新名称',
                      trim_spaces=True, collapse_spaces=True, normalize_separators=False,
                      dry_run=True, cancelled=None, **unused):
    legacy._check_cancelled(cancelled)
    root = Path(root_dir).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f'文件夹不存在：{root}')
    if file_type not in legacy.FILE_TYPE_EXTENSIONS:
        raise ValueError(f'不支持的文件类型：{file_type}')
    warnings, names, mapping = [], [], {}
    if mode in ('excel', 'excel_map'):
        if not excel_path:
            raise ValueError('请选择 Excel 名单或映射表')
        candidates = legacy._list_items_for_excel_rename(root, file_type=file_type, file_extensions=None,
                                                        excel_path=Path(excel_path).resolve())
        if mode == 'excel':
            names = legacy._read_names_from_excel(Path(excel_path), name_column, header_row)
            if not names:
                raise ValueError('Excel 中没有可用姓名')
            if len(names) != len(candidates):
                warnings.append(f'Excel 共 {len(names)} 个姓名，所选类型共 {len(candidates)} 项；未配对项目已排除，可手动编辑后勾选。')
        else:
            mapping = _read_explicit_mapping(excel_path, source_column, target_column)
            found = {item.name.casefold() for item in candidates}
            warnings.extend(f'Excel 原文件名未匹配所选类型：{key}' for key in mapping if key not in found)
    elif mode == 'replace':
        target_name, replacement_name = target_name.strip(), replacement_name.strip()
        legacy._validate_folder_name(target_name)
        if not replacement_name:
            raise ValueError('请填写替换后的名称')
        candidates = [root / target_name]
        if not candidates[0].exists():
            raise FileNotFoundError(f'未找到要替换的项目：{target_name}')
    elif mode in ('append', 'remove', 'replace_text', 'normalize'):
        if mode in ('append', 'remove'):
            text, target_name = text.strip(), target_name.strip()
            if target_name:
                legacy._validate_folder_name(target_name)
        if mode != 'normalize' and not text:
            raise ValueError('请填写要处理的原文字')
        if mode == 'replace_text' and not replacement_name:
            raise ValueError('请填写替换后的文字')
        candidates = legacy._iter_target_items(root, target_name if mode in ('append', 'remove') else '', file_type)
    else:
        raise ValueError(f'不支持的改名模式：{mode}')

    rows = []
    for index, source in enumerate(candidates):
        legacy._check_cancelled(cancelled)
        is_dir = source.is_dir()
        suffix = '' if is_dir else source.suffix
        stem = source.name if is_dir else source.stem
        new_name, included, note = source.name, True, ''
        if mode == 'excel':
            included = index < len(names)
            if included:
                new_name = names[index] + suffix
            else:
                note = 'Excel 没有对应姓名'
        elif mode == 'excel_map':
            included = source.name.casefold() in mapping
            if included:
                # Explicit mapping values are complete names, or stems with extension omitted.
                value = mapping[source.name.casefold()]
                new_name = value if not suffix or value.endswith(suffix) else value + suffix
                if not value:
                    new_name = ''
            else:
                note = 'Excel 未提供此项目的映射'
        elif mode == 'replace':
            new_name = replacement_name if not suffix or replacement_name.endswith(suffix) else replacement_name + suffix
        elif mode == 'append':
            new_name = source.name if stem.endswith(text) else stem + text + suffix
        elif mode == 'remove':
            matched = legacy._matching_remove_suffix(stem, legacy._remove_suffix_candidates(text))
            if matched:
                new_name = stem[:-len(matched)] + suffix
                if not stem[:-len(matched)]:
                    new_name = ''
        elif mode == 'replace_text':
            new_name = stem.replace(text, replacement_name) + suffix
        elif mode == 'normalize':
            value = stem.strip() if trim_spaces else stem
            if collapse_spaces:
                value = re.sub(r'[\s\u3000]+', ' ', value)
            if normalize_separators:
                value = re.sub(r'[_＿－-]+', '_', value)
            new_name = value + suffix
        source_error, fingerprint = '', ''
        try:
            fingerprint = _fingerprint(source, cancelled)
        except (OSError, ValueError) as exc:
            source_error = str(exc)
        rows.append({'row_id': str(index), 'order': index + 1, 'source_name': source.name,
                     'relative_path': source.name, 'target_name': new_name, 'is_dir': is_dir,
                     'suffix': suffix, 'included': included, 'fingerprint': fingerprint,
                     'source_error': source_error, 'note': note})
    return validate_plan({'schema': 1, 'root_dir': str(root), 'mode': mode, 'file_type': file_type,
                          'existing_names': _existing(root), 'rows': rows, 'warnings': warnings,
                          'rules': {'text': text, 'replacement': replacement_name, 'target': target_name,
                                    'excel': str(excel_path or ''), 'trim_spaces': trim_spaces,
                                    'collapse_spaces': collapse_spaces, 'normalize_separators': normalize_separators}})


def _journal(stream, event):
    stream.write(json.dumps(event, ensure_ascii=False) + '\n')
    stream.flush()
    os.fsync(stream.fileno())


def execute_rename_plan(root_dir, *, plan, output_dir=None, dry_run=False, cancelled=None, progress_callback=None):
    """Validate and execute ONLY explicit confirmed mappings, including safe name cycles."""
    root = Path(root_dir).resolve()
    if plan.get('schema') != 1:
        raise ValueError('不支持的改名预览版本，请重新预览')
    checked = copy.deepcopy(plan)
    checked['existing_names'] = _existing(root)
    checked = validate_plan(checked)
    if checked['conflict_count']:
        raise ValueError('预览存在冲突，请修正或排除后重试：' + '\n'.join(
            row['source_name'] + '：' + row['issue'] for row in checked['rows'] if row['status'] == 'conflict'))
    for row in checked['rows']:
        if row['included']:
            legacy._check_cancelled(cancelled)
            if _fingerprint(root / row['source_name'], cancelled) != row['fingerprint']:
                raise ValueError(f'{row["source_name"]} 在预览后已变化，请重新预览')
    operations = [row for row in checked['rows'] if row['status'] == 'ready']
    if not operations:
        raise ValueError('没有已选中且需要改名的项目')
    if dry_run:
        return checked
    output = Path(output_dir) if output_dir is not None else root.parent
    output.mkdir(parents=True, exist_ok=True)
    token = uuid4().hex
    ledger = output / ('改名记录_' + token + '.jsonl')
    receipt = output / ('改名清单_' + token + '.csv')
    staging = root / ('.hr-rename-' + token)
    states = {row['row_id']: 'pending' for row in operations}
    def stage_path(row):
        return staging / row['row_id']
    def log_move(stream, row, source, target, state):
        _journal(stream, {'event': 'intent', 'row_id': row['row_id'], 'from': str(source.relative_to(root)),
                          'to': str(target.relative_to(root)), 'state': state})
        legacy._rename_text_no_replace(source, target)
        states[row['row_id']] = state
        _journal(stream, {'event': 'done', 'row_id': row['row_id'], 'state': state})
    failure = None
    with ledger.open('x', encoding='utf-8') as stream:
        _journal(stream, {'event': 'plan', 'schema': 1, 'created_at': datetime.now(timezone.utc).isoformat(),
                          'working_root': str(root), 'staging': staging.name, 'plan': checked})
        staging.mkdir(exist_ok=False)
        try:
            for index, row in enumerate(operations):
                legacy._check_cancelled(cancelled)
                log_move(stream, row, root / row['source_name'], stage_path(row), 'staged')
                if progress_callback:
                    progress_callback(index + 1, len(operations) * 2, f'准备改名 {index + 1}/{len(operations)}')
            for index, row in enumerate(operations):
                legacy._check_cancelled(cancelled)
                log_move(stream, row, stage_path(row), root / row['target_name'], 'completed')
                if progress_callback:
                    progress_callback(len(operations) + index + 1, len(operations) * 2,
                                      f'已改名 {index + 1}/{len(operations)}')
            legacy._check_cancelled(cancelled)
            _journal(stream, {'event': 'finished', 'status': 'success'})
        except BaseException as exc:
            failure = exc
            # Never overwrite a path created by another process while restoring.
            for row in reversed(operations):
                if states[row['row_id']] == 'completed':
                    try:
                        if _fingerprint(root / row['target_name']) != row['fingerprint']:
                            raise ValueError('完成后的项目已变化，保留现场')
                        log_move(stream, row, root / row['target_name'], stage_path(row), 'staged')
                    except (OSError, ValueError):
                        pass  # Ledger plus retained paths provide the recovery evidence.
            for row in operations:
                if states[row['row_id']] == 'staged':
                    try:
                        log_move(stream, row, stage_path(row), root / row['source_name'], 'restored')
                    except OSError:
                        pass
            try:
                _journal(stream, {'event': 'finished', 'status': 'failed', 'error': str(exc), 'states': states})
            except OSError:
                pass  # Keep the last durable intent and every retained data path.
        finally:
            try:
                staging.rmdir()  # Only the empty, tool-created staging directory can be removed.
            except OSError:
                pass
    try:
        with receipt.open('x', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(['原名称', '确认名称', '相对路径', '执行状态'])
            for row in checked['rows']:
                # Escape spreadsheet formulas without changing the authoritative JSON journal.
                def cell(value):
                    return "'" + value if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value
                status = states.get(row['row_id'], row['status'])
                status = {'completed': '已完成', 'pending': '未执行', 'staged': '已暂存待恢复',
                          'restored': '已恢复原名', 'unchanged': '名称不变', 'excluded': '已排除'}.get(status, status)
                writer.writerow([cell(row['source_name']), cell(row['target_name']), cell(row['relative_path']), status])
    except OSError as exc:
        if failure is None:
            raise RuntimeError(f'改名已完成，但 CSV 清单保存失败。请查看完整改名记录：{ledger}。原因：{exc}') from exc
    if failure is not None:
        retained = sum(state in {'staged', 'completed'} for state in states.values())
        raise RuntimeError(f'改名未全部完成，已尝试安全恢复原名；{retained} 项尚未恢复。'
                           f'请保留结果目录（含临时项目）并查看改名记录：{ledger}。原因：{failure}') from failure
    return {'mode': checked['mode'], 'operation_count': len(operations),
            'excluded_count': sum(not row['included'] for row in checked['rows']),
            'unchanged_count': sum(row['status'] == 'unchanged' for row in checked['rows']),
            'operations': [{'source': str(root / row['source_name']), 'target': str(root / row['target_name'])}
                           for row in operations], 'warnings': checked['warnings'],
            'rename_ledger': str(ledger), 'rename_receipt': str(receipt), 'renamed_root': str(root)}
