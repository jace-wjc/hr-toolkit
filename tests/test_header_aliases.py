"""名称规则只改变输入定位，不改变工资计算和输出布局。"""
from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from hr_toolkit.common.header_aliases import matching_columns
from hr_toolkit.tools.salary_headers import (
    ALIAS_PROFILE_KEY, alias_rules, inspect_workbook, profile_from_selection,
)
from hr_toolkit.tools.salary_merge import _detect_source_layout, inspect_salary_templates, merge_monthly_salary


def rules(*, names=None, sheets=None):
    return {ALIAS_PROFILE_KEY: {"fields": {"name": names or ["姓名", "名字", "name"]},
                                "sheets": {"detail": sheets} if sheets else {}}}


def workbook(name="姓名", sheet="工资明细", amount="应发小计"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["序号", name, "身份证号码", amount])
    ws.append([1, "测试人员", "TEST-001", 1250.5])
    return wb


class HeaderAliasTest(unittest.TestCase):
    def test_alternatives_match_one_name_per_template(self):
        for name in ("姓名", "名字", " ＮＡＭＥ "):
            with self.subTest(name=name):
                group = inspect_workbook(workbook(name), profiles=rules())
                self.assertTrue(group["ready"])
                self.assertEqual(group["selections"], {"name": 2, "id_card": 3, "amount": 4})

    def test_missing_name_requires_confirmation(self):
        group = inspect_workbook(workbook("员工称呼"), profiles=rules())
        self.assertFalse(group["ready"])
        self.assertEqual(group["selections"]["name"], 0)

    def test_multiple_matching_columns_require_explicit_choice(self):
        wb = workbook()
        wb.active["E1"] = "名字"
        group = inspect_workbook(wb, profiles=rules())
        self.assertFalse(group["ready"])
        profile = profile_from_selection(group, {"name": 5, "id_card": 3, "amount": 4})
        saved = {**rules(), group["key"]: profile}
        result = inspect_workbook(wb, profiles=saved)
        self.assertTrue(result["ready"])
        self.assertEqual(result["selections"]["name"], 5)
        # 调整名称规则后不能继续沿用旧的手选列。
        saved[ALIAS_PROFILE_KEY]["fields"]["name"] = ["姓名"]
        self.assertEqual(inspect_workbook(wb, profiles=saved)["selections"]["name"], 2)

    def test_normalization_deduplicates_and_rejects_cross_field_conflict(self):
        self.assertEqual(alias_rules(rules(names=["name", "ＮＡＭＥ", "Name"]))["fields"]["name"], ["姓名", "name"])
        with self.assertRaises(ValueError):
            alias_rules(rules(names=["身份证号码"]))
        columns = [{"column": 1, "label": "姓名 / 名字", "leaves": ["姓名", "名字"]}]
        self.assertEqual(matching_columns(columns, ["姓名", "名字"]), [1])

    def test_summary_requires_only_its_two_fields(self):
        wb = workbook("名字", sheet="汇总")
        wb.active.delete_cols(4)
        group = inspect_workbook(wb, role="summary", profiles=rules())
        self.assertTrue(group["ready"])
        self.assertEqual(set(group["selections"]), {"name", "id_card"})

    def test_sheet_alternatives_missing_and_ambiguous(self):
        profiles = rules(sheets=["工资明细", "Sheet1", "工资数据"])
        wb = workbook("名字", sheet="Sheet1")
        self.assertTrue(inspect_workbook(wb, profiles=profiles)["ready"])
        wb.active.title = "新工资页"
        missing = inspect_workbook(wb, profiles=profiles)
        self.assertTrue(missing["sheet_needs_confirmation"])
        with self.assertRaises(ValueError):
            profile_from_selection(missing, missing["selections"])
        wb.active.title = "Sheet1"
        extra = wb.copy_worksheet(wb.active)
        extra.title = "工资明细"
        self.assertTrue(inspect_workbook(wb, profiles=profiles)["sheet_needs_confirmation"])
        confirmed = inspect_workbook(wb, profiles=profiles, hint={"sheet": "Sheet1", "header_row": 1})
        self.assertTrue(confirmed["ready"])
        self.assertFalse(confirmed.get("sheet_needs_confirmation", False))

    def test_sheet_name_can_be_shared_by_separately_supplied_roles(self):
        result = alias_rules({ALIAS_PROFILE_KEY: {"sheets": {"detail": ["Sheet1"], "summary": ["Sheet1"]}}})
        self.assertEqual(result["sheets"]["detail"], ["明细", "Sheet1"])

    def test_header_position_and_column_reordering(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "工资明细"
        ws.cell(40, 1, "身份证号码")
        ws.cell(40, 2, "应发小计")
        ws.cell(40, 3, "name")
        result = inspect_workbook(wb, profiles=rules())
        self.assertTrue(result["ready"])
        self.assertEqual(result["header_row"], 40)
        self.assertEqual(result["selections"], {"name": 3, "id_card": 1, "amount": 2})

    def test_original_saved_profile_still_works_without_rules(self):
        wb = workbook("原表姓名", sheet="自定义工资")
        group = inspect_workbook(wb)
        profile = profile_from_selection(group, {"name": 2, "id_card": 3, "amount": 4})
        profile.pop("alias_signature")  # 升级前已保存的设置没有这个键。
        self.assertTrue(inspect_workbook(wb, profiles={group["key"]: profile})["saved"])

    def test_original_amount_fallback_unchanged_and_custom_amount_is_strict(self):
        wb = workbook("名字")
        self.assertTrue(_detect_source_layout(wb, header_profiles=rules()).legacy_amount_fallback)
        wb.active["D1"] = "工资合计"
        profiles = rules()
        profiles[ALIAS_PROFILE_KEY]["fields"]["amount"] = ["工资合计"]
        self.assertFalse(_detect_source_layout(wb, header_profiles=profiles).legacy_amount_fallback)

    def test_sheet_inventory_differences_are_confirmed_separately(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first, second = root / "甲_202608.xlsx", root / "乙_202608.xlsx"
            wb = workbook(sheet="Sheet1")
            wb.save(first)
            wb.copy_worksheet(wb.active).title = "工资明细"
            wb.save(second)
            inspected = inspect_salary_templates([first, second], header_profiles=rules(sheets=["Sheet1", "工资明细"]))
            self.assertFalse(inspected["issues"])
            self.assertEqual(len(inspected["groups"]), 2)
            self.assertEqual(sum(g["ready"] for g in inspected["groups"]), 1)

    def test_standard_template_output_and_sources_unchanged(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "工资_202608.xlsx"
            workbook().save(source)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            baseline = merge_monthly_salary(source, root / "before", header_profiles={}, strict_headers=True)
            revised = merge_monthly_salary(source, root / "after", header_profiles=rules(), strict_headers=True)
            self.assertEqual(baseline.record_count, revised.record_count)
            self.assertEqual(baseline.warnings, revised.warnings)
            before, after = load_workbook(baseline.output_file), load_workbook(revised.output_file)
            try:
                self.assertEqual(before.sheetnames, after.sheetnames)
                for name in before.sheetnames:
                    left, right = before[name], after[name]
                    self.assertEqual(left.calculate_dimension(), right.calculate_dimension())
                    self.assertEqual(str(left.merged_cells), str(right.merged_cells))
                    self.assertEqual(left.freeze_panes, right.freeze_panes)
                    for row in left:
                        for cell in row:
                            other = right[cell.coordinate]
                            self.assertEqual((cell.value, cell.data_type, cell._style),
                                             (other.value, other.data_type, other._style))
            finally:
                before.close()
                after.close()
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)


if __name__ == "__main__":
    unittest.main()
