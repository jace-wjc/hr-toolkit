from __future__ import annotations

import io
import json
import os
import runpy
import tempfile
import threading
import time
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("HR_TOOLKIT_SKIP_UPDATE", "1")

try:
    from hr_toolkit.gui_qt.compat import QCoreApplication
    from hr_toolkit.gui_qt.controller import AppController
except ImportError:
    AppController = None

from hr_toolkit.run_coordinator import ProjectRunCoordinator, RunCallbacks, RunRequest
from hr_toolkit.background_process import PROCESS_FILE_THRESHOLD_BYTES, PROCESS_WORKSHEET_THRESHOLD_BYTES, should_use_process
from hr_toolkit.tools import material_collector as mc


class SalaryMergePerformanceTests(unittest.TestCase):
    def test_releasing_source_cells_preserves_cross_workbook_snapshot_styles(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from hr_toolkit.common.excel import apply_row_snapshot, snapshot_row
        from hr_toolkit.tools.salary_split import _release_workbook_cells

        source = Workbook()
        source.active["A1"] = "=SUM(B1:C1)"
        source.active["A1"].font = Font(name="宋体", bold=True, color="FF0088")
        source.active["A1"].fill = PatternFill("solid", fgColor="FFFF00")
        source.active["A1"].number_format = "0.000"
        source.active.row_dimensions[1].height = 29
        snapshot = snapshot_row(source.active, 1, 3)
        outputs = []
        for release in (False, True):
            if release:
                _release_workbook_cells(source)
                self.assertEqual(len(source.active._cells), 0)
            destination = Workbook()
            # Force different style indices in the destination workbook.
            destination.active["E1"].font = Font(name="Arial", italic=True)
            apply_row_snapshot(destination.active, 2, snapshot)
            output = io.BytesIO()
            destination.save(output)
            with zipfile.ZipFile(output) as archive:
                outputs.append({name: archive.read(name) for name in archive.namelist()
                                if name != "docProps/core.xml"})
            destination.close()
        self.assertEqual(outputs[0], outputs[1])

    def test_indexed_merges_preserve_openpyxl_xml_and_borders(self):
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Border, Protection, Side
        from hr_toolkit.tools.salary_split import _RowMergeWriter

        packages = []
        for indexed in (False, True):
            workbook = Workbook()
            worksheet = workbook.active
            for row in range(1, 20):
                for column in range(1, 9):
                    cell = worksheet.cell(row, column, f"{row}-{column}")
                    cell.border = Border(right=Side(style="thin"), bottom=Side(style="double"))
                    cell.protection = Protection(locked=False)
                    cell.comment = Comment("合并前备注", "测试")
            worksheet.merge_cells("A1:D3")
            writer = _RowMergeWriter(worksheet) if indexed else None
            # Header containment, duplicate/subset/overlapping ranges and
            # independent subtotal rows all use openpyxl's original semantics.
            for row, left, right in [(2, 2, 3), (5, 1, 4), (5, 1, 4), (5, 2, 3),
                                     (5, 3, 6), (6, 1, 4), (7, 2, 7), (8, 1, 1)]:
                if writer:
                    writer.merge(row, left, right)
                else:
                    worksheet.merge_cells(start_row=row, end_row=row, start_column=left, end_column=right)
            output = io.BytesIO()
            workbook.save(output)
            with zipfile.ZipFile(output) as archive:
                packages.append({name: archive.read(name) for name in archive.namelist()
                                 if name != "docProps/core.xml"})
            workbook.close()
        self.assertEqual(packages[0], packages[1])


class PackagingPerformanceTests(unittest.TestCase):
    def test_incremental_windows_build_preserves_release_gates(self):
        from scripts import build_windows
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            options = dict(version=build_windows.read_project_version(), output_dir=root / "dist",
                           work_dir=root / "build", version_file=root / "version.txt")
            for option, names in (("seven_zip_dir", build_windows.WIN7_REQUIRED_7ZIP_FILES),
                                  ("ucrt_dir", build_windows.WIN7_REQUIRED_UCRT_FILES),
                                  ("vc_runtime_dir", build_windows.WIN7_REQUIRED_VC_RUNTIME_FILES)):
                directory = root / option
                directory.mkdir()
                for name in names:
                    (directory / name).write_bytes(b"runtime fixture")
                options[option] = directory
            for target in (build_windows.WINDOWS_TARGET_MODERN, build_windows.WINDOWS_TARGET_WIN7):
                normal = build_windows.pyinstaller_commands(**options, target=target)
                incremental = build_windows.pyinstaller_commands(**options, target=target, clean=False)
                for before, after in zip(normal, incremental):
                    self.assertIn("--clean", before)
                    self.assertEqual([arg for arg in before if arg != "--clean"], after)

    def test_qml_filter_runs_before_plugin_analysis_and_keeps_runtime_files(self):
        helper = Path(__file__).resolve().parents[1] / "packaging/qt/hooks/qml_payload.py"
        collect = runpy.run_path(str(helper))["collect_required_qml_files"]
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            for module in ("QtQuick", "UnusedWebEngine"):
                directory = root / module
                directory.mkdir()
                (directory / "qmldir").touch()
                (directory / "plugins.qmltypes").touch()
            processed = []

            def process(qmldir):
                processed.append(qmldir.parent.name)
                return [qmldir.parent / "plugin.dll"], [qmldir, qmldir.parent / "plugins.qmltypes"]

            info = SimpleNamespace(version=[6, 6], location={"QmlImportsPath": str(root)},
                                   qt_rel_dir="PySide6/Qt", _process_qml_plugin=process)
            binaries, data = collect(info, lambda entry: entry[1].replace("\\", "/").endswith("/QtQuick"))
            self.assertEqual(processed, ["QtQuick"])
            destination = str(Path("PySide6/Qt/qml/QtQuick"))
            self.assertEqual(binaries, [(str(root / "QtQuick/plugin.dll"), destination)])
            self.assertEqual(data, [(str(root / "QtQuick/qmldir"), destination)])


class WindowsFileChangeTests(unittest.TestCase):
    def tearDown(self):
        mc._windows_file_change_reader.cache_clear()

    def test_bindings_are_reused_but_each_file_change_is_queried(self):
        import ctypes

        tokens = iter([101, 202])
        def query(_handle, _kind, info, _size):
            info._obj.ChangeTime = next(tokens)
            return 1
        kernel = SimpleNamespace(CreateFileW=Mock(return_value=123),
                                 GetFileInformationByHandleEx=Mock(side_effect=query),
                                 CloseHandle=Mock(return_value=1))
        source = Path("same-file.png")
        mc._windows_file_change_reader.cache_clear()
        with patch.object(ctypes, "WinDLL", return_value=kernel, create=True) as dll, patch.object(mc.os, "name", "nt"):
            self.assertEqual(mc._windows_file_change_time(source), 101)
            self.assertEqual(mc._windows_file_change_time(source), 202)
        self.assertEqual(dll.call_count, 1)
        self.assertEqual(kernel.CreateFileW.call_count, 2)
        self.assertEqual(kernel.GetFileInformationByHandleEx.call_count, 2)
        self.assertEqual(kernel.CloseHandle.call_count, 2)

    def test_query_failures_close_handles_and_do_not_cache_file_results(self):
        import ctypes
        from ctypes import wintypes

        kernel = SimpleNamespace(CreateFileW=Mock(side_effect=[wintypes.HANDLE(-1).value, 123, 124]),
                                 GetFileInformationByHandleEx=Mock(side_effect=[0, OSError("无法查询")]),
                                 CloseHandle=Mock(return_value=1))
        source = Path("unreadable.png")
        mc._windows_file_change_reader.cache_clear()
        with patch.object(ctypes, "WinDLL", return_value=kernel, create=True), patch.object(mc.os, "name", "nt"):
            for _ in range(3):
                self.assertIsNone(mc._windows_file_change_time(source))
        self.assertEqual(kernel.CreateFileW.call_count, 3)
        self.assertEqual(kernel.GetFileInformationByHandleEx.call_count, 2)
        self.assertEqual([call.args[0] for call in kernel.CloseHandle.call_args_list], [123, 124])


class CachePerformanceTests(unittest.TestCase):
    def test_checkpoint_timestamp_reuse_preserves_expiry_and_invalid_value_fallback(self):
        from datetime import datetime, timedelta
        now = datetime.now(tz=mc._BEIJING_TZ)
        fresh = now.strftime("%Y-%m-%d %H:%M:%S")
        old = (now - timedelta(days=91)).strftime("%Y-%m-%d %H:%M:%S")
        data = mc._new_ocr_cache()
        data["entries"] = {str(index): {"verified_at": fresh} for index in range(1000)}
        data["entries"].update({
            "expired": {"verified_at": old},
            "expired_iso": {"verified_at": (now - timedelta(days=91)).isoformat()},
            "fresh_iso": {"verified_at": now.isoformat()},
            "invalid_string": {"verified_at": "无效时间"},
            "invalid_list": {"verified_at": [fresh]},
            "missing": {},
        })
        data["paths"] = {key + ".jpg": {"cache_key": key} for key in data["entries"]}
        expected = set(data["entries"]) - {"expired", "expired_iso", "missing"}
        mc._trim_cache_by_age_and_size(data)
        self.assertEqual(set(data["entries"]), expected)
        self.assertEqual(set(data["paths"]), {key + ".jpg" for key in expected})

    def test_chunked_cache_json_is_byte_identical_to_original_serializer(self):
        data = {"version": 5, "entries": {
            f"员工-{index}": {"ocr_text": '姓名：“测试”\n\\', "score": 0.123, "flag": True, "names": [None, "张三"]}
            for index in range(5000)
        }, "paths": {"照片.jpg": {"source_mtime_ns": 1234567890123456789}}, "extra": [1, False, "中文"]}
        output = io.StringIO()
        mc._write_ocr_cache_json(output, data)
        self.assertEqual(output.getvalue(), json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n")

    def test_failed_checkpoint_keeps_original_and_removes_partial(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "cache.json"
            original = b'{"original":true}'
            path.write_bytes(original)

            def fail(handle, _data):
                handle.write("partial")
                raise OSError("disk full")

            with patch.object(mc, "_write_ocr_cache_json", side_effect=fail):
                self.assertFalse(mc._save_ocr_cache(path, mc._new_ocr_cache()))
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(path.parent.iterdir()), [path])

    def test_filename_rule_order_and_custom_materials_are_preserved(self):
        cases = [
            ("张三身份证正面.jpg", ["身份证"], "身份证"),
            ("id.jpg", ["身份证"], None),
            ("张三资质证明扫描.jpg", ["资质证明", "证明"], "资质证明"),
            ("张三资质证明扫描.jpg", ["证明", "资质证明"], "资质证明"),
            ("000012.jpg", ["证件照片"], None),
        ]
        for filename, requested, expected in cases:
            with self.subTest(filename=filename, requested=requested):
                self.assertEqual(mc._classify_material_from_filename(filename, requested)[0], expected)


@unittest.skipUnless(AppController is not None, "Qt runtime is not installed")
class GuiPerformanceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QCoreApplication.instance() or QCoreApplication([])

    def setUp(self):
        self.controller = AppController()
        self.controller._save_workspace_preferences = lambda: True

    def tearDown(self):
        self.controller.close()
        self.app.processEvents()

    def pump_until(self, condition):
        deadline = time.monotonic() + 3
        while not condition() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.001)
        self.app.processEvents()
        self.assertTrue(condition())

    def test_slow_metadata_is_off_gui_and_cannot_restore_cleared_inputs(self):
        entered, release = threading.Event(), threading.Event()
        threads = []
        gui_thread = threading.get_ident()

        def stat(_path):
            threads.append(threading.get_ident())
            entered.set()
            release.wait(2)
            return True

        try:
            with patch.object(Path, "is_dir", stat):
                self.controller._set_inputs([Path("/synthetic-folder")], replace=True)
                self.assertTrue(entered.wait(1))
                self.controller.clearInputs()
                release.set()
                self.pump_until(lambda: not self.controller._input_scan_running)
        finally:
            release.set()
        self.assertEqual(self.controller._input_model.rowCount(), 0)
        self.assertFalse(self.controller._input_metadata)
        self.assertNotIn(gui_thread, threads)

    def test_rapid_reselection_keeps_one_metadata_worker_and_latest_order(self):
        entered, release = threading.Event(), threading.Event()
        calls = []

        def stat(path):
            calls.append((str(path), threading.get_ident()))
            entered.set()
            release.wait(2)
            return True

        try:
            with patch.object(Path, "is_dir", stat):
                self.controller._set_inputs([Path("/first")], replace=True)
                self.assertTrue(entered.wait(1))
                for index in range(20):
                    self.controller._set_inputs([Path(f"/next-{index}")], replace=True)
                release.set()
                self.pump_until(lambda: not self.controller._input_scan_running)
        finally:
            release.set()
        self.assertEqual(len({thread for _path, thread in calls}), 1)
        self.assertEqual([item["path"] for item in self.controller._input_model.items()], [str(Path("/next-19"))])
        self.assertEqual(self.controller._input_model.item_at(0)["kind"], "folder")

    def test_cancel_during_slow_validation_never_launches_business(self):
        entered, release = threading.Event(), threading.Event()
        threads = []
        launched = []
        self.controller._project_path = Path("/synthetic-project")

        def validate(*_args, **_kwargs):
            threads.append(threading.get_ident())
            entered.set()
            release.wait(2)
            return SimpleNamespace(tool_id="social_security", preview=False)

        self.controller._start_project_run = launched.append
        try:
            with patch("hr_toolkit.gui_qt.controller.build_invocation", side_effect=validate):
                self.controller._prepare_invocation(preview=False)
                self.assertTrue(entered.wait(1))
                self.controller.runOrCancel()
                release.set()
                self.pump_until(lambda: not self.controller.busy)
        finally:
            release.set()
        self.assertEqual(launched, [])
        self.assertNotIn(threading.get_ident(), threads)

    def test_progress_storm_is_bounded_and_final_completion_is_retained(self):
        self.controller._run_progress_visible = True
        changes = []
        self.controller.runProgressChanged.connect(lambda: changes.append(1))

        def produce():
            for index in range(20000):
                self.controller._queue_run_progress(1, 1, f"【阶段{index}】完成")

        worker = threading.Thread(target=produce)
        worker.start()
        worker.join(2)
        self.assertFalse(worker.is_alive())
        self.assertEqual(changes, [])
        self.controller._drain_run_progress()
        self.controller._flush_material_progress()
        self.assertEqual(len(changes), 1)
        self.assertEqual(self.controller._run_progress_message, "【阶段19999】完成")

    def test_startup_disk_checks_do_not_block_gui_or_overwrite_unread_settings(self):
        entered, release = threading.Event(), threading.Event()
        original_is_dir = Path.is_dir
        threads = []
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            settings = root / "workspace-ui.json"
            remote = root / "remote"
            original = json.dumps({"recent_projects": [str(remote)], "last_selected_dir": str(remote)})
            settings.write_text(original, encoding="utf-8")

            def stat(path):
                if path == remote:
                    threads.append(threading.get_ident())
                    entered.set()
                    release.wait(2)
                    return True
                return original_is_dir(path)

            try:
                with patch.object(AppController, "_settings_path", return_value=settings), patch.object(Path, "is_dir", stat):
                    self.controller.start()
                    self.assertTrue(entered.wait(1))
                    self.assertFalse(AppController._save_workspace_preferences(self.controller))
                    self.controller.close()
                    self.assertEqual(settings.read_text(encoding="utf-8"), original)
                    release.set()
                    self.pump_until(lambda: not self.controller.busy)
            finally:
                release.set()
        self.assertNotIn(threading.get_ident(), threads)

    def test_workspace_queue_bounds_threads_and_discards_obsolete_reads(self):
        release = threading.Event()
        started = []
        self.controller._workspace_read_limit = 2

        def read():
            started.append(threading.get_ident())
            release.wait(2)

        try:
            for _ in range(50):
                self.controller._schedule_workspace_read(0, read)
            self.pump_until(lambda: len(started) == 2)
            self.assertEqual(self.controller._workspace_read_workers, 2)
            self.controller._workspace_generation = 1
            release.set()
            self.pump_until(lambda: self.controller._workspace_read_workers == 0)
        finally:
            release.set()
        self.assertEqual(len(started), 2)

    def test_directory_result_survives_another_folders_row_insertion(self):
        self.controller._workspace_items = [
            {"path": "/first", "depth": 0, "expanded": True},
            {"path": "/first/child", "depth": 1, "expanded": False},
            {"path": "/second", "depth": 0, "expanded": True},
        ]
        self.controller._workspace_model.set_items(self.controller._workspace_items)
        self.controller._apply_workspace_children(0, 1, "/second", 0, [{"path": "/second/child", "depth": 1}])
        self.assertEqual(self.controller._workspace_items[-1]["path"], "/second/child")

    def test_packaging_smoke_never_reads_or_rewrites_user_settings(self):
        from importlib import import_module
        from hr_toolkit.gui_qt import smoke
        main = import_module("hr_toolkit.gui_qt.main")
        with tempfile.TemporaryDirectory() as temporary:
            user_settings = Path(temporary) / "settings.json"
            user_settings.write_text('{"current_project":"user-project"}', encoding="utf-8")
            seen = []

            def fake_main():
                path = AppController._settings_path()
                seen.append(path)
                self.assertNotEqual(path, user_settings)
                self.assertFalse(path.exists())
                path.write_text("{}", encoding="utf-8")
                return 0

            with patch.object(AppController, "_settings_path", staticmethod(lambda: user_settings)), \
                    patch.object(main, "main", side_effect=fake_main), \
                    patch.dict(os.environ), patch.object(smoke.sys, "argv", ["smoke"]):
                self.assertEqual(smoke.run(), 0)
                self.assertEqual(AppController._settings_path(), user_settings)
            self.assertEqual(user_settings.read_text(encoding="utf-8"), '{"current_project":"user-project"}')
            self.assertFalse(seen[0].parent.exists())


class LazyToolImportTests(unittest.TestCase):
    def test_small_compressed_workbooks_and_multi_file_archives_are_isolated(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook = root / "large.xlsx"
            with zipfile.ZipFile(workbook, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("xl/worksheets/sheet1.xml", b" " * PROCESS_WORKSHEET_THRESHOLD_BYTES)
            self.assertLess(workbook.stat().st_size, PROCESS_FILE_THRESHOLD_BYTES // 100)
            with patch.object(zipfile.ZipFile, "read", side_effect=AssertionError("must not decompress input")):
                self.assertTrue(should_use_process("salary_split", (workbook,), {}))
            bundle = root / "batch.zip"
            with zipfile.ZipFile(bundle, "w") as archive:
                for index in range(5):
                    archive.writestr(f"{index}.xlsx", b"fixture")
            self.assertTrue(should_use_process("salary_merge", (bundle,), {}))

    def test_import_failure_is_reported_and_finished_without_project_mutation(self):
        errors, finished, threads = [], [], []
        request = RunRequest("probe", "导入测试", "测试", "测试", None, (), {}, "missing_tool", "run")
        store = SimpleNamespace(create_draft=lambda **kwargs: self.fail("must not mutate project"))

        def fail(_name):
            threads.append(threading.get_ident())
            raise ImportError("dependency unavailable")

        coordinator = ProjectRunCoordinator()
        with patch("hr_toolkit.run_coordinator.importlib.import_module", side_effect=fail):
            self.assertTrue(coordinator.start(store, request, RunCallbacks(error=errors.append, finished=lambda: finished.append(True))))
            self.assertTrue(coordinator.wait(2))
        self.assertEqual(len(errors), 1)
        self.assertEqual(finished, [True])
        self.assertNotIn(threading.get_ident(), threads)
        self.assertFalse(coordinator.running)
