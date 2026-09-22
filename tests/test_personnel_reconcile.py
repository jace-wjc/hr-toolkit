"""Synthetic reconciliation fixtures; never use client personnel data."""
from __future__ import annotations

from datetime import date
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from hr_toolkit.tools.personnel_reconcile import (
    company_mapping, parse_company_aliases, reconcile_personnel_changes,
)


ID_A = "110101199001010011"
ID_B = "110101199001010029"
ID_C = "110101199001010037"
SUMMARY_HEADERS = ["公司", "姓名", "身份证号码", "工号", "入职日期", "离职日期",
                   "学历", "毕业学校", "专业", "岗位", "家庭住址", "联系方式"]
JOIN_HEADERS = ["入职公司", "姓名", "身份证号", "工号", "入职日期", "最高学历文化程度",
                "最高学历毕业学校", "最高学历所学专业", "岗位", "户籍地址", "手机号码",
                "第一学历文化程度", "职位", "现住址", "人员状态", "办结时间"]
LEAVE_HEADERS = ["入职公司", "姓名", "身份证号", "工号", "入职日期", "离职日期",
                 "预计离职日期", "岗位", "人员状态", "办结时间"]


def summary_row(identity=ID_A, **changes):
    return {"公司": "唐人", "姓名": "测试甲", "身份证号码": identity,
            "入职日期": date(2026, 7, 1), **changes}


def join_row(identity=ID_A, **changes):
    return {"入职公司": "浙江唐人人力资源有限公司", "姓名": "测试甲", "身份证号": identity,
            "入职日期": date(2026, 7, 1), "最高学历文化程度": "本科",
            "最高学历毕业学校": "测试大学", "最高学历所学专业": "测试专业",
            "岗位": "测试岗位", "户籍地址": "测试户籍地址", "手机号码": "测试电话",
            "第一学历文化程度": "中专", "职位": "不应读取的职位", "现住址": "不应读取的住址",
            "人员状态": "退回", "办结时间": None, **changes}


class PersonnelReconcileTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.summary = self.root / "文件名不代表月份.xlsx"
        self.flow = self.root / "流程.xlsx"
        self.output = self.root / "results"

    def write_summary(self, joins=(), leaves=()):
        book = Workbook()
        book.remove(book.active)
        for name, rows in (("增员", joins), ("减员", leaves)):
            ws = book.create_sheet(name)
            ws.append(["异动汇总表"])
            ws.append(SUMMARY_HEADERS)
            for row in rows:
                ws.append([row.get(field) for field in SUMMARY_HEADERS])
        book.create_sheet("调动")["A1"] = "原有内容"
        book.save(self.summary)
        book.close()

    def write_flow(self, rows, *, leave=False, status=False, headers=None):
        book = Workbook()
        ws = book.active
        ws.title = "Sheet1"
        headers = headers or (LEAVE_HEADERS if leave else JOIN_HEADERS)
        if status:
            headers = [*headers, "流程状态"]
        ws.append(headers)
        for row in rows:
            ws.append([row.get(field) for field in headers])
        book.save(self.flow)
        book.close()

    def run_check(self, **changes):
        options = dict(month="2026-07")
        options.update(changes)
        return reconcile_personnel_changes([self.flow], self.summary, self.output, **options)

    def read_result(self, result):
        report = load_workbook(result.output_file, data_only=False)
        filled = load_workbook(result.filled_output_file, data_only=False)
        self.addCleanup(report.close)
        self.addCleanup(filled.close)
        notices = list(report["预警明细"].values)[1:]
        return notices, filled, report

    def test_fill_uses_highest_education_and_never_edits_original_or_existing_values(self):
        self.write_summary([summary_row(岗位="原岗位", 联系方式="=1+1")])
        self.write_flow([join_row()])
        originals = self.summary.read_bytes(), self.flow.read_bytes()
        result = self.run_check(highlight=False, template_rules={})
        notices, filled, report = self.read_result(result)
        ws = filled["增员"]
        self.assertEqual(ws["G3"].value, "本科")
        self.assertEqual(ws["H3"].value, "测试大学")
        self.assertEqual(ws["I3"].value, "测试专业")
        self.assertEqual(ws["J3"].value, "原岗位")
        self.assertEqual(ws["K3"].value, "测试户籍地址")
        self.assertEqual(ws["L3"].value, "=1+1")
        self.assertEqual(ws["L3"].data_type, "f")
        self.assertEqual(ws["G3"].fill.patternType, None)
        self.assertEqual(filled["调动"]["A1"].value, "原有内容")
        self.assertTrue(any(row[0] == "字段差异" for row in notices))
        self.assertEqual(report["补入记录"].max_row, result.filled_count + 1)
        self.assertEqual(originals, (self.summary.read_bytes(), self.flow.read_bytes()))

    def test_missing_status_logs_once_and_continues_without_confirmation(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row(), join_row(ID_B, 姓名="测试乙")])
        progress = []
        result = self.run_check(progress_callback=lambda *args: progress.append(args))
        self.assertGreater(result.filled_count, 0)
        self.assertEqual(len(progress), 1)
        self.assertIn("未进行状态筛选", progress[0][2])
        self.assertIn("办结时间不作为状态依据", progress[0][2])
        self.assertFalse(any("未进行状态筛选" in warning for warning in result.warnings))
        notices, _, report = self.read_result(result)
        self.assertTrue(any(row[0] == "流程有、异动表无" for row in notices))
        info = dict(list(report["核对说明"].values)[1:])
        self.assertIn("未进行状态筛选", info["入职流程状态来源"])
        self.assertNotIn("用户已确认", info["核对范围"])

    def test_pending_status_is_valid_and_only_named_exclusions_are_removed(self):
        self.write_summary([summary_row(), summary_row(ID_B, 姓名="测试乙"), summary_row(ID_C, 姓名="测试丙")])
        self.write_flow([join_row(流程状态="审批中"), join_row(ID_B, 姓名="测试乙", 流程状态="未发起"),
                         join_row(ID_C, 姓名="测试丙", 流程状态="退回")], status=True)
        progress = []
        result = self.run_check(progress_callback=lambda *args: progress.append(args))
        self.assertEqual(progress, [])
        notices, _, _ = self.read_result(result)
        self.assertEqual(result.matched_count, 1)
        self.assertEqual(sum(row[0] == "异动表有、流程无" for row in notices), 2)

    def test_both_directions_are_reported(self):
        self.write_summary([summary_row(ID_B, 姓名="测试乙")])
        self.write_flow([join_row()])
        notices, _, _ = self.read_result(self.run_check())
        self.assertTrue({"异动表有、流程无", "流程有、异动表无"}.issubset({row[0] for row in notices}))

    def test_mixed_exports_still_filter_records_with_status(self):
        self.write_summary([summary_row(), summary_row(ID_B, 姓名="测试乙"),
                            summary_row(ID_C, 姓名="测试丙")])
        self.write_flow([join_row()])
        no_status = self.root / "无状态流程.xlsx"
        self.flow.rename(no_status)
        self.write_flow([join_row(ID_B, 姓名="测试乙", 流程状态="审批中"),
                         join_row(ID_C, 姓名="测试丙", 流程状态="退回")], status=True)
        progress = []
        result = reconcile_personnel_changes(
            [no_status, self.flow], self.summary, self.output, month="2026-07",
            progress_callback=lambda *args: progress.append(args))
        notices, _, report = self.read_result(result)
        self.assertEqual(result.matched_count, 2)
        self.assertEqual(len(progress), 1)
        missing = [row for row in notices if row[0] == "异动表有、流程无"]
        self.assertEqual(len(missing), 1)
        self.assertEqual(missing[0][3], ID_C)
        info = dict(list(report["核对说明"].values)[1:])
        self.assertIn("按状态列排除未发起/退回 1 条", info["入职范围内记录"])

    def test_duplicate_admissions_with_different_dates_are_not_auto_selected(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row(), join_row(入职日期=date(2026, 7, 2))])
        result = self.run_check()
        notices, filled, _ = self.read_result(result)
        self.assertEqual(result.filled_count, 0)
        self.assertIsNone(filled["增员"]["G3"].value)
        self.assertEqual(sum(row[0] == "重复流程待确认" for row in notices), 2)

    def test_expected_departure_requires_explicit_choice(self):
        self.write_summary(leaves=[summary_row(离职日期=date(2026, 7, 31))])
        self.write_flow([join_row(离职日期=None, 预计离职日期=date(2026, 7, 31))], leave=True)
        result = self.run_check()
        notices, _, _ = self.read_result(result)
        self.assertEqual(result.matched_count, 0)
        self.assertTrue(any(row[0] == "流程待确认" for row in notices))
        self.assertFalse(any(row[0] == "异动表有、流程无" for row in notices))
        self.assertEqual(self.run_check(leave_date_field="预计离职日期").matched_count, 1)

    def test_company_keywords_preserve_geography_and_report_ambiguity(self):
        mapping, ambiguous = company_mapping(["北京春苗", "唐人"],
            ["春苗人力资源（北京）有限公司", "春苗人力资源（南昌）有限公司",
             "杭州唐人有限公司", "北京唐人有限公司"], {})
        self.assertEqual(mapping["北京春苗"], "春苗人力资源北京有限公司")
        self.assertEqual(ambiguous, {"唐人"})
        mapping, ambiguous = company_mapping(["唐人"], ["杭州唐人有限公司", "北京唐人有限公司"],
                                             parse_company_aliases("唐人=杭州唐人有限公司"))
        self.assertEqual(mapping["唐人"], "杭州唐人有限公司")
        self.assertFalse(ambiguous)

    def test_masked_or_numeric_ids_do_not_trigger_fill(self):
        for identity in ("脱敏编号00000000000001", 110101199001010000):
            with self.subTest(identity=identity):
                self.write_summary([summary_row(identity)])
                self.write_flow([join_row(identity)])
                result = self.run_check()
                notices, _, _ = self.read_result(result)
                self.assertEqual(result.filled_count, 0)
                self.assertTrue(any(row[0] == "身份证待确认" for row in notices))

    def test_cross_month_date_difference_is_not_misreported_as_no_workflow(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row(入职日期=date(2026, 8, 1))])
        notices, _, _ = self.read_result(self.run_check())
        self.assertTrue(any(row[0] == "日期差异" for row in notices))
        self.assertFalse(any(row[0] == "异动表有、流程无" for row in notices))

    def test_other_company_workflow_is_not_swallowed_by_date_difference(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row(入职日期=date(2026, 7, 2)),
                         join_row(入职公司="另一公司", 入职日期=date(2026, 7, 3))])
        notices, _, _ = self.read_result(self.run_check())
        self.assertTrue(any(row[0] == "日期差异" for row in notices))
        self.assertTrue(any(row[0] == "待确认" and "第3行" in row[8] for row in notices))

    def test_missing_date_at_other_company_does_not_block_unique_match(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row(), join_row(入职公司="另一公司", 入职日期=None)])
        result = self.run_check()
        notices, filled, _ = self.read_result(result)
        self.assertEqual(filled["增员"]["G3"].value, "本科")
        self.assertTrue(any(row[0] == "流程待确认" for row in notices))

    def test_header_only_flow_is_zero_records_but_missing_flow_is_unchecked(self):
        self.write_summary([summary_row()])
        self.write_flow([])
        notices, _, _ = self.read_result(self.run_check())
        self.assertTrue(any(row[0] == "异动表有、流程无" for row in notices))
        self.assertTrue(any(row[0] == "未核对" and "离职" in row[9] for row in notices))

    def test_custom_headers_use_existing_template_mapping(self):
        self.write_summary([summary_row()])
        self.write_flow([{"单位名称": "浙江唐人人力资源有限公司", "人员名字": "测试甲",
                          "证件号码": ID_A, "入职日期": date(2026, 7, 1), "最高学历文化程度": "本科"}],
                        headers=["单位名称", "人员名字", "证件号码", "入职日期", "最高学历文化程度"])
        result = self.run_check(template_rules={"fields": {
            "flow_join|公司": ["单位名称"], "flow_join|姓名": ["人员名字"], "flow_join|身份证号码": ["证件号码"],
        }})
        self.assertEqual(result.filled_count, 1)

    def test_imported_formula_like_text_stays_text(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row()])
        book = load_workbook(self.flow)
        book.active["I2"] = "=DANGEROUS()"
        book.active["I2"].data_type = "s"
        book.save(self.flow)
        book.close()
        _, filled, _ = self.read_result(self.run_check())
        self.assertEqual(filled["增员"]["J3"].value, "=DANGEROUS()")
        self.assertEqual(filled["增员"]["J3"].data_type, "s")

    def test_cancellation_writes_no_results(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row()])
        with self.assertRaisesRegex(RuntimeError, "停止"):
            self.run_check(cancelled=lambda: True)
        self.assertFalse(self.output.exists())

    def test_blank_month_uses_row_dates_and_warns_for_empty_category(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row()])
        result = self.run_check(month="")
        notices, _, report = self.read_result(result)
        self.assertEqual(dict(list(report["核对说明"].values)[1:])["入职核对月份"], "2026-07")
        self.assertEqual(result.matched_count, 1)
        self.assertTrue(any(row[0] == "未核对" and "月份" in row[9] for row in notices))

    def test_duplicate_report_contains_company_employee_and_event_date(self):
        self.write_summary([summary_row()])
        self.write_flow([join_row(工号="001"), join_row(工号="002")])
        notices, _, _ = self.read_result(self.run_check())
        duplicates = [row for row in notices if row[0] == "重复流程待确认"]
        self.assertEqual({row[13] for row in duplicates}, {"001", "002"})
        self.assertTrue(all(row[11] == "浙江唐人人力资源有限公司" for row in duplicates))
        self.assertTrue(all(row[15].date() == date(2026, 7, 1) for row in duplicates))
        self.assertTrue(all(str(self.root) not in row[8] for row in duplicates))
