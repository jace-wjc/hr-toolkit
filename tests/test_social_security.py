from __future__ import annotations

import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from hr_toolkit.tools.social_security import (
    DetailRecord,
    _read_payment_file,
    _read_payment_headers,
    _read_xls_payment_file,
    _source_context,
    _with_sheet_fee_period,
    _payment_periods_from_row,
    _write_detail_workbook,
    generate_social_security_reports,
)


class SocialSecurityTest(unittest.TestCase):
    def test_difference_headers_use_only_each_category_periods(self) -> None:
        from hr_toolkit.tools.social_security import DIFFERENCE_COLUMNS, _write_difference_headers
        workbook = Workbook()
        self.addCleanup(workbook.close)
        ws = workbook.active
        record = SimpleNamespace(difference_periods={
            "医疗": {"202603", "202605"}, "养老": {"202601"},
            "失业": {"202512", "202601"}, "补充工伤": {"202604"},
        })
        _write_difference_headers(ws, [record])
        expected = {"医疗": "医疗2026年3月-5月补差", "养老": "养老2026年1月补差",
                    "失业": "失业2025年12月-2026年1月补差", "工伤": "工伤补差",
                    "补充工伤": "补充工伤2026年4月补差"}
        for category, text in expected.items():
            self.assertEqual(ws.cell(2, DIFFERENCE_COLUMNS[category]["基数"]).value, text)
        _write_difference_headers(ws, [SimpleNamespace(difference_periods={"工伤": {"未知"}})])
        for category, columns in DIFFERENCE_COLUMNS.items():
            self.assertEqual(ws.cell(2, columns["基数"]).value, f"{category}补差")

    def test_amount_diagnostics_preserve_parsing_and_do_not_log_cell_contents(self) -> None:
        from hr_toolkit.tools import social_security as social
        context = _source_context(Path("2026年6月社保.xlsx"))
        for value in (None, "", "  ", 0, 12.5, "1,200元", "金额隐私样例"):
            with self.subTest(value=value), patch.object(social.runlog, "log_line") as log:
                self.assertEqual(social._payment_amount(value, context, 7, 4), social._to_number(value))
                if value == "金额隐私样例":
                    message = log.call_args[0][0]
                    self.assertIn("第 7 行", message)
                    self.assertIn("列 4", message)
                    self.assertNotIn(value, message)
                else:
                    log.assert_not_called()
        for first, second in ((0, 12), ("", 12), ("异常内容", 12), (None, None)):
            with patch.object(social.runlog, "log_line"):
                actual = social._row_payment_amount({"本期应缴费额": first, "应缴费额(元)": second},
                                                   ("本期应缴费额", "应缴费额(元)"), context, 7)
            self.assertEqual(actual, social._to_number(first or second))

    def test_supplementary_injury_difference_is_separate_from_normal_and_arrears(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年6月社保单位缴费明细.xlsx"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(source, [
                ["张三", "360111199001010011", "补充工伤保险费", "补充工伤保险", date(2026, month, 1), base, 0.0012, amount, nature]
                for month, base, amount, nature in (
                    (3, 1000, 1.2, "补缴"), (4, 1000, 1.2, "补缴"),
                    (5, 100, 0.12, "补差"), (6, 1000, 1.2, "正常缴费"),
                )
            ])
            result = generate_social_security_reports(source, roster, root / "output")
            self.assertNotIn("模板没有对应补差明细列", "\n".join(result.warnings))
            for path in [result.detail_output_file, *result.detail_output_files]:
                with self.subTest(file=path.name):
                    wb = load_workbook(path)
                    try:
                        ws = wb["社保明细表"]
                        self.assertEqual(ws.max_column, 79)
                        self.assertEqual(ws["BE2"].value, "补充工伤2026年5月补差")
                        self.assertEqual([ws.cell(3, c).value for c in (57, 58, 59)], ["基数", "单位", "单位  金额"])
                        self.assertEqual(ws["H4"].value, "202603-202604")
                        self.assertEqual(ws["AA4"].value, 1000)
                        self.assertEqual(ws["AC4"].value, 2.4)
                        self.assertIsNone(ws["BG4"].value)
                        self.assertIsNone(ws["BN4"].value)
                        self.assertEqual(ws["H5"].value, "202606")
                        self.assertEqual(ws["AC5"].value, "=ROUND(AA5*AB5,2)")
                        self.assertEqual(ws["BE5"].value, 100)
                        self.assertEqual(ws["BF5"].value, 0.0012)
                        self.assertEqual(ws["BF5"].number_format, "0.00%")
                        self.assertEqual(ws["BG5"].value, 0.12)
                        self.assertEqual(ws["BN5"].value, "=AV5+BA5+BD5+BL5+BG5")
                        self.assertEqual(ws["BS5"].value, "=M5+R5+W5+Z5+AK5+BN5+AC5+BQ5")
                        self.assertIsNone(ws["BX4"].value)
                        self.assertEqual(ws["BX5"].value, 20)
                    finally:
                        wb.close()

    def test_supplementary_injury_multiple_bases_and_negative_difference(self) -> None:
        for second_base, second_amount, expected_base, expected_amount in (
            (100, 0.12, 100, 0.24), (200, 0.24, None, 0.36), (100, -0.24, 100, -0.12),
        ):
            with self.subTest(base=second_base, amount=second_amount), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = root / "北京春苗抚州账户2026年6月社保单位缴费明细.xlsx"
                roster = root / "参保人员花名册.xlsx"
                _write_roster(roster)
                _write_long_payment_rows_with_nature(source, [
                    ["张三", "360111199001010011", "补充工伤保险费", "补充工伤保险", date(2026, 3, 1), 100, 0.0012, 0.12, "补差"],
                    ["张三", "360111199001010011", "补充工伤保险费", "补充工伤保险", date(2026, 4, 1), second_base, 0.0012, second_amount, "补差"],
                ])
                result = generate_social_security_reports(source, roster, root / "output")
                wb = load_workbook(result.detail_output_file)
                try:
                    ws = wb["社保明细表"]
                    self.assertEqual(ws["BE2"].value, "补充工伤2026年3月-4月补差")
                    self.assertEqual(ws["BE4"].value, expected_base)
                    self.assertAlmostEqual(ws["BG4"].value, expected_amount)
                    self.assertIsNone(ws["AC4"].value)
                    self.assertEqual(ws["BN4"].value, "=AV4+BA4+BD4+BL4+BG4")
                    self.assertEqual(any("补充工伤补差含多种基数" in warning for warning in result.warnings), expected_base is None)
                finally:
                    wb.close()

    def test_two_level_payment_headers_read_both_sides(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for difference in (False, True):
                for personal_first in (False, True):
                    with self.subTest(difference=difference, personal_first=personal_first):
                        source = root / f"测试2024年10月{'补差' if difference else '正常'}明细.xlsx"
                        _write_two_level_payment_file(source, personal_first=personal_first)
                        lines = _read_payment_file(source)
                        self.assertEqual(len(lines), 9)
                        amounts = {(line.source_row, line.category, line.side): line.amount for line in lines}
                        self.assertEqual(amounts, {
                            (5, "养老", "个人"): 384.96, (5, "养老", "单位"): 769.92,
                            (5, "失业", "个人"): 24.06, (5, "失业", "单位"): 24.06,
                            (5, "医疗", "个人"): 96.24, (5, "医疗", "单位"): 384.96,
                            (5, "工伤", "单位"): 12.03,
                            (6, "养老", "个人"): -10, (6, "养老", "单位"): -20,
                        })
                        self.assertEqual({line.fee_period for line in lines}, {"202409"})
                        self.assertEqual({line.nature_hint for line in lines}, {"difference" if difference else None})
                        self.assertTrue(all(line.base is None and line.rate is None for line in lines))
                        self.assertAlmostEqual(sum(line.amount for line in lines if line.source_row == 5 and line.side == "个人"), 505.26)

    def test_two_level_payment_xls_path_matches_xlsx(self) -> None:
        # 同一单元格网格送入 xlrd 适配入口，核对零基列号与一基列号一致。
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "测试2024年10月补差明细.xlsx"
            _write_two_level_payment_file(source, personal_first=True)
            expected = _read_payment_file(source)
            wb = load_workbook(source, data_only=True)
            try:
                sheet_name = wb.active.title
                values = list(wb.active.iter_rows(values_only=True))
            finally:
                wb.close()
            sheet = SimpleNamespace(
                name=sheet_name,
                nrows=len(values), row_values=lambda row: list(values[row]),
                cell_value=lambda row, col: values[row][col],
            )
            book = SimpleNamespace(sheets=lambda: [sheet])
            with patch("xlrd.open_workbook", return_value=book):
                actual = _read_xls_payment_file(source.with_suffix(".xls"), _source_context(source))
            self.assertEqual(actual, expected)

    def test_payment_subheaders_do_not_change_single_headers_or_cross_groups(self) -> None:
        for first_column in (0, 1):
            with self.subTest(first_column=first_column):
                single = ["姓名", "证件号码", "基本养老保险(个人缴纳)应缴费额", "基本养老保险(单位缴纳)应缴费额"]
                expected = {name: index for index, name in enumerate(single, start=first_column)}
                self.assertEqual(_read_payment_headers(single, ["张三", "360111199001010011", 8, 16], first_column=first_column), expected)
                self.assertEqual(_read_payment_headers(single, [], first_column=first_column), expected)
                headers = _read_payment_headers(
                    ["姓名", "证件号码", "养老应缴费额", None, "合计", None, "未知应缴费额", None],
                    [None, None, "单位部分", "个人部分", "单位部分", "个人部分", "单位部分", "个人部分"],
                    first_column=first_column,
                )
                self.assertEqual(headers["养老应缴费额(单位部分)"], first_column + 2)
                self.assertEqual(headers["养老应缴费额(个人部分)"], first_column + 3)
                self.assertNotIn("养老应缴费额", headers)
                self.assertEqual(headers["合计"], first_column + 4)
                self.assertEqual(headers["未知应缴费额"], first_column + 6)
                self.assertEqual(len(headers), 6)

    def test_medical_difference_rollup_includes_medical_but_not_arrears(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年6月社保单位缴费明细.xlsx"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            rows = [
                ["张三", "360111199001010011", kind, item, date(2026, 5, 1), base, rate, amount, nature]
                for kind, item, base, rate, amount, nature in (
                    ("养老保险费", "养老保险个人", 1085, 0.08, 86.8, "补差"),
                    ("失业保险费", "失业保险个人", 1084, 0.005, 5.42, "补差"),
                    ("医疗保险费", "医疗保险个人", 253, 0.02, 5.06, "补差"),
                    ("医疗保险费", "医疗保险个人", 5000, 0.02, 100, "补缴"),
                )
            ]
            _write_long_payment_rows_with_nature(source, rows)
            result = generate_social_security_reports(source, roster, root / "output")
            wb = load_workbook(result.detail_output_file)
            try:
                ws = wb["社保明细表"]
                self.assertEqual(ws["H4"].value, "202605")
                self.assertEqual(ws["P4"].value, "=ROUND(N4*O4,2)")
                self.assertIsNone(ws["BM4"].value)
                self.assertEqual(ws["H5"].value, "202606")
                self.assertEqual(ws["BM5"].value, "=AT5+AY5+BJ5")
                self.assertAlmostEqual(sum(ws[cell].value for cell in ("AT5", "AY5", "BJ5")), 97.28)
                self.assertEqual(ws["BM2"].value, "个人社保\n补缴合计")
                self.assertEqual(ws["BN2"].value, "单位社保\n补缴合计")
                self.assertIsNone(ws["BX4"].value)
                self.assertEqual(ws["BX5"].value, 20)
            finally:
                wb.close()

    def test_arrears_merge_actual_periods_and_split_different_profiles(self) -> None:
        for second_base, second_rate, expected in (
            (500, 0.01, [("202603-202604", 500, 0.01, 10), ("202605", 500, 0.01, 5)]),
            (400, 0.01, [("202603", 500, 0.01, 5), ("202604", 400, 0.01, 4), ("202605", 500, 0.01, 5)]),
            (500, 0.02, [("202603", 500, 0.01, 5), ("202604", 500, 0.02, 10), ("202605", 500, 0.01, 5)]),
        ):
            with self.subTest(base=second_base, rate=second_rate), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = root / "北京春苗抚州账户2026年5月社保单位缴费明细.xlsx"
                roster = root / "参保人员花名册.xlsx"
                _write_roster(roster)
                _write_long_payment_rows_with_nature(source, [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 3, 1), 500, 0.01, 5, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 4, 1), second_base, second_rate, second_base * second_rate, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 5, 1), 500, 0.01, 5, "正常缴费"],
                ])
                result = generate_social_security_reports(source, roster, root / "output")
                self.assertEqual(result.detail_record_count, len(expected))
                wb = load_workbook(result.detail_output_file)
                try:
                    ws = wb["社保明细表"]
                    for row, (period, base, rate, amount) in enumerate(expected, start=4):
                        self.assertEqual(ws.cell(row, 8).value, period)
                        self.assertEqual(ws.cell(row, 24).value, base)
                        self.assertEqual(ws.cell(row, 25).value, rate)
                        self.assertEqual(ws.cell(row, 26).value, amount if "-" in period else f"=ROUND(X{row}*Y{row},2)")
                        self.assertIsNone(ws.cell(row, 66).value)
                    self.assertEqual(sum(ws.cell(row, 76).value or 0 for row in range(4, 4+len(expected))), 20)
                finally:
                    wb.close()

    def test_sheet_period_hint_keeps_bill_month_and_row_period_precedence(self) -> None:
        context = _source_context(Path("测试2024年10月补差明细.xlsx"))
        for rows, expected in (
            ([["费款所属期：", None, "2024-09"]], ("202409", "202409")),
            ([["费款所属期：2023-12至2024-02"]], ("202312", "202402")),
            ([["打印日期：", "2024-09"]], ("202410", "202410")),
            ([["费款所属期：", "2024-20"]], ("202410", "202410")),
            ([["费款所属期：", "2024-13"]], ("202410", "202410")),
        ):
            with self.subTest(rows=rows):
                amended = _with_sheet_fee_period(context, rows)
                self.assertEqual(amended.billing_period_hint, "202410")
                self.assertEqual(_payment_periods_from_row({}, amended), expected)
                self.assertEqual(_payment_periods_from_row({"费款所属期起": "2024-08"}, amended), ("202408", "202408"))

    def test_normal_payments_with_later_dates_still_keep_bill_month(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年5月社保单位缴费明细.xlsx"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(source, [
                ["张三", "360111199001010011", "工伤保险费", "工伤保险", day, 500, 0.01, 5, "正常缴费"]
                for day in (date(2026, 5, 1), date(2026, 5, 2), date(2026, 6, 1), date(2026, 6, 2), date(2026, 6, 3))
            ])
            result = generate_social_security_reports(source, roster, root / "output")
            self.assertEqual(result.detail_record_count, 1)
            self.assertEqual(result.period_counts, {"202605": 1})
            wb = load_workbook(result.detail_output_file)
            try:
                self.assertEqual(wb["社保明细表"]["H4"].value, "202605")
                self.assertEqual(wb["社保明细表"]["X4"].value, 500)
                self.assertEqual(wb["社保明细表"]["Z4"].value, 25)
            finally:
                wb.close()

    def test_large_detail_output_does_not_rescan_sheet_width_per_record(self) -> None:
        records = [
            DetailRecord(
                id_card=f"3601111990{index:08d}"[:18],
                name=f"员工{index}",
                period="202601",
                billing_period="202601",
                period_split_input=False,
                account="测试账户",
                account_display="测试账户",
                company="测试公司",
                insured_place="测试地",
                project="测试项目",
                project_display="测试项目",
                cost_center="测试成本中心",
                start_period="202601",
                management_fee=0,
            )
            for index in range(100)
        ]
        original_getter = Worksheet.max_column.fget
        access_count = 0

        def counted_max_column(worksheet):
            nonlocal access_count
            access_count += 1
            return original_getter(worksheet)

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch.object(Worksheet, "max_column", new=property(counted_max_column)):
                _write_detail_workbook(records, root / "明细.xlsx", root)

        self.assertLessEqual(access_count, 3)

    def test_generate_social_security_reports_from_mixed_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_file(input_dir / "北京春苗抚州账户2026年5月社保单位缴费明细.xlsx")
            _write_single_kind_file(input_dir / "2026-04——工伤保险（单位缴纳部分）职工明细.xlsx")
            progress = []

            result = generate_social_security_reports(
                input_dir,
                roster,
                output_dir,
                progress_callback=lambda current, total, message: progress.append(
                    (current, total, message)
                ),
            )
            payload = result.to_dict()

            self.assertEqual(payload["source_file_count"], 2)
            self.assertEqual(payload["source_record_count"], 3)
            self.assertEqual(payload["detail_record_count"], 2)
            self.assertEqual(payload["employee_count"], 2)
            self.assertEqual(payload["account_counts"], {"北京抚州": 1, "唐人四川": 1})
            self.assertEqual(payload["period_counts"], {"202605": 1, "202604": 1})
            self.assertTrue(result.detail_output_file and result.detail_output_file.exists())
            self.assertEqual(len(result.detail_output_files), 2)
            split_names = {path.name for path in result.detail_output_files}
            self.assertEqual(split_names, {"北京抚州-社保明细表.xlsx", "唐人四川-社保明细表.xlsx"})
            self.assertTrue(result.summary_output_file and result.summary_output_file.exists())
            self.assertEqual(progress[0][:2], (0, 5))
            self.assertEqual(progress[-1], (5, 5, "社保报表生成完成"))

            detail_wb = load_workbook(result.detail_output_file, data_only=False)
            detail_ws = detail_wb["社保明细表"]
            rows = {detail_ws.cell(row, 6).value: row for row in range(4, 6)}
            zhang_row = rows["360111199001010011"]
            li_row = rows["360111199002020022"]
            self.assertEqual(detail_ws.cell(zhang_row, 2).value, "北京春苗")
            self.assertEqual(detail_ws.cell(zhang_row, 3).value, "抚州")
            self.assertEqual(detail_ws.cell(zhang_row, 8).value, "202605")
            self.assertEqual(detail_ws.cell(zhang_row, 11).value, f"=ROUND(I{zhang_row}*J{zhang_row},2)")
            self.assertEqual(detail_ws.cell(zhang_row, 13).value, f"=ROUND(I{zhang_row}*L{zhang_row},2)")
            self.assertEqual(detail_ws.cell(zhang_row, 10).number_format, "0.00%")
            self.assertEqual(detail_ws.cell(zhang_row, 70).number_format, "0.00_ ")
            self.assertEqual(detail_ws.cell(zhang_row, 76).value, 20)
            self.assertEqual(detail_ws.cell(li_row, 2).value, "唐人数智")
            self.assertEqual(detail_ws.cell(li_row, 8).value, "202604")
            self.assertEqual(detail_ws.cell(li_row, 26).value, f"=ROUND(X{li_row}*Y{li_row},2)")
            template_wb = load_workbook(
                Path(__file__).resolve().parents[1] / "hr_toolkit" / "templates" / "social_security_detail_template.xlsx",
                data_only=False,
            )
            template_ws = template_wb["社保明细表模板"]
            for col_index in range(1, 77):
                output_col = col_index + 3 if col_index >= 57 else col_index
                self.assertEqual(detail_ws.cell(zhang_row, output_col)._style, template_ws.cell(4, col_index)._style)
            template_wb.close()
            detail_wb.close()

            split_detail = next(path for path in result.detail_output_files if path.name == "北京抚州-社保明细表.xlsx")
            split_wb = load_workbook(split_detail, data_only=False)
            split_ws = split_wb["社保明细表"]
            self.assertEqual(split_ws.cell(1, 1).value, "北京春苗2026年5月社保明细表")
            self.assertEqual(split_ws.max_row, 4)
            self.assertEqual(split_ws.cell(4, 5).value, "张三")
            split_wb.close()

            summary_wb = load_workbook(result.summary_output_file, data_only=False)
            self.assertIn("社保汇总表", summary_wb.sheetnames)
            self.assertIn("北京春苗", summary_wb.sheetnames)
            self.assertIn("唐人数智", summary_wb.sheetnames)
            self.assertIn("数据分析", summary_wb.sheetnames)
            self.assertIn("异常提醒", summary_wb.sheetnames)
            summary_wb.close()

    def test_generate_social_security_reports_from_zip(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            roster = root / "参保人员花名册.xlsx"
            source = root / "北京春苗抚州账户2026年5月社保单位缴费明细.xlsx"
            archive = root / "社保清单.zip"
            output_dir = root / "output"
            _write_roster(roster)
            _write_long_payment_file(source)
            with zipfile.ZipFile(archive, "w") as zip_file:
                zip_file.write(source, arcname=source.name)

            result = generate_social_security_reports([archive], roster, output_dir)

            self.assertEqual(result.source_record_count, 2)
            self.assertEqual(result.detail_record_count, 1)
            self.assertTrue(result.detail_output_file and result.detail_output_file.exists())
            self.assertTrue(result.summary_output_file and result.summary_output_file.exists())

    def test_zip_name_supplies_context_for_root_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            roster = root / "参保人员花名册.xlsx"
            source = root / "2026-04——工伤保险（单位缴纳部分）职工明细.xlsx"
            archive = root / "北京春苗抚州账户2026年5月社保单位缴费明细.zip"
            output_dir = root / "output"
            _write_roster(roster)
            _write_single_kind_file(source)
            with zipfile.ZipFile(archive, "w") as zip_file:
                zip_file.write(source, arcname=source.name)

            result = generate_social_security_reports([archive], roster, output_dir)

            self.assertEqual(result.period_counts, {"202604": 1})
            joined_warnings = "\n".join(result.warnings)
            self.assertIn("参保账户与花名册不一致", joined_warnings)
            self.assertIn("参保地与花名册不一致", joined_warnings)
            wb = load_workbook(result.detail_output_file, data_only=True)
            ws = wb["社保明细表"]
            self.assertEqual(ws.cell(4, 2).value, "北京春苗")
            self.assertEqual(ws.cell(4, 3).value, "抚州")
            self.assertEqual(ws.cell(4, 8).value, "202604")
            wb.close()

    def test_file_fee_month_overrides_container_month(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir = root / "唐人四川2026年5月社保单位缴费明细"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            input_dir.mkdir()
            _write_roster(roster)
            _write_single_kind_file(input_dir / "2026-04——工伤保险（单位缴纳部分）职工明细.xlsx")

            result = generate_social_security_reports(input_dir, roster, output_dir)

            self.assertEqual(result.period_counts, {"202604": 1})
            wb = load_workbook(result.detail_output_file, data_only=True)
            ws = wb["社保明细表"]
            self.assertEqual(ws.cell(4, 8).value, "202604")
            wb.close()

    def test_distinguishes_arrears_and_difference_in_same_month_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir = root / "唐人四川2026年5月社保单位缴费明细"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            input_dir.mkdir()
            _write_roster(
                roster,
                extra_rows=[
                    ["王五", "360111199003030033", "正常", date(2026, 2, 1), "唐人四川", "唐人数智科技股份有限公司", "四川项目部", "项目（成都）", "成本二", 30],
                ],
            )
            zhang = ("张三", "360111199001010011")
            li = ("李四", "360111199002020022")
            wang = ("王五", "360111199003030033")
            _write_single_kind_rows(
                input_dir / "2026-01——工伤保险（单位缴纳部分）职工明细.xlsx",
                [(*li, 4588, 0.001, 4.58), (*wang, 4588, 0.001, 4.58)],
            )
            _write_single_kind_rows(
                input_dir / "2026-02——工伤保险（单位缴纳部分）职工明细.xlsx",
                [(*li, 4588, 0.001, 4.58), (*wang, 4588, 0.001, 4.58)],
            )
            _write_single_kind_rows(
                input_dir / "2026-03——工伤保险（单位缴纳部分）职工明细.xlsx",
                [(*li, 4588, 0.001, 4.58), (*wang, 4588, 0.001, 4.58), (*zhang, 4588, 0.003, 13.76)],
            )
            _write_single_kind_rows(
                input_dir / "2026-04——工伤保险（单位缴纳部分）职工明细.xlsx",
                [(*li, 4588, 0.003, 13.76), (*wang, 4588, 0.003, 13.76), (*zhang, 4588, 0.003, 13.76)],
            )

            result = generate_social_security_reports(input_dir, roster, output_dir)

            self.assertEqual(len(result.source_files), 4)
            self.assertEqual(result.source_record_count, 10)
            self.assertEqual(result.detail_record_count, 4)
            self.assertEqual(result.period_counts, {"202604": 3, "202603": 1})
            self.assertNotIn("待确认历史缴费", "\n".join(result.warnings))
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            rows = {ws.cell(row, 6).value: row for row in range(4, 8) if ws.cell(row, 8).value == "202604"}
            li_row = rows[li[1]]
            zhang_row = rows[zhang[1]]
            self.assertEqual(ws["A1"].value, "唐人数智2026年4月社保明细表")
            self.assertEqual(ws["BB2"].value, "工伤2026年1月-3月补差")
            self.assertEqual(ws["BM2"].value, "个人社保\n补缴合计")
            self.assertEqual(ws["BN2"].value, "单位社保\n补缴合计")
            self.assertEqual(ws.cell(li_row, 8).value, "202604")
            self.assertEqual(ws.cell(li_row, 26).value, f"=ROUND(X{li_row}*Y{li_row},2)")
            self.assertEqual(ws.cell(li_row, 54).value, 4588)
            self.assertEqual(ws.cell(li_row, 55).value, 0.001)
            self.assertEqual(ws.cell(li_row, 55).number_format, "0.00%")
            self.assertEqual(ws.cell(li_row, 56).value, 13.74)
            self.assertIsNone(ws.cell(li_row, 65).value)
            self.assertEqual(ws.cell(li_row, 66).value, f"=AV{li_row}+BA{li_row}+BD{li_row}+BL{li_row}")
            self.assertEqual(
                ws.cell(li_row, 71).value,
                f"=M{li_row}+R{li_row}+W{li_row}+Z{li_row}+AK{li_row}+BN{li_row}+AC{li_row}+BQ{li_row}",
            )
            self.assertEqual(ws.cell(li_row, 70).value, f"=K{li_row}+P{li_row}+U{li_row}+AI{li_row}+BM{li_row}")
            self.assertEqual(ws.cell(li_row, 72).value, f"=BR{li_row}+BS{li_row}")
            self.assertEqual(ws.cell(li_row, 77).value, f"=ROUND((BT{li_row}+BX{li_row})*6.72%,2)")
            self.assertEqual(ws.cell(li_row, 78).value, f"=BT{li_row}+BX{li_row}+BY{li_row}")
            self.assertIsNone(ws.cell(li_row, 79).value)
            self.assertEqual(ws.cell(zhang_row, 8).value, "202604")
            self.assertEqual(ws.cell(zhang_row, 24).value, 4588)
            self.assertEqual(ws.cell(zhang_row, 25).value, 0.003)
            self.assertEqual(ws.cell(zhang_row, 26).value, f"=ROUND(X{zhang_row}*Y{zhang_row},2)")
            arrears_row = next(row for row in range(4, 8) if ws.cell(row, 8).value == "202603")
            self.assertEqual(ws.cell(arrears_row, 6).value, zhang[1])
            self.assertEqual(ws.cell(arrears_row, 26).value, f"=ROUND(X{arrears_row}*Y{arrears_row},2)")
            self.assertIsNone(ws.cell(arrears_row, 76).value)
            self.assertIsNone(ws.cell(zhang_row, 56).value)
            self.assertIsNone(ws.cell(zhang_row, 66).value)
            self.assertIsNone(ws.cell(zhang_row, 79).value)
            wb.close()

    def test_compatible_arrears_and_normal_still_use_separate_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年4月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 3, 1), 5000, 0.01, 50, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 4, 1), 5000, 0.01, 50, "正常缴费"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertEqual(result.detail_record_count, 2)
            self.assertEqual(result.period_counts, {"202603": 1, "202604": 1})
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["H4"].value, "202603")
            self.assertEqual(ws["H5"].value, "202604")
            self.assertEqual(ws["X4"].value, 5000)
            self.assertEqual(ws["Y4"].value, 0.01)
            self.assertEqual(ws["Z4"].value, "=ROUND(X4*Y4,2)")
            self.assertEqual(ws["Z5"].value, "=ROUND(X5*Y5,2)")
            self.assertIsNone(ws["BD4"].value)
            self.assertIsNone(ws["BN4"].value)
            wb.close()

            summary_wb = load_workbook(result.summary_output_file, data_only=False)
            analysis = summary_wb["数据分析"]
            category_rows = {
                analysis.cell(row, 1).value: row
                for row in range(1, analysis.max_row + 1)
            }
            self.assertEqual(analysis.cell(category_rows["工伤"], 3).value, 5000)
            summary_wb.close()

    def test_same_base_is_not_multiplied_and_arrears_keep_actual_months(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "唐人长春2026年5月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 5, 1), 500, 0.01, 5, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 5, 2), 500, 0.01, 5, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 6, 1), 500, 0.01, 5, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 6, 2), 500, 0.01, 5, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 6, 3), 500, 0.01, 5, "补缴"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertEqual(result.detail_record_count, 1)
            self.assertEqual(result.period_counts, {"202605-202606": 1})
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["H4"].value, "202605-202606")
            self.assertEqual(ws["X4"].value, 500)
            self.assertEqual(ws["Y4"].value, 0.01)
            self.assertEqual(ws["Z4"].value, 25)
            self.assertEqual(ws["BX4"].value, 20)
            wb.close()

    def test_different_bases_split_and_arrears_keep_actual_months(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "2026-05——工伤保险（单位缴纳部分）职工明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 5, 1), 400, 0.01, 4, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 6, 1), 500, 0.01, 5, "补缴"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertEqual(result.detail_record_count, 2)
            self.assertEqual(result.period_counts, {"202605": 1, "202606": 1})
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            rows = {ws.cell(row, 24).value: row for row in range(4, 6)}
            self.assertEqual(set(rows), {400, 500})
            self.assertEqual(ws.cell(rows[400], 8).value, "202605")
            self.assertEqual(ws.cell(rows[500], 8).value, "202606")
            self.assertEqual(ws.cell(rows[400], 26).value, f"=ROUND(X{rows[400]}*Y{rows[400]},2)")
            self.assertEqual(ws.cell(rows[500], 26).value, f"=ROUND(X{rows[500]}*Y{rows[500]},2)")
            self.assertEqual(ws.cell(rows[400], 76).value, 20)
            self.assertIsNone(ws.cell(rows[500], 76).value)
            wb.close()

    def test_explicit_difference_marker_handles_single_history_month(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年4月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 3, 1), 4588, 0.001, 4.58, "调整补收"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 4, 1), 4588, 0.003, 13.76, "正常缴费"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["A1"].value, "北京春苗2026年4月社保明细表")
            self.assertEqual(ws["H4"].value, "202604")
            self.assertEqual(ws["Z4"].value, "=ROUND(X4*Y4,2)")
            self.assertEqual(ws["BB4"].value, 4588)
            self.assertEqual(ws["BC4"].value, 0.001)
            self.assertEqual(ws["BC4"].number_format, "0.00%")
            self.assertEqual(ws["BD4"].value, 4.58)
            self.assertEqual(ws["BN4"].value, "=AV4+BA4+BD4+BL4")
            self.assertIsNone(ws["CA4"].value)
            wb.close()

    def test_explicit_arrears_marker_never_moves_amount_to_difference_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年4月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 3, 1), 4588, 0.001, 4.58, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 4, 1), 4588, 0.003, 13.76, "正常缴费"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertEqual(result.warnings, [])
            self.assertEqual(result.detail_record_count, 2)
            self.assertEqual(result.period_counts, {"202603": 1, "202604": 1})
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            rows = {ws.cell(row, 25).value: row for row in range(4, 6)}
            arrears_row = rows[0.001]
            current_row = rows[0.003]
            self.assertEqual(ws.cell(arrears_row, 8).value, "202603")
            self.assertEqual(ws.cell(current_row, 8).value, "202604")
            self.assertEqual(ws.cell(arrears_row, 24).value, 4588)
            self.assertEqual(ws.cell(arrears_row, 25).value, 0.001)
            self.assertEqual(ws.cell(arrears_row, 26).value, 4.58)
            self.assertIsNone(ws.cell(arrears_row, 56).value)
            self.assertIsNone(ws.cell(arrears_row, 66).value)
            self.assertIsNone(ws.cell(arrears_row, 76).value)
            self.assertEqual(ws.cell(current_row, 26).value, f"=ROUND(X{current_row}*Y{current_row},2)")
            self.assertIsNone(ws.cell(current_row, 56).value)
            self.assertIsNone(ws.cell(current_row, 66).value)
            self.assertEqual(ws.cell(current_row, 76).value, 20)
            self.assertIsNone(ws.cell(current_row, 79).value)
            wb.close()

    def test_explicit_arrears_and_difference_remain_separate_for_same_person(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年4月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 2, 1), 4588, 0.003, 13.76, "补缴"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 3, 1), 4588, 0.001, 4.58, "调整补收"],
                    ["张三", "360111199001010011", "工伤保险费", "工伤保险", date(2026, 4, 1), 4588, 0.003, 13.76, "正常缴费"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertEqual(result.warnings, [])
            self.assertEqual(result.detail_record_count, 2)
            self.assertEqual(result.period_counts, {"202602": 1, "202604": 1})
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["H4"].value, "202602")
            self.assertEqual(ws["H5"].value, "202604")
            self.assertEqual(ws["X4"].value, 4588)
            self.assertEqual(ws["Y4"].value, 0.003)
            self.assertEqual(ws["Z4"].value, "=ROUND(X4*Y4,2)")
            self.assertEqual(ws["Z5"].value, "=ROUND(X5*Y5,2)")
            self.assertIsNone(ws["BD4"].value)
            self.assertIsNone(ws["BN4"].value)
            self.assertEqual(ws["BD5"].value, 4.58)
            self.assertEqual(
                ws["BN5"].value,
                "=AV5+BA5+BD5+BL5",
            )
            self.assertIsNone(ws["CA4"].value)
            wb.close()

    def test_unsupported_difference_category_is_kept_visible_in_normal_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "北京春苗抚州账户2026年4月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_long_payment_rows_with_nature(
                source,
                [
                    ["张三", "360111199001010011", "大病医疗保险费", "大病医疗保险", date(2026, 3, 1), 4588, 0.001, 5, "补差"],
                    ["张三", "360111199001010011", "大病医疗保险费", "大病医疗保险", date(2026, 4, 1), 4588, 0.005, 25, "正常缴费"],
                ],
            )

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertIn("模板没有对应补差明细列", "\n".join(result.warnings))
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["AK4"].value, 30)
            self.assertIsNone(ws["BL4"].value)
            self.assertIsNone(ws["BN4"].value)
            self.assertIsNone(ws["CA4"].value)
            wb.close()

    def test_does_not_guess_single_person_historical_rate_change_as_difference(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir = root / "唐人四川2026年4月社保单位缴费明细"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            input_dir.mkdir()
            _write_roster(roster)
            person = ("李四", "360111199002020022")
            for month in (1, 2):
                _write_single_kind_rows(
                    input_dir / f"2026-{month:02d}——工伤保险（单位缴纳部分）职工明细.xlsx",
                    [(*person, 4588, 0.001, 4.58)],
                )
            _write_single_kind_rows(
                input_dir / "2026-04——工伤保险（单位缴纳部分）职工明细.xlsx",
                [(*person, 4588, 0.003, 13.76)],
            )

            result = generate_social_security_reports(input_dir, roster, output_dir)

            self.assertIn("待确认历史缴费", "\n".join(result.warnings))
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["Z4"].value, 22.92)
            self.assertIsNone(ws["BD4"].value)
            self.assertIsNone(ws["BN4"].value)
            self.assertIsNone(ws["CA4"].value)
            wb.close()

    def test_combined_wide_file_keeps_its_bill_month(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "唐人长春2026年5月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_wide_payment_file(source)

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertEqual(result.source_record_count, 10)
            self.assertEqual(result.detail_record_count, 1)
            self.assertEqual(result.period_counts, {"202605": 1})
            self.assertNotIn("待确认历史缴费", "\n".join(result.warnings))
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["A1"].value, "唐人数智2026年5月社保明细表")
            self.assertEqual(ws["H4"].value, "202605")
            self.assertEqual(ws["M4"].value, 100)
            self.assertEqual(ws["K4"].value, 50)
            self.assertEqual(ws["R4"].value, 120)
            self.assertEqual(ws["P4"].value, 20)
            self.assertEqual(ws["W4"].value, 10)
            self.assertEqual(ws["U4"].value, 5)
            self.assertEqual(ws["Z4"].value, 8)
            self.assertEqual(ws["AK4"].value, 5)
            self.assertEqual(ws["AI4"].value, 2)
            self.assertIsNone(ws["CA4"].value)
            wb.close()

    def test_wide_history_without_nature_or_basis_is_not_guessed_as_arrears(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "唐人四川2026年5月社保单位缴费明细.xlsx"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            _write_roster(roster)
            _write_wide_amount_only_history_file(source)

            result = generate_social_security_reports(source, roster, output_dir)

            self.assertIn("待确认历史缴费", "\n".join(result.warnings))
            wb = load_workbook(result.detail_output_file, data_only=False)
            ws = wb["社保明细表"]
            self.assertEqual(ws["M4"].value, 200)
            self.assertIsNone(ws["BN4"].value)
            self.assertIsNone(ws["CA4"].value)
            wb.close()

    def test_warns_when_bill_account_differs_from_roster(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir = root / "北京春苗抚州账户2026年5月社保单位缴费明细"
            output_dir = root / "output"
            roster = root / "参保人员花名册.xlsx"
            input_dir.mkdir()
            _write_roster(roster)
            _write_single_kind_file(input_dir / "2026-05——工伤保险（单位缴纳部分）职工明细.xlsx")

            result = generate_social_security_reports(input_dir, roster, output_dir)

            joined_warnings = "\n".join(result.warnings)
            self.assertIn("参保账户与花名册不一致", joined_warnings)
            self.assertIn("参保地与花名册不一致", joined_warnings)
            wb = load_workbook(result.detail_output_file, data_only=True)
            ws = wb["社保明细表"]
            self.assertEqual(ws.cell(4, 2).value, "北京春苗")
            self.assertEqual(ws.cell(4, 3).value, "抚州")
            wb.close()


def _write_two_level_payment_file(path: Path, *, personal_first: bool = False) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.append(["职工全险种明细"])
    ws.append(["费款所属期", "2024-09"])
    ws.append(["序号", "姓名", "证件类型", "证件号码", "基本养老应缴费额", None, "失业应缴费额", None, "工伤应缴费额", "基本医疗应缴费额", None, "合计"])
    sides = ["个 人\n部分", "单位部分"] if personal_first else ["单位部分", "个 人\n部分"]
    ws.append([None, None, None, None, *sides, *sides, "单位部分", *sides, None])
    for start, end in ((5, 6), (7, 8), (10, 11)):
        ws.merge_cells(start_row=3, start_column=start, end_row=3, end_column=end)
    for col in (1, 2, 3, 4, 12):
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
    pair = lambda unit, person: [person, unit] if personal_first else [unit, person]
    ws.append([1, "张三", "身份证", "360111199001010011", *pair(769.92, "384.96"), *pair(24.06, 24.06), 12.03, *pair(384.96, 96.24), 9999])
    ws.append([2, "李四", "身份证", "360111199002020022", *pair(-20, -10), *pair(0, None), 0, *pair(None, 0), -30])
    ws.append([None, "合计", None, None, 9999])
    workbook.save(path)
    workbook.close()


def _write_roster(path: Path, extra_rows: list[list[object]] | None = None) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "花名册"
    headers = [
        "*姓名.简体中文",
        "*身份证",
        "*参保状态",
        "*参保日期",
        "*参保方案.名称",
        "*参保单位.名称",
        "*责任部门.名称",
        "项目.项目名称",
        "成本中心.名称",
        "管理费",
    ]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(1, col_index).value = header
    rows = [
        ["张三", "360111199001010011", "正常", date(2026, 1, 1), "北京春苗抚州", "春苗人力资源（北京）有限公司", "抚州项目部", "项目（上饶市）", "成本一", 20],
        ["李四", "360111199002020022", "正常", date(2026, 2, 1), "唐人四川", "唐人数智科技股份有限公司", "四川项目部", "项目（成都）", "成本二", 30],
    ]
    rows.extend(extra_rows or [])
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index).value = value
    workbook.save(path)
    workbook.close()


def _write_long_payment_file(path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "缴费明细"
    headers = ["姓名", "身份证件号码", "参保费种", "征收品目", "费款所属日期起", "缴费基数", "费率", "本期应缴费额"]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(1, col_index).value = header
    rows = [
        ["张三", "360111199001010011", "城镇企业职工基本养老保险", "个人缴纳部分", date(2026, 5, 1), 3000, 0.08, 240],
        ["张三", "360111199001010011", "城镇企业职工基本养老保险", "单位缴纳部分", date(2026, 5, 1), 3000, 0.16, 480],
    ]
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index).value = value
    workbook.save(path)
    workbook.close()


def _write_single_kind_file(path: Path, name: str = "李四", id_card: str = "360111199002020022") -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "职工明细"
    headers = ["姓名", "证件号码", "缴费基数", "费率", "应缴费额(元)"]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(1, col_index).value = header
    row = [name, id_card, 3600, 0.01, 36]
    for col_index, value in enumerate(row, start=1):
        ws.cell(2, col_index).value = value
    workbook.save(path)
    workbook.close()


def _write_single_kind_rows(path: Path, rows: list[tuple[str, str, float, float, float]]) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "职工明细"
    headers = ["姓名", "证件号码", "缴费基数", "费率", "应缴费额(元)"]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(1, col_index).value = header
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index).value = value
    workbook.save(path)
    workbook.close()


def _write_long_payment_rows_with_nature(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "缴费明细"
    headers = ["姓名", "身份证件号码", "参保费种", "征收品目", "费款所属日期起", "缴费基数", "费率", "本期应缴费额", "业务类型"]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(1, col_index).value = header
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index).value = value
    workbook.save(path)
    workbook.close()


def _write_wide_payment_file(path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "缴费明细"
    ws.merge_cells("A1:F1")
    ws.merge_cells("G1:J1")
    ws.merge_cells("K1:L1")
    ws.merge_cells("M1:N1")
    ws["G1"] = "基本医疗保险费"
    ws["K1"] = "企业职工基本养老保险费"
    ws["M1"] = "失业保险费"
    ws["O1"] = "工伤保险费"
    headers = [
        "序号",
        "姓名",
        "证件类型",
        "证件号码",
        "费款所属期起",
        "费款所属期止",
        "职工基本医疗保险(单位缴纳)应缴费额(元)",
        "职工基本医疗保险(个人缴纳)应缴费额(元)",
        "职工大额医疗互助保险(单位缴纳)应缴费额(元)",
        "职工大额医疗互助保险(个人缴纳)应缴费额(元)",
        "职工基本养老保险(单位缴纳)应缴费额(元)",
        "职工基本养老保险(个人缴纳)应缴费额(元)",
        "失业保险(单位缴纳)应缴费额(元)",
        "失业保险(个人缴纳)应缴费额(元)",
        "工伤保险应缴费额(元)",
    ]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(2, col_index).value = header
    rows = [
        [1, "李四", "居民身份证", "360111199002020022", "2026-05", "2026-05", 70, 20, 5, 2, 100, 50, 10, 5, 8],
        [2, "李四", "居民身份证", "360111199002020022", "2026-06", "2026-06", 50, None, None, None, None, None, None, None, None],
    ]
    for row_index, row in enumerate(rows, start=3):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index).value = value
    workbook.save(path)
    workbook.close()


def _write_wide_amount_only_history_file(path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "缴费明细"
    headers = [
        "姓名",
        "证件号码",
        "费款所属期起",
        "费款所属期止",
        "职工基本养老保险(单位缴纳)应缴费额(元)",
    ]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(1, col_index).value = header
    rows = [
        ["李四", "360111199002020022", "2026-04", "2026-04", 100],
        ["李四", "360111199002020022", "2026-05", "2026-05", 100],
    ]
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index).value = value
    workbook.save(path)
    workbook.close()


if __name__ == "__main__":
    unittest.main()
