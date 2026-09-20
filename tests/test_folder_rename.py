from __future__ import annotations

import hashlib
import io
import json
import shutil
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from hr_toolkit.cli import main as cli_main
from hr_toolkit.tools.folder_rename import (
    FILE_TYPE_ALL,
    FILE_TYPE_EXTENSIONS,
    FILE_TYPE_FOLDER,
    FILE_TYPE_IMAGE,
    MODE_APPEND,
    MODE_EXCEL_BATCH,
    MODE_REMOVE,
    MODE_REPLACE,
    MODE_REPLACE_TEXT,
    _rename_text_no_replace,
    rename_files_by_excel,
    rename_person_folders,
)


class FolderRenameTest(unittest.TestCase):
    def test_replace_text_all_types_preserve_contents_extensions_and_unmatched_items(self) -> None:
        for file_type in FILE_TYPE_EXTENSIONS:
            with self.subTest(file_type=file_type), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                folder = root / "张三-劳动合同.v2"
                folder.mkdir()
                (folder / "劳动合同.txt").write_bytes(b"nested-content")
                names = ["王五-劳动合同.PDF", "李四-劳动合同.xlsx", "赵六-劳动合同.JpG",
                         "其他-劳动合同.bin", "劳动合同", "不匹配.pdf",
                         ".隐藏-劳动合同.pdf", "~$劳动合同.xlsx"]
                names += ["资料-劳动合同" + ext for ext in
                          dict.fromkeys(ext for extensions in FILE_TYPE_EXTENSIONS.values() for ext in extensions)]
                for name in names:
                    (root / name).write_bytes(name.encode("utf-8"))
                expected = {}
                if file_type in ("folder", "all"):
                    expected[folder.name] = "张三-资金合同.v2"
                for name in names:
                    path = root / name
                    if name.startswith((".", "~$")) or "劳动合同" not in path.stem:
                        continue
                    if file_type == "all" or (file_type != "folder" and path.suffix.lower() in FILE_TYPE_EXTENSIONS[file_type]):
                        expected[name] = name.replace("劳动合同", "资金合同")
                kwargs = dict(mode=MODE_REPLACE_TEXT, text="劳动合同", replacement_name="资金合同",
                              file_type=file_type, target_name="../ignored-hidden-field")
                preview = rename_person_folders(root, **kwargs, dry_run=True)
                self.assertEqual({op.source.name: op.target.name for op in preview.operations}, expected)
                self.assertTrue(all((root / name).exists() for name in names))
                result = rename_person_folders(root, **kwargs,
                    expected_operations=[(op.source.name, op.target.name) for op in preview.operations],
                    expected_warnings=preview.warnings)
                self.assertEqual(result.operations, preview.operations)
                for name in names:
                    target = root / expected.get(name, name)
                    self.assertEqual(target.read_bytes(), name.encode("utf-8"))
                    self.assertEqual(target.suffix, Path(name).suffix)
                    if name in expected:
                        self.assertFalse((root / name).exists())
                target_folder = root / expected.get(folder.name, folder.name)
                self.assertEqual((target_folder / "劳动合同.txt").read_bytes(), b"nested-content")

    def test_replace_text_matches_all_occurrences_and_preserves_literal_spaces(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name in ("Old Old.PDF", "old Old.v2.pdf", "Only.pdf"):
                (root / name).write_bytes(b"content")
            result = rename_person_folders(root, mode=MODE_REPLACE_TEXT, file_type="pdf",
                                          text="Old", replacement_name="New")
            self.assertEqual(result.operation_count, 2)
            self.assertTrue((root / "New New.PDF").is_file())
            self.assertTrue((root / "old New.v2.pdf").is_file())
            rename_person_folders(root, mode=MODE_REPLACE_TEXT, file_type="pdf",
                                  text=" New", replacement_name=" - New")
            self.assertTrue((root / "New - New.PDF").is_file())
            preview = rename_person_folders(root, mode=MODE_REPLACE_TEXT, file_type="pdf",
                                            text="PDF", replacement_name="txt", dry_run=True)
            self.assertEqual(preview.operation_count, 0)
            self.assertIn("没有找到", preview.warnings[0])

    def test_replace_text_does_not_prioritize_an_exact_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name in ("张三", "张三-劳动合同", "王五-身份证"):
                (root / name).mkdir()
            result = rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="张三", replacement_name="王五")
            self.assertEqual(result.operation_count, 2)
            self.assertEqual({path.name for path in root.iterdir()}, {"王五", "王五-劳动合同", "王五-身份证"})

    def test_replace_text_blocks_entire_batch_for_existing_or_duplicate_targets(self) -> None:
        cases = [
            (["a-old.pdf", "b-old.pdf", "b-new.pdf"], "old", "new", "pdf", "已存在"),
            (["a-old.pdf", "B-NEW.PDF", "b-old.pdf"], "old", "new", "pdf", "已存在"),
            (["xaaa.pdf", "xaaaa.pdf"], "aa", "a", "pdf", "重复"),
            (["xaaaA.pdf", "xaaaaa.PDF"], "aa", "a", "pdf", "重复"),
        ]
        for names, text, replacement, file_type, message in cases:
            with self.subTest(names=names), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                for name in names:
                    (root / name).write_bytes(name.encode())
                before = {p.name: p.read_bytes() for p in root.iterdir()}
                for preview in (True, False):
                    with self.assertRaisesRegex(ValueError, message):
                        rename_person_folders(root, mode=MODE_REPLACE_TEXT, text=text,
                                              replacement_name=replacement, file_type=file_type, dry_run=preview)
                    self.assertEqual({p.name: p.read_bytes() for p in root.iterdir()}, before)

    def test_replace_text_existing_folder_blocks_pdf_and_empty_folder_is_not_overwritten(self) -> None:
        for folder_source in (False, True):
            with self.subTest(folder_source=folder_source), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = root / "old.pdf"
                source.mkdir() if folder_source else source.write_bytes(b"source")
                (root / "new.pdf").mkdir()
                with self.assertRaisesRegex(ValueError, "已存在"):
                    rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new",
                                          file_type="folder" if folder_source else "pdf")
                self.assertTrue(source.exists())
                self.assertEqual(list((root / "new.pdf").iterdir()), [])

    def test_replace_text_invalid_names_and_inputs_never_mutate(self) -> None:
        for replacement in ("bad/name", "bad\\name", "bad:name", 'bad"name', "bad|name", "bad?name", "bad*name",
                            "bad<name", "bad>name", "bad\x00name", "CON", "NUL.txt", "..", ".", "bad.", "bad ", ""):
            with self.subTest(replacement=replacement), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                (root / "old").mkdir()
                with self.assertRaises(ValueError):
                    rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name=replacement)
                self.assertEqual([p.name for p in root.iterdir()], ["old"])
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for kwargs in (dict(text="", replacement_name="new"), dict(text="old", replacement_name="new", file_type="unknown")):
                with self.assertRaises(ValueError):
                    rename_person_folders(root, mode=MODE_REPLACE_TEXT, **kwargs)
            unchanged = rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="old")
            self.assertEqual(unchanged.operation_count, 0)
            self.assertIn("相同", unchanged.warnings[0])

    def test_replace_text_refuses_to_add_an_extension_to_extensionless_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "old").write_bytes(b"content")
            with self.assertRaisesRegex(ValueError, "扩展名"):
                rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new.txt", file_type="all")
            self.assertEqual((root / "old").read_bytes(), b"content")

    def test_replace_text_rejects_changed_preview_before_any_rename(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "old.pdf").write_bytes(b"one")
            kwargs = dict(mode=MODE_REPLACE_TEXT, text="old", replacement_name="new", file_type="pdf")
            preview = rename_person_folders(root, **kwargs, dry_run=True)
            (root / "another-old.pdf").write_bytes(b"two")
            with self.assertRaisesRegex(RuntimeError, "重新预览"):
                rename_person_folders(root, **kwargs, expected_operations=[(op.source.name, op.target.name) for op in preview.operations],
                                      expected_warnings=preview.warnings)
            self.assertEqual({p.name for p in root.iterdir()}, {"old.pdf", "another-old.pdf"})
            with self.assertRaisesRegex(RuntimeError, "重新预览"):
                rename_person_folders(root, **kwargs, expected_warnings=["changed"])

    def test_replace_text_native_rename_refuses_existing_targets(self) -> None:
        for is_folder in (False, True):
            with self.subTest(is_folder=is_folder), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source, target = root / "old", root / "new"
                if is_folder:
                    source.mkdir()
                    target.mkdir()
                    (source / "record.txt").write_bytes(b"source")
                else:
                    source.write_bytes(b"source")
                    target.write_bytes(b"target")
                with self.assertRaises(OSError):
                    _rename_text_no_replace(source, target)
                self.assertTrue(source.exists())
                if is_folder:
                    self.assertEqual((source / "record.txt").read_bytes(), b"source")
                    self.assertEqual(list(target.iterdir()), [])
                else:
                    self.assertEqual(source.read_bytes(), b"source")
                    self.assertEqual(target.read_bytes(), b"target")

    def test_replace_text_runtime_race_stops_without_overwriting_or_processing_remaining(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name in ("a-old.pdf", "b-old.pdf", "c-old.pdf"):
                (root / name).write_bytes(name.encode())
            def race(source, target):
                if source.name == "b-old.pdf":
                    target.write_bytes(b"appeared-after-check")
                _rename_text_no_replace(source, target)
            with patch("hr_toolkit.tools.folder_rename._rename_text_no_replace", side_effect=race):
                with self.assertRaisesRegex(RuntimeError, "已完成 1 项.*未处理 2 项"):
                    rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new", file_type="pdf")
            self.assertEqual((root / "a-new.pdf").read_bytes(), b"a-old.pdf")
            self.assertEqual((root / "b-new.pdf").read_bytes(), b"appeared-after-check")
            self.assertEqual((root / "b-old.pdf").read_bytes(), b"b-old.pdf")
            self.assertEqual((root / "c-old.pdf").read_bytes(), b"c-old.pdf")

    def test_replace_text_windows_uses_existing_non_overwriting_os_rename(self) -> None:
        source, target = Path("old.pdf"), Path("new.pdf")
        with patch("hr_toolkit.tools.folder_rename.sys.platform", "win32"), \
             patch("hr_toolkit.tools.folder_rename.os.rename", side_effect=FileExistsError) as rename:
            with self.assertRaises(FileExistsError):
                _rename_text_no_replace(source, target)
            rename.assert_called_once_with(source, target)

    def test_replace_text_cancellation_stops_remaining_items(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "a-old.pdf").write_bytes(b"first")
            (root / "b-old.pdf").write_bytes(b"second")
            with self.assertRaisesRegex(RuntimeError, "已停止.*已完成 1 项.*未处理 1 项"):
                rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new",
                                      file_type="pdf", cancelled=lambda: (root / "a-new.pdf").exists())
            self.assertEqual((root / "a-new.pdf").read_bytes(), b"first")
            self.assertEqual((root / "b-old.pdf").read_bytes(), b"second")

    def test_replace_text_unsupported_exclusive_rename_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "old.pdf").write_bytes(b"original")
            with patch("hr_toolkit.tools.folder_rename.sys.platform", "unsupported"), \
                 self.assertRaisesRegex(RuntimeError, "不支持安全.*已完成 0 项"):
                rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new", file_type="pdf")
            self.assertEqual((root / "old.pdf").read_bytes(), b"original")
            self.assertFalse((root / "new.pdf").exists())

    def test_replace_text_links_are_blocked_without_modifying_referents(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "items"
            root.mkdir()
            outside = base / "outside"
            outside.mkdir()
            (outside / "record.txt").write_bytes(b"original")
            try:
                (root / "old").symlink_to(outside, target_is_directory=True)
            except (OSError, NotImplementedError):
                self.skipTest("Symlink creation is unavailable")
            with self.assertRaisesRegex(ValueError, "链接"):
                rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new")
            self.assertTrue((root / "old").is_symlink())
            self.assertEqual((outside / "record.txt").read_bytes(), b"original")
            (root / "old").unlink()
            (root / "old").mkdir()
            (root / "new").symlink_to(base / "missing")
            with self.assertRaisesRegex(ValueError, "已存在"):
                rename_person_folders(root, mode=MODE_REPLACE_TEXT, text="old", replacement_name="new")
            self.assertTrue((root / "old").is_dir())
            self.assertTrue((root / "new").is_symlink())

    def test_cli_replace_text_preview_and_apply_keep_type_selection(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "old.PDF").write_bytes(b"pdf")
            (root / "old.jpg").write_bytes(b"image")
            args = ["folder-rename", "--root", str(root), "--mode", "replace_text", "--file-type", "pdf",
                    "--text", "old", "--replacement", "new", "--json"]
            for apply in (False, True):
                output = io.StringIO()
                with redirect_stdout(output):
                    self.assertEqual(cli_main(args + (["--apply"] if apply else [])), 0)
                payload = json.loads(output.getvalue())
                self.assertEqual(payload["operation_count"], 1)
                self.assertEqual(payload["dry_run"], not apply)
                self.assertEqual(payload["operations"][0]["target_name"], "new.PDF")
                self.assertEqual((root / ("new.PDF" if apply else "old.PDF")).read_bytes(), b"pdf")
                self.assertEqual((root / "old.jpg").read_bytes(), b"image")

    @staticmethod
    def _write_name_workbook(path: Path, names: list[str], *, header_row: int = 1) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.cell(header_row, 1, "姓名")
        for row, name in enumerate(names, start=header_row + 1):
            worksheet.cell(row, 1, name)
        workbook.save(path)
        workbook.close()

    @staticmethod
    def _sha256(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def test_append_suffix_to_all_folders(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三").mkdir()
            (root / "李四").mkdir()
            (root / "说明.txt").write_text("ignore", encoding="utf-8")

            # “-劳动合同”带前缀追加;右侧注释验证“劳动合同”不带前缀时直传
            preview = rename_person_folders(root, mode=MODE_APPEND, text="-劳动合同", dry_run=True)
            self.assertEqual(preview.operation_count, 2)
            self.assertTrue((root / "张三").exists())

            result = rename_person_folders(root, mode=MODE_APPEND, text="-劳动合同")

            self.assertEqual(result.operation_count, 2)
            self.assertTrue((root / "张三-劳动合同").exists())
            self.assertTrue((root / "李四-劳动合同").exists())
            self.assertTrue((root / "说明.txt").exists())

    def test_append_text_passed_through_unchanged(self) -> None:
        """bug4 修复:用户输入什么就追加什么，不自动加分隔符"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三").mkdir()

            result = rename_person_folders(root, mode=MODE_APPEND, text="劳动合同")

            self.assertEqual(result.operation_count, 1)
            self.assertTrue((root / "张三劳动合同").exists())
            self.assertFalse((root / "张三-劳动合同").exists())

    def test_append_suffix_to_one_person(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三").mkdir()
            (root / "李四").mkdir()

            result = rename_person_folders(root, mode=MODE_APPEND, text="-身份证", target_name="张三")

            self.assertEqual(result.operation_count, 1)
            self.assertTrue((root / "张三-身份证").exists())
            self.assertTrue((root / "李四").exists())

    def test_remove_suffix_variants_from_all_folders(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三-劳动合同").mkdir()
            (root / "李四_劳动合同").mkdir()
            (root / "赵露思劳动合同").mkdir()
            (root / "王五-身份证").mkdir()

            result = rename_person_folders(root, mode=MODE_REMOVE, text="_劳动合同")

            self.assertEqual(result.operation_count, 3)
            self.assertTrue((root / "张三").exists())
            self.assertTrue((root / "李四").exists())
            self.assertTrue((root / "赵露思").exists())
            self.assertTrue((root / "王五-身份证").exists())

    def test_remove_suffix_from_one_person(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三_身份证").mkdir()
            (root / "李四_身份证").mkdir()

            result = rename_person_folders(root, mode=MODE_REMOVE, text="身份证", target_name="张三")

            self.assertEqual(result.operation_count, 1)
            self.assertTrue((root / "张三").exists())
            self.assertTrue((root / "李四_身份证").exists())

    def test_replace_one_folder_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三").mkdir()
            (root / "李四").mkdir()

            result = rename_person_folders(
                root,
                mode=MODE_REPLACE,
                target_name="张三",
                replacement_name="章五",
            )

            self.assertEqual(result.operation_count, 1)
            self.assertFalse((root / "张三").exists())
            self.assertTrue((root / "章五").exists())
            self.assertTrue((root / "李四").exists())

    def test_manual_modes_reject_paths_outside_the_selected_root(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "张三").mkdir()
            outside = base / "外部目录"
            outside.mkdir()

            for mode, kwargs in (
                (MODE_APPEND, {"text": "-合同", "target_name": "../外部目录"}),
                (MODE_REMOVE, {"text": "合同", "target_name": "../外部目录"}),
                (MODE_REPLACE, {"target_name": "../外部目录", "replacement_name": "新名称"}),
            ):
                with self.subTest(mode=mode), self.assertRaisesRegex(ValueError, "Windows 不支持的字符"):
                    rename_person_folders(root, mode=mode, **kwargs)

            self.assertTrue(outside.is_dir())
            self.assertTrue((root / "张三").is_dir())

    def test_manual_modes_reject_windows_reserved_target_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三").mkdir()

            with self.assertRaisesRegex(ValueError, "Windows 保留名称"):
                rename_person_folders(
                    root,
                    mode=MODE_REPLACE,
                    target_name="张三",
                    replacement_name="CON",
                )

    def test_skip_existing_target(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "张三").mkdir()
            (root / "张三-劳动合同").mkdir()

            result = rename_person_folders(root, mode=MODE_APPEND, text="-劳动合同")

            self.assertEqual(result.operation_count, 0)
            self.assertTrue(any("已存在" in warning or "已包含后缀" in warning for warning in result.warnings))

    def test_excel_batch_preview_and_execute_change_only_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            sources = {
                "10.png": b"image-ten",
                "2.jpg": b"image-two",
                "1.pdf": b"pdf-one",
            }
            for name, content in sources.items():
                (root / name).write_bytes(content)
            (root / ".DS_Store").write_bytes(b"hidden")
            (root / "~$临时.docx").write_bytes(b"temporary")
            roster = base / "人员名单.xlsx"
            self._write_name_workbook(roster, ["张三", "李四", "王五"])

            source_hashes = {name: self._sha256(root / name) for name in sources}
            roster_hash = self._sha256(roster)
            preview = rename_files_by_excel(root, roster, file_type=FILE_TYPE_ALL, dry_run=True)

            self.assertEqual(preview.mode, MODE_EXCEL_BATCH)
            self.assertEqual(
                [operation.source.name for operation in preview.operations],
                ["1.pdf", "2.jpg", "10.png"],
            )
            self.assertEqual(
                [operation.target.name for operation in preview.operations],
                ["张三.pdf", "李四.jpg", "王五.png"],
            )
            self.assertEqual({name: self._sha256(root / name) for name in sources}, source_hashes)
            self.assertEqual(self._sha256(roster), roster_hash)

            result = rename_files_by_excel(
                root,
                roster,
                file_type=FILE_TYPE_ALL,
                expected_operations=[
                    (operation.source.name, operation.target.name)
                    for operation in preview.operations
                ],
                expected_warnings=list(preview.warnings),
            )

            self.assertEqual(result.operation_count, 3)
            for operation in preview.operations:
                self.assertFalse(operation.source.exists())
                self.assertEqual(
                    self._sha256(operation.target),
                    source_hashes[operation.source.name],
                )
                self.assertEqual(operation.source.suffix, operation.target.suffix)
            self.assertTrue((root / ".DS_Store").is_file())
            self.assertTrue((root / "~$临时.docx").is_file())
            self.assertEqual(self._sha256(roster), roster_hash)

    def test_excel_batch_fewer_names_keeps_extra_item_and_warns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.pdf").write_bytes(b"one")
            (root / "2.pdf").write_bytes(b"two")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三"])

            result = rename_files_by_excel(root, roster)

            self.assertTrue((root / "张三.pdf").is_file())
            self.assertEqual((root / "张三.pdf").read_bytes(), b"one")
            self.assertEqual((root / "2.pdf").read_bytes(), b"two")
            self.assertTrue(any("少 1 个" in warning and "2.pdf" in warning for warning in result.warnings))

    def test_excel_batch_more_names_reports_every_unmatched_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.pdf").write_bytes(b"one")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三", "李四", "王五"])

            preview = rename_files_by_excel(root, roster, dry_run=True)

            self.assertEqual(preview.operation_count, 1)
            self.assertTrue(
                any("多 2 人" in warning and "李四" in warning and "王五" in warning for warning in preview.warnings)
            )
            self.assertTrue((root / "1.pdf").is_file())

    def test_excel_batch_file_type_filter_keeps_nonmatching_items(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.jpg").write_bytes(b"jpg")
            (root / "2.pdf").write_bytes(b"pdf")
            (root / "3.png").write_bytes(b"png")
            (root / "4").mkdir()
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三", "李四"])

            result = rename_files_by_excel(root, roster, file_type=FILE_TYPE_IMAGE)

            self.assertEqual(result.operation_count, 2)
            self.assertEqual((root / "张三.jpg").read_bytes(), b"jpg")
            self.assertEqual((root / "李四.png").read_bytes(), b"png")
            self.assertEqual((root / "2.pdf").read_bytes(), b"pdf")
            self.assertTrue((root / "4").is_dir())

    def test_excel_batch_folder_filter_preserves_folder_contents(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            (root / "档案10").mkdir(parents=True)
            (root / "档案2").mkdir()
            (root / "档案2" / "说明.txt").write_text("record", encoding="utf-8")
            (root / "忽略.pdf").write_bytes(b"pdf")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三", "李四"])

            result = rename_files_by_excel(root, roster, file_type=FILE_TYPE_FOLDER)

            self.assertEqual(result.operation_count, 2)
            self.assertEqual((root / "张三" / "说明.txt").read_text(encoding="utf-8"), "record")
            self.assertTrue((root / "李四").is_dir())
            self.assertTrue((root / "忽略.pdf").is_file())

    def test_excel_batch_excludes_roster_copy_inside_root(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.pdf").write_bytes(b"one")
            (root / "2.pdf").write_bytes(b"two")
            external_roster = base / "留存" / "名单.xlsx"
            external_roster.parent.mkdir()
            self._write_name_workbook(external_roster, ["张三", "李四"])
            roster_copy = root / external_roster.name
            shutil.copy2(external_roster, roster_copy)
            roster_hash = self._sha256(roster_copy)

            preview = rename_files_by_excel(root, external_roster, file_type=FILE_TYPE_ALL, dry_run=True)

            self.assertEqual([item.source.name for item in preview.operations], ["1.pdf", "2.pdf"])
            self.assertEqual(self._sha256(roster_copy), roster_hash)

    def test_excel_batch_existing_target_is_never_overwritten(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.pdf").write_bytes(b"source")
            (root / "张三.pdf").mkdir()
            (root / "张三.pdf" / "原内容.txt").write_text("keep", encoding="utf-8")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三"])

            result = rename_files_by_excel(root, roster)

            self.assertEqual(result.operation_count, 0)
            self.assertEqual((root / "1.pdf").read_bytes(), b"source")
            self.assertEqual((root / "张三.pdf" / "原内容.txt").read_text(encoding="utf-8"), "keep")
            self.assertTrue(any("已存在" in warning for warning in result.warnings))

    def test_excel_batch_duplicate_and_invalid_names_do_not_shift_pairing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            for index in range(1, 4):
                (root / f"{index}.pdf").write_bytes(str(index).encode())
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["坏/名", "李四", "李四"])

            result = rename_files_by_excel(root, roster)

            self.assertEqual(result.operation_count, 1)
            self.assertEqual((root / "1.pdf").read_bytes(), b"1")
            self.assertEqual((root / "李四.pdf").read_bytes(), b"2")
            self.assertEqual((root / "3.pdf").read_bytes(), b"3")
            self.assertTrue(any("不能用于改名" in warning for warning in result.warnings))
            self.assertTrue(any("目标名称重复" in warning for warning in result.warnings))

    def test_excel_batch_no_matching_items_returns_explicit_warning(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "说明.txt").write_text("keep", encoding="utf-8")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三", "李四"])

            preview = rename_files_by_excel(root, roster, file_type=FILE_TYPE_IMAGE, dry_run=True)

            self.assertEqual(preview.operation_count, 0)
            self.assertTrue(any("多 2 人" in warning for warning in preview.warnings))
            self.assertTrue((root / "说明.txt").is_file())

    def test_excel_batch_finds_name_header_below_title_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.pdf").write_bytes(b"one")
            roster = base / "名单.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.cell(1, 1, "项目人员姓名名单")
            worksheet.cell(3, 1, "姓名")
            worksheet.cell(4, 1, "张三")
            workbook.save(roster)
            workbook.close()

            preview = rename_files_by_excel(root, roster, dry_run=True)

            self.assertEqual(preview.operation_count, 1)
            self.assertEqual(preview.operations[0].target.name, "张三.pdf")

    def test_excel_batch_rejects_windows_reserved_and_case_duplicate_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            for index in range(1, 4):
                (root / f"{index}.pdf").write_bytes(str(index).encode())
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["CON", "Alice", "ALICE"])

            preview = rename_files_by_excel(root, roster, dry_run=True)

            self.assertEqual(preview.operation_count, 1)
            self.assertEqual(preview.operations[0].source.name, "2.pdf")
            self.assertEqual(preview.operations[0].target.name, "Alice.pdf")
            self.assertTrue(any("Windows 保留名称" in warning for warning in preview.warnings))
            self.assertTrue(any("目标名称重复" in warning for warning in preview.warnings))

    def test_excel_batch_aborts_when_confirmed_preview_changes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            (root / "1.pdf").write_bytes(b"one")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三"])
            preview = rename_files_by_excel(root, roster, dry_run=True)
            expected_operations = [
                (operation.source.name, operation.target.name)
                for operation in preview.operations
            ]
            (root / "0.pdf").write_bytes(b"new")

            with self.assertRaisesRegex(RuntimeError, "预览确认后发生了变化"):
                rename_files_by_excel(
                    root,
                    roster,
                    expected_operations=expected_operations,
                    expected_warnings=list(preview.warnings),
                )

            self.assertEqual((root / "0.pdf").read_bytes(), b"new")
            self.assertEqual((root / "1.pdf").read_bytes(), b"one")
            self.assertFalse((root / "张三.pdf").exists())

    def test_cli_excel_mode_previews_without_changing_source(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "待改名"
            root.mkdir()
            source = root / "1.pdf"
            source.write_bytes(b"source")
            roster = base / "名单.xlsx"
            self._write_name_workbook(roster, ["张三"])
            output = io.StringIO()

            with redirect_stdout(output):
                exit_code = cli_main(
                    [
                        "folder-rename",
                        "--root",
                        str(root),
                        "--mode",
                        "excel",
                        "--excel",
                        str(roster),
                        "--file-type",
                        "pdf",
                        "--json",
                    ]
                )

            payload = json.loads(output.getvalue())
            self.assertEqual(exit_code, 0)
            self.assertTrue(payload["dry_run"])
            self.assertEqual(payload["operations"][0]["source_name"], "1.pdf")
            self.assertEqual(payload["operations"][0]["target_name"], "张三.pdf")
            self.assertEqual(source.read_bytes(), b"source")


if __name__ == "__main__":
    unittest.main()
