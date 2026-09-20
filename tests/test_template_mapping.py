"""输入适配入口及只读映射的回归用例；不依赖业务原始数据。"""
from __future__ import annotations

import inspect
from io import BytesIO
import json
import re
import tempfile
import unittest
from unittest.mock import patch
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook

from hr_toolkit.common.excel import SheetGrid
from hr_toolkit.common.template_mapping import (
    SUPPORTED_TOOLS, TemplateSelectionRequired, active, catalog, choose_sheet,
    clean_rules, ignored_sheet, map_sheet, request_selection, save_choice, template_tool,
    resolve_sheet_roles, request_sheet_selection, unused_sheet_notices, choose_content_sheet,
)


def sheet(headers, data=None, name="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws.append(headers)
    if data is not None:
        ws.append(data)
    return ws


def invoke(tool, callback, rules=None):
    return template_tool(tool)(callback)(template_rules={} if rules is None else rules)


class TemplateMappingTest(unittest.TestCase):
    def test_same_basename_files_do_not_share_absence_confirmation(self):
        from hr_toolkit.tools.personnel_change_merge import _read_change_file, TARGET_SHEETS
        with tempfile.TemporaryDirectory() as tmp:
            paths = [Path(tmp) / company / "异动.xlsx" for company in ("甲公司", "乙公司")]
            for path in paths:
                path.parent.mkdir()
                wb = Workbook()
                wb.active.append(["序号", "姓名", "入职日期"])
                wb.save(path)
                wb.close()
            with self.assertRaises(TemplateSelectionRequired) as first:
                invoke("personnel_change_merge", lambda: _read_change_file(paths[0]))
            saved = save_choice("personnel_change_merge", {}, first.exception.payload,
                                {"sheet_selections": dict.fromkeys(TARGET_SHEETS)})
            invoke("personnel_change_merge", lambda: _read_change_file(paths[0]), saved)
            with self.assertRaises(TemplateSelectionRequired):
                invoke("personnel_change_merge", lambda: _read_change_file(paths[1]), saved)

    def test_archive_and_conversion_identity_survives_retry(self):
        from hr_toolkit.common.inputs import extract_archive_excel_files
        from hr_toolkit.common.excel_compat import ensure_xlsx_workbook
        from hr_toolkit.tools.personnel_change_merge import _read_change_file, TARGET_SHEETS
        with tempfile.TemporaryDirectory() as tmp:
            archive = Path(tmp) / "资料.zip"
            wb = Workbook()
            wb.active.append(["说明"])
            content = BytesIO()
            wb.save(content)
            wb.close()
            with ZipFile(archive, "w") as packed:
                # XLSX content with an XLS suffix exercises the conversion path.
                for name in ("甲/异动.xls", "乙/异动.xls"):
                    packed.writestr(name, content.getvalue())
            def read(index):
                with tempfile.TemporaryDirectory() as work:
                    files = extract_archive_excel_files(archive, Path(work), [])
                    return _read_change_file(ensure_xlsx_workbook(files[index], Path(work)))
            with self.assertRaises(TemplateSelectionRequired) as first:
                invoke("personnel_change_merge", lambda: read(0))
            saved = save_choice("personnel_change_merge", {}, first.exception.payload,
                                {"sheet_selections": dict.fromkeys(TARGET_SHEETS)})
            invoke("personnel_change_merge", lambda: read(0), saved)
            with self.assertRaises(TemplateSelectionRequired):
                invoke("personnel_change_merge", lambda: read(1), saved)

    def test_content_selection_recovers_declared_range_before_choosing_sheet(self):
        wb = Workbook()
        wb.active.title = "本期缴费"
        wb.active.append(["姓名", "证件号码"])
        wb.active.append(["本期员工", "123"])
        wb.create_sheet("历史备查").append(["姓名", "证件号码"])
        original, damaged = BytesIO(), BytesIO()
        wb.save(original)
        wb.close()
        with ZipFile(original) as source, ZipFile(damaged, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = re.sub(rb'<dimension[^>]+>', b'<dimension ref="A1:A1"/>', data)
                target.writestr(item, data)
        damaged.seek(0)
        loaded = load_workbook(damaged, read_only=True, data_only=True)
        self.addCleanup(loaded.close)
        selected = invoke("social_security", lambda: choose_content_sheet(loaded.worksheets, "payment", loaded.worksheets[0]))
        self.assertEqual(selected.title, "本期缴费")
        grid = SheetGrid(selected)
        self.assertEqual(grid.value(2, 1), "本期员工")
        self.assertTrue(grid.dimension_recovered)

    def test_extra_sheet_is_logged_without_interrupting_resolved_roles(self):
        ws = sheet(["说明"], name="说明")
        wb = ws.parent
        self.addCleanup(wb.close)
        defaults = {name: wb.create_sheet(name) for name in ("增员", "减员", "转正", "调动")}
        def read():
            selected = resolve_sheet_roles(wb.worksheets, defaults, file="异动.xlsx", optional=tuple(defaults))
            unused_sheet_notices(wb.worksheets, {page.title for page in selected.values()}, "异动.xlsx")
            unused_sheet_notices(wb.worksheets, {page.title for page in selected.values()}, "异动.xlsx")
            return SimpleNamespace(warnings=[])
        result = invoke("personnel_change_merge", read)
        self.assertEqual(len(result.warnings), 1)
        self.assertIn("「说明」", result.warnings[0])

    def test_missing_sheets_are_confirmed_together_and_absence_is_temporary(self):
        ws = sheet(["任意列"], name="新增人员")
        wb = ws.parent
        self.addCleanup(wb.close)
        defaults = {"增员": None, "减员": None, "转正": wb.create_sheet("转正"), "调动": wb.create_sheet("调动")}
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("personnel_change_merge", lambda: resolve_sheet_roles(wb.worksheets, defaults, file="异动.xlsx", optional=tuple(defaults)))
        issue = caught.exception.payload
        self.assertEqual(issue["kind"], "worksheets")
        self.assertEqual([r["key"] for r in issue["sheet_requests"]], ["增员", "减员"])
        rules = save_choice("personnel_change_merge", {}, issue,
                            {"sheet_selections": {"增员": "新增人员", "减员": None}, "remember": True})
        selected = invoke("personnel_change_merge", lambda: resolve_sheet_roles(wb.worksheets, defaults, file="异动.xlsx", optional=tuple(defaults)), rules)
        self.assertIs(selected["增员"], ws)
        self.assertIsNone(selected["减员"])
        persistent = {k: v for k, v in rules.items() if k != "sheet_choices"}
        self.assertIn("新增人员", persistent["sheets"]["增员"])
        with self.assertRaises(TemplateSelectionRequired) as later:
            invoke("personnel_change_merge", lambda: resolve_sheet_roles(wb.worksheets, defaults, file="下月.xlsx", optional=tuple(defaults)), persistent)
        self.assertEqual([r["key"] for r in later.exception.payload["sheet_requests"]], ["减员"])

    def test_sheet_selection_rejects_required_absence_and_duplicate_consumption(self):
        ws = sheet(["列"], name="资料")
        self.addCleanup(ws.parent.close)
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("salary_split", lambda: resolve_sheet_roles(ws.parent.worksheets, {"detail": None, "summary": None}))
        for selection in ({"detail": "资料", "summary": None}, {"detail": "资料", "summary": "资料"}):
            with self.subTest(selection=selection), self.assertRaises(ValueError):
                save_choice("salary_split", {}, caught.exception.payload, {"sheet_selections": selection})

    def test_single_role_tools_reuse_page_choice_without_binding_column_names(self):
        for tool, role in (("archive_import", "transfer"), ("insurance_ledger", "policy"),
                           ("social_security", "payment"), ("data_statistics", "staff")):
            with self.subTest(tool=tool):
                ws = sheet(["旧列"], name="1")
                self.addCleanup(ws.parent.close)
                ws.parent.create_sheet("说明")
                with self.assertRaises(TemplateSelectionRequired) as caught:
                    invoke(tool, lambda: request_sheet_selection(ws.parent.worksheets, [role], file="来源.xlsx"))
                rules = save_choice(tool, {}, caught.exception.payload,
                                    {"sheet_selections": {role: "1"}, "remember": True})
                rules.pop("sheet_choices")
                ws.cell(1, 1).value = "另一套列名"
                self.assertIs(invoke(tool, lambda: choose_sheet(ws.parent.worksheets, role, file="下月.xlsx"), rules), ws)
                self.assertEqual(rules["profiles"], [])

    def test_content_reader_keeps_usable_original_default(self):
        ws = sheet(["姓名", "证件号码"], name="原表")
        self.addCleanup(ws.parent.close)
        other = ws.parent.create_sheet("另一个有效页")
        other.append(["姓名", "证件号码"])
        selected = invoke("social_security", lambda: choose_content_sheet(ws.parent.worksheets, "payment", ws))
        self.assertIs(selected, ws)

    def test_statistics_extra_instruction_sheet_does_not_change_business_rows(self):
        from hr_toolkit.tools.data_statistics import _read_statistics_file
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "考勤202608.xlsx"
            ws = sheet(["姓名", "应出勤天数", "实际出勤天数"], ["测试人员", 20, 19], name="考勤表")
            self.addCleanup(ws.parent.close)
            ws.parent.save(path)
            baseline = invoke("data_statistics", lambda: _read_statistics_file(path, []))
            extra = ws.parent.create_sheet("说明", 0)
            extra.append(["姓名", "备注"])
            extra.append(["示例", "这里只是填写说明"])
            ws.parent.save(path)
            original = path.read_bytes()
            result = invoke("data_statistics", lambda: SimpleNamespace(value=_read_statistics_file(path, []), warnings=[]))
            self.assertEqual(result.value, baseline)
            self.assertEqual(len(result.warnings), 1)
            self.assertIn("「说明」", result.warnings[0])
            self.assertEqual(path.read_bytes(), original)

    def test_archive_export_preserves_multiple_company_pages_and_logs_extra(self):
        from hr_toolkit.tools.archive_import import _read_archive_summary_records
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "档案汇总.xlsx"
            ws = sheet(["姓名", "身份证"], ["甲", "TEST-001"], name="甲公司")
            self.addCleanup(ws.parent.close)
            other = ws.parent.create_sheet("乙公司")
            other.append(["姓名", "身份证"])
            other.append(["乙", "TEST-002"])
            ws.parent.save(path)
            baseline, _ = _read_archive_summary_records([path], [])
            ws.parent.create_sheet("说明", 0).append(["填写说明"])
            ws.parent.save(path)
            result = invoke("archive_export", lambda: SimpleNamespace(value=_read_archive_summary_records([path], []), warnings=[]))
            self.assertEqual(result.value[0], baseline)
            self.assertEqual(len(result.value[0]), 2)
            self.assertEqual(len(result.warnings), 1)
            self.assertIn("「说明」", result.warnings[0])

    def test_summary_with_renamed_name_suggests_summary_without_changing_parser(self):
        from hr_toolkit.tools.data_statistics import _adapt_statistics_grid
        ws = sheet(["222", "应出勤天数", "实际出勤天数", "请假天数"], ["张三", 19, 19, 0], name="考勤表")
        self.addCleanup(ws.parent.close)
        grid = SheetGrid(ws)
        self.assertIsNone(invoke("data_statistics", lambda: _adapt_statistics_grid(grid, "考勤.xlsx")))
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("data_statistics", lambda: request_selection([grid], ["attendance", "attendance_summary", "weekly", "monthly"], file="考勤.xlsx"))
        issue = caught.exception.payload
        self.assertEqual(issue["suggested_role"], "attendance_summary")
        spec = next(role for role in issue["roles"] if role["key"] == issue["suggested_role"])
        self.assertEqual(spec["required"], ["姓名"])
        rules = save_choice("data_statistics", {}, issue,
                            {"role": "attendance_summary", "sheet": "考勤表", "row": 1, "columns": {"姓名": 1}})
        mapped = invoke("data_statistics", lambda: _adapt_statistics_grid(grid, "考勤.xlsx"), rules)
        self.assertEqual(mapped.mapping_role, "attendance_summary")
        self.assertEqual(mapped.cell(2, 1).value, "张三")
        self.assertEqual(ws.cell(1, 1).value, "222")

    def test_role_suggestion_does_not_guess_ambiguous_weekly_monthly(self):
        ws = sheet(["汇报编号", "汇报人", "汇报时间"])
        self.addCleanup(ws.parent.close)
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("data_statistics", lambda: request_selection([ws], ["weekly", "monthly"]))
        self.assertEqual(caught.exception.payload["suggested_role"], "")

    def test_saved_choice_can_be_edited_using_headers_without_saving_personal_rows(self):
        from hr_toolkit.common.template_mapping import profile_issue, saved_choices
        ws = sheet(["222", "名字", "应出勤天数"], [28, "测试人员", 20], name="考勤表")
        self.addCleanup(ws.parent.close)
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("data_statistics", lambda: map_sheet(ws, "attendance_summary"))
        rules = save_choice("data_statistics", {}, caught.exception.payload,
                            {"role": "attendance_summary", "sheet": "考勤表", "row": 1, "columns": {"姓名": 1}})
        profile = rules["profiles"][0]
        self.assertNotIn("测试人员", json.dumps(profile, ensure_ascii=False))
        self.assertIn("姓名 ← 222", saved_choices("data_statistics", rules)[0]["description"])
        issue = profile_issue("data_statistics", rules, profile["key"])
        updated = save_choice("data_statistics", rules, issue,
                              {"role": "attendance_summary", "sheet": "考勤表", "row": 1, "columns": {"姓名": 2}})
        self.assertEqual(len(updated["profiles"]), 1)
        mapped = invoke("data_statistics", lambda: map_sheet(ws, "attendance_summary"), updated)
        self.assertEqual(mapped.cell(1, 1).value, "222")
        self.assertEqual(mapped.cell(1, 2).value, "姓名")
        self.assertEqual(mapped.cell(2, 2).value, "测试人员")
        legacy = {"profiles": [{k: v for k, v in profile.items() if k not in {"headers", "required", "one_of", "file"}}]}
        self.assertFalse(saved_choices("data_statistics", legacy)[0]["editable"])
        self.assertEqual(invoke("data_statistics", lambda: map_sheet(ws, "attendance_summary"), legacy).cell(1, 1).value, "姓名")
        with self.assertRaisesRegex(ValueError, "删除后重新处理"):
            profile_issue("data_statistics", legacy, profile["key"])

    def test_normalization_is_reused_without_changing_source_or_column_priority(self):
        from hr_toolkit.common import template_mapping as mapping
        marker = " 唯一未使用表头 "
        ws = sheet(["公司", "姓名", "姓名", "身份证", marker], ["甲", "原姓名", "后姓名", "TEST-001", "值"])
        with patch.object(mapping, "normalize_alias", wraps=mapping.normalize_alias) as normalize:
            mapped = invoke("archive_import", lambda: map_sheet(ws, "transfer"))
        self.assertEqual(sum(call.args == (marker,) for call in normalize.call_args_list), 1)
        self.assertEqual(mapped.header_row, 1)
        self.assertEqual(mapped.changes, {})
        self.assertEqual(ws.cell(1, 5).value, marker)
        self.assertEqual(mapped.cell(2, 2).value, "原姓名")
        self.assertEqual(mapped.cell(2, 3).value, "后姓名")
        ws.parent.close()

    def test_prompt_labels_do_not_change_internal_fields_or_source(self):
        ws = sheet(["名字", "证件编号"], ["示例人员", "TEST-001"], name="增员")
        ws.parent.create_sheet("说明", 0)
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("social_security", lambda: map_sheet(ws, "roster", file="异动表.xlsx", source_sheets=ws.parent.worksheets))
        payload = caught.exception.payload
        self.assertEqual(payload["selected_sheet"], "增员")
        self.assertEqual(payload["sheets"][0]["name"], "说明")
        role = payload["roles"][0]
        self.assertEqual(role["field_labels"]["*姓名.简体中文"], "姓名")
        self.assertEqual(role["field_labels"]["*身份证"], "身份证号码")
        self.assertIn("*姓名.简体中文", role["required"])
        self.assertEqual(role["fields"]["*姓名.简体中文"], ["*姓名.简体中文"])
        self.assertNotIn("简体中文", payload["message"])
        self.assertEqual(ws.cell(1, 1).value, "名字")

    def test_prompt_still_rejects_missing_and_duplicate_required_columns(self):
        ws = sheet(["名字", "证件编号"], ["示例人员", "TEST-001"])
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("social_security", lambda: map_sheet(ws, "roster"))
        payload = caught.exception.payload
        base = {"role": "roster", "sheet": "Sheet1", "row": 1}
        with self.assertRaisesRegex(ValueError, "身份证号码"):
            save_choice("social_security", {}, payload, {**base, "columns": {"*姓名.简体中文": 1}})
        with self.assertRaisesRegex(ValueError, "同一列"):
            save_choice("social_security", {}, payload, {**base, "columns": {"*姓名.简体中文": 1, "*身份证": 1}})

    def test_all_entrypoints_accept_serializable_rules(self):
        from hr_toolkit.tools.registry import get_tool_by_id
        self.assertEqual(set(SUPPORTED_TOOLS), {"salary_split", "personnel_change_merge", "roster_update", "archive_import", "archive_export", "insurance_ledger", "data_statistics", "social_security"})
        for tool in SUPPORTED_TOOLS:
            with self.subTest(tool=tool):
                function = get_tool_by_id(tool).entry_point
                self.assertIn("template_rules", inspect.signature(function).parameters)
                self.assertTrue(catalog(tool))
                self.assertEqual(clean_rules(tool, {}), {"fields": {}, "sheets": {}, "profiles": []})

    def test_three_alternative_names_and_read_only_grid_view(self):
        rules = {"fields": {"transfer|姓名": ["姓名", "名字", "name"]}}
        for name in ("姓名", "名字", "ＮＡＭＥ"):
            ws = sheet(["公司", name, "身份证"], ["甲公司", "测试人员", "TEST-001"])
            grid = SheetGrid(ws)
            mapped = invoke("archive_import", lambda: map_sheet(grid, "transfer"), rules)
            self.assertEqual(mapped.cell(1, 2).value, "姓名")
            self.assertEqual(mapped.value(2, 2), "测试人员")
            self.assertEqual(grid.cell(1, 2).value, name)
            self.assertEqual(ws.cell(1, 2).value, name)

    def test_missing_name_cannot_be_swallowed_as_nonbusiness_value_error(self):
        ws = sheet(["公司", "新称呼", "身份证"])
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("archive_import", lambda: map_sheet(ws, "transfer"))
        self.assertNotIsInstance(caught.exception, ValueError)
        self.assertEqual(caught.exception.payload["roles"][0]["required"], ["公司", "姓名", "身份证"])
        self.assertFalse(active())

    def test_ambiguous_columns_require_choice_and_saved_choice_is_structural(self):
        ws = sheet(["公司", "姓名", "名字", "身份证"], ["甲公司", "甲", "乙", "TEST-001"])
        rules = {"fields": {"transfer|姓名": ["姓名", "名字"]}}
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("archive_import", lambda: map_sheet(ws, "transfer"), rules)
        saved = save_choice("archive_import", rules, caught.exception.payload,
                            {"role": "transfer", "sheet": "Sheet1", "row": 1,
                             "columns": {"公司": 1, "姓名": 3, "身份证": 4}})
        mapped = invoke("archive_import", lambda: map_sheet(ws, "transfer"), saved)
        self.assertEqual(mapped.cell(1, 3).value, "姓名")
        self.assertEqual(mapped.cell(1, 2).value, "")
        self.assertEqual(mapped.cell(2, 3).value, "乙")
        self.assertEqual(ws.cell(1, 2).value, "姓名")
        ws.cell(1, 4).value = "更换证件列"
        with self.assertRaises(TemplateSelectionRequired):
            invoke("archive_import", lambda: map_sheet(ws, "transfer"), saved)

    def test_optional_fields_only_become_required_when_configured(self):
        ws = sheet(["姓名", "身份证号码"])
        self.assertIsNotNone(invoke("insurance_ledger", lambda: map_sheet(ws, "policy")))
        with self.assertRaises(TemplateSelectionRequired):
            invoke("insurance_ledger", lambda: map_sheet(ws, "policy"),
                   {"fields": {"policy|每人伤残死亡限额": ["保额"]}})

    def test_missing_and_multiple_sheets_prompt_and_single_alias_matches(self):
        ws = sheet(["公司", "姓名", "身份证"], name="明细A")
        wb = ws.parent
        rules = {"sheets": {"transfer": ["明细A", "明细B"]}}
        self.assertIs(invoke("archive_import", lambda: choose_sheet(wb.worksheets, "transfer"), rules), ws)
        wb.create_sheet("明细B")
        with self.assertRaises(TemplateSelectionRequired):
            invoke("archive_import", lambda: choose_sheet(wb.worksheets, "transfer"), rules)
        with self.assertRaises(TemplateSelectionRequired):
            invoke("archive_import", lambda: choose_sheet(wb.worksheets, "transfer"), {"sheets": {"transfer": ["不存在"]}})

    def test_optional_change_type_absent_does_not_require_all_four_sheets(self):
        ws = sheet(["序号", "姓名", "离职日期"], name="减员")
        self.assertIsNone(invoke("personnel_change_merge", lambda: choose_sheet(ws.parent.worksheets, "增员", required=False, allow_absent=True),
                                 {"sheets": {"增员": ["入职人员"]}}))

    def test_xls_and_grid_header_values_are_equal(self):
        class XlsSheet:
            name = "Sheet1"
            nrows, ncols = 2, 3
            rows = [["名字", "证件号码", "缴费基数"], ["测试人员", "TEST-001", 500]]
            def row_values(self, r): return list(self.rows[r])
            def cell_value(self, r, c): return self.rows[r][c]
        rules = {"fields": {"payment|姓名": ["名字"]}}
        old = XlsSheet()
        xls = invoke("social_security", lambda: map_sheet(old, "payment"), rules)
        grid = SheetGrid(sheet(old.rows[0], old.rows[1]))
        xlsx = invoke("social_security", lambda: map_sheet(grid, "payment"), rules)
        self.assertEqual(xls.row_values(0), [xlsx.cell(1, c).value for c in range(1, 4)])
        self.assertEqual(xls.cell_value(1, 2), xlsx.cell(2, 3).value)
        self.assertEqual(old.rows[0][0], "名字")

    def test_context_resets_even_when_mapping_fails(self):
        with self.assertRaises(RuntimeError):
            invoke("archive_import", lambda: (_ for _ in ()).throw(RuntimeError("stop")))
        self.assertFalse(active())
        ws = sheet(["任意列"])
        self.assertIs(map_sheet(ws, "transfer"), ws)

    def test_name_aliases_do_not_change_salary_layout_or_source_headers(self):
        from hr_toolkit.tools.salary_split import _detect_layout
        ws = sheet(["序号", "入职公司", "姓名", "身份证号码", "应发小计"], name="明细")
        ws.parent.create_sheet("汇总")
        before = _detect_layout(ws.parent)
        self.assertEqual(invoke("salary_split", lambda: _detect_layout(ws.parent)), before)
        ws.cell(1, 3).value = "name"
        after = invoke("salary_split", lambda: _detect_layout(ws.parent), {"fields": {"detail|姓名": ["name"]}})
        self.assertEqual(after, before)
        self.assertEqual(ws.cell(1, 3).value, "name")

    def test_reader_results_equal_with_empty_rules(self):
        from hr_toolkit.tools.archive_import import _read_transfer_file
        from hr_toolkit.tools.insurance_ledger import _read_policy_file
        from hr_toolkit.tools.personnel_change_merge import _read_change_file
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            examples = [
                ("archive_import", _read_transfer_file, sheet(["公司", "姓名", "身份证"], ["甲公司", "测试人员", "TEST-001"])),
                ("insurance_ledger", lambda p: _read_policy_file(p, []), sheet(["姓名", "身份证号码", "每人伤残死亡限额"], ["测试人员", "TEST-001", 600000])),
                ("personnel_change_merge", _read_change_file, sheet(["序号", "姓名", "入职日期"], [1, "测试人员", datetime(2026, 8, 1)], name="增员")),
            ]
            for tool, reader, ws in examples:
                path = root / (tool + "_202608.xlsx")
                ws.parent.save(path)
                original = path.read_bytes()
                baseline = reader(path)
                rules = {}
                if tool == "personnel_change_merge":
                    with self.assertRaises(TemplateSelectionRequired) as caught:
                        invoke(tool, lambda: reader(path))
                    rules = save_choice(tool, {}, caught.exception.payload,
                                        {"sheet_selections": {"减员": None, "转正": None, "调动": None}})
                self.assertEqual(invoke(tool, lambda: reader(path), rules), baseline)
                self.assertEqual(path.read_bytes(), original)

    def test_report_sheet_renaming_does_not_guess_weekly_or_monthly(self):
        from hr_toolkit.tools.data_statistics import _adapt_statistics_grid
        grid = SheetGrid(sheet(["汇报编号", "汇报人", "汇报时间"], ["R001", "测试人员", datetime(2026, 8, 1)], name="数据"))
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("data_statistics", lambda: _adapt_statistics_grid(grid, "数据.xlsx"))
        self.assertEqual([r["key"] for r in caught.exception.payload["roles"]], ["weekly", "monthly"])

    def test_ignore_is_explicit_and_bound_to_file_and_preview(self):
        ws = sheet(["填写说明"], ["说明文字"])
        with self.assertRaises(TemplateSelectionRequired) as caught:
            invoke("data_statistics", lambda: request_selection([ws], ["weekly"], file="说明.xlsx", allow_ignore=True))
        saved = save_choice("data_statistics", {}, caught.exception.payload,
                            {"role": "_ignore", "sheet": "Sheet1", "row": 1, "columns": {}})
        saved = json.loads(json.dumps(saved))
        self.assertTrue(invoke("data_statistics", lambda: ignored_sheet(ws, "说明.xlsx"), saved))
        self.assertFalse(invoke("data_statistics", lambda: ignored_sheet(ws, "别的文件.xlsx"), saved))
        ws.cell(1, 1).value = "汇报人"
        self.assertFalse(invoke("data_statistics", lambda: ignored_sheet(ws, "说明.xlsx"), saved))


if __name__ == "__main__":
    unittest.main()
