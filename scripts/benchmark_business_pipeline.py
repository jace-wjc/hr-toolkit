"""Same-input stress runner with source hashes and full XLSX package comparison.

Fixtures are synthetic. No client attachments are read or changed. Run prepare
once, then run against two source roots using the same --root. Only volatile
OOXML creation/modification timestamps are omitted from the package manifest.
"""
from __future__ import annotations

import argparse
import cProfile
import hashlib
import json
import sys
import time
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path


CASES = ("social", "insurance", "attendance", "salary_merge", "salary_split", "changes", "roster", "archive", "archive_import", "rename")


def peak_rss_bytes() -> int:
    try:
        import resource
        return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * (1 if sys.platform == "darwin" else 1024)
    except ImportError:
        import psutil
        memory = psutil.Process().memory_info()
        return int(getattr(memory, "peak_wset", memory.rss))


def prepare(root: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from tests import test_social_security as social
    from tests import test_insurance_ledger as insurance
    from tests import test_data_statistics as attendance
    from tests import test_salary_merge as salary
    from tests import test_salary_split as split
    from tests import test_personnel_change_merge as changes
    from tests import test_archive_import as archive
    from benchmark_archive_generation import _build_summary_fixture

    source = root / "sources"
    source.mkdir(parents=True, exist_ok=False)
    for case in CASES:
        (source / case).mkdir()
    people = [(f"测{chr(0x4e00 + i // 80)}{chr(0x4e00 + i % 80)}", f"11010119900101{i:04d}") for i in range(1000)]
    social._write_roster(source / "social" / "roster.xlsx", [
        [name, identity, "正常", date(2026, 1, 1), "唐人四川", "唐人数智科技股份有限公司", "四川项目部", "项目（成都）", "成本二", 30]
        for name, identity in people
    ])
    (source / "social" / "payments").mkdir()
    for index in range(50):
        social._write_single_kind_rows(source / "social" / "payments" / f"工伤保险_{index:03d}.xlsx",
                                      [(name, identity, 3600, 0.01, 36) for name, identity in people[index * 20:(index + 1) * 20]])
    insurance._write_roster(source / "insurance" / "roster.xlsx")
    workbook = load_workbook(source / "insurance" / "roster.xlsx")
    for name, identity in people:
        workbook.active.append([name, identity, "运维一部", "在职"])
    workbook.save(source / "insurance" / "roster.xlsx")
    workbook.close()
    (source / "insurance" / "policies").mkdir()
    for index in range(30):
        path = source / "insurance" / "policies" / f"PZDX{index:04d}.xlsx"
        insurance._write_pzdx_policy(path)
        workbook = load_workbook(path)
        workbook.active.cell(5, 1, f"保单号码:PZDX{index:04d}")
        for row, (name, identity) in enumerate(people, 12):
            for col, value in enumerate([row - 6, name, "男性", identity, None, "CNY", "600,000.00", "60,000.00"], 1):
                workbook.active.cell(row, col, value)
        workbook.save(path)
        workbook.close()
    attendance._write_large_attendance_file(source / "attendance" / "日结果.xlsx", 1000, 30)
    for index in range(30):
        salary._write_salary_file(source / "salary_merge" / f"2026年8月工资表_{index:02d}.xlsx",
                                  [(name, identity, 5000 + index) for name, identity in people], month_date=date(2026, 8, 1))
    split._write_many_areas_sample(source / "salary_split" / "salary.xlsx", num_areas=100)
    for index in range(50):
        changes._write_change_file(source / "changes" / f"2026年4月项目{index:02d}异动表.xlsx", {
            "增员": [[name, identity, "生产人员", "2026-04-01"] for name, identity in people[index * 20:(index + 1) * 20]]
        })
    changes._write_analysis_template(source / "roster" / "analysis.xlsx")
    changes._write_change_file(source / "roster" / "2026年4月异动汇总.xlsx", {
        "增员": [[name, identity, "生产人员", "2026-04-01"] for name, identity in people]
    })
    _build_summary_fixture(source / "archive" / "archive.xlsx", 6000)
    for index in range(30):
        archive._write_transfer_file(source / "archive_import" / f"移交表_{index:03d}.xlsx")
    for index in range(10000):
        (source / "rename" / f"人员{index:05d}").mkdir()


def digest(path: Path) -> str:
    sha = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()


def output_manifest(root: Path) -> dict:
    manifest = {}
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        relative = str(path.relative_to(root))
        if path.suffix != ".xlsx":
            manifest[relative] = digest(path)
            continue
        parts = {}
        with zipfile.ZipFile(path) as archive:
            for name in sorted(archive.namelist()):
                data = archive.read(name)
                if name == "docProps/core.xml":
                    xml = ET.fromstring(data)
                    for child in list(xml):
                        if child.tag.rsplit("}", 1)[-1] in {"created", "modified"}:
                            xml.remove(child)
                    data = ET.tostring(xml)
                parts[name] = hashlib.sha256(data).hexdigest()
        manifest[relative] = parts
    return manifest


def run_case(root: Path, case: str):
    from hr_toolkit.tools import social_security, insurance_ledger, data_statistics, salary_merge, salary_split, personnel_change_merge, archive_import, folder_rename
    source = root / "sources" / case
    output = root / "result" / case
    if case == "social":
        return social_security.generate_social_security_reports(source / "payments", source / "roster.xlsx", output)
    if case == "insurance":
        return insurance_ledger.generate_insurance_ledger(source / "policies", source / "roster.xlsx", output)
    if case == "attendance":
        return data_statistics.generate_data_statistics_reports(source, output)
    if case == "salary_merge":
        return salary_merge.merge_monthly_salary(source, output)
    if case == "salary_split":
        return salary_split.split_salary_by_company(source / "salary.xlsx", output)
    if case == "changes":
        return personnel_change_merge.merge_personnel_changes(source, output)
    if case == "roster":
        return personnel_change_merge.update_roster_from_change_summaries(source / "2026年4月异动汇总.xlsx", source / "analysis.xlsx", output)
    if case == "archive":
        return archive_import.export_company_archive_tables(source / "archive.xlsx", output)
    if case == "archive_import":
        return archive_import.import_archive_transfers(source, None, output)
    return folder_rename.rename_person_folders(root_dir=source, mode="append", text="_资料", dry_run=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--label", default="before")
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--case", choices=CASES)
    parser.add_argument("--profile", action="store_true")
    args = parser.parse_args()
    sys.path.insert(0, str(args.source_root.resolve()))
    sys.path.insert(1, str(args.source_root.resolve() / "scripts"))
    root = args.root.resolve()
    root.mkdir(parents=True, exist_ok=True)
    if args.prepare:
        prepare(root)
        return
    sources = {str(p.relative_to(root)): digest(p) for p in (root / "sources").rglob("*") if p.is_file()}
    reports = root / args.label
    reports.mkdir(exist_ok=True)
    for case in (args.case,) if args.case else CASES:
        output = root / "result" / case
        if output.exists():
            raise RuntimeError(f"Output already exists: {output}")
        profile = cProfile.Profile() if args.profile else None
        started = time.perf_counter()
        if profile:
            profile.enable()
        result = run_case(root, case)
        if profile:
            profile.disable()
            profile.dump_stats(str(reports / f"{case}.prof"))
        seconds = time.perf_counter() - started
        payload = result.to_dict()
        manifest = output_manifest(output) if output.exists() else {}
        report = {"seconds": seconds, "peak_rss_bytes": peak_rss_bytes(),
                  "result": payload, "outputs": manifest, "sources": sources}
        (reports / f"{case}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        if output.exists():
            output.rename(reports / case)
        print(json.dumps({"case": case, "seconds": round(seconds, 4), "output_files": len(manifest)}, ensure_ascii=False), flush=True)
    assert sources == {str(p.relative_to(root)): digest(p) for p in (root / "sources").rglob("*") if p.is_file()}, "Source changed"


if __name__ == "__main__":
    main()
