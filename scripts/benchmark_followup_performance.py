"""Same-input benchmarks for footer appends, many-company splits and cache checkpoints.

Cache mode measures persistence only, with synthetic OCR entries; no inference.
Use --prepare once and the same --root with different --source-root/--label.
"""
from __future__ import annotations

import argparse
import cProfile
import json
import platform
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch


def prepare(source: Path, case: str, companies: int) -> None:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from tests import test_personnel_change_merge as changes
    from tests import test_salary_split as salary
    source.mkdir(parents=True, exist_ok=False)
    if case == "footer":
        changes._write_real_summary_template(source / "summary.xlsx")
        wb = load_workbook(source / "summary.xlsx")
        ws = wb["增员"]
        for col in range(23, 31):
            ws.cell(2, col, f"附加字段{col}")
        for index in range(5000):
            row = index + 3
            for col in range(1, 31):
                ws.cell(row, col, "既有资料")
            ws.cell(row, 1, index + 1)
            ws.cell(row, 4, f"原有{index}")
            ws.cell(row, 5, f"11010119900101{index:04d}")
            ws.cell(row, 19, "2026-04-01")
            ws.cell(row, 22).value = None
        for row in range(5003, 5008):
            ws.cell(row, 1, row - 2)
            ws.cell(row, 4).fill = PatternFill("solid", fgColor="FFE0FFE0")
        ws.cell(5008, 1, "制表人：测试")
        ws.cell(5008, 1).font = Font(name="宋体", bold=True)
        ws.merge_cells("A5009:C5009")
        ws.cell(5009, 1, "审核人：测试")
        ws.cell(5010, 31, "=SUM(A3:A5002)")
        ws.cell(5010, 31).comment = Comment("表尾批注", "测试")
        ws.cell(5010, 31).hyperlink = "https://example.com/footer"
        ws.row_dimensions[5008].height = 31
        wb.save(source / "summary.xlsx")
        wb.close()
        rows = [[f"新增{i}", f"11010219900101{i:04d}", "生产人员", "2026-04-01"] for i in range(1000)]
        rows += rows[:50]
        rows += [[f"原有{i}", f"11010119900101{i:04d}", "生产人员", "2026-04-01"] for i in range(50)]
        changes._write_change_file(source / "2026年4月新增.xlsx", {"增员": rows})
    elif case == "salary":
        path = source / "salary.xlsx"
        salary._write_many_areas_sample(path, num_areas=1000)
        wb = load_workbook(path)
        ws = wb["明细表"]
        employee = 0
        for row in ws.iter_rows(min_row=6):
            if row[1].value:
                row[39].value = f"公司{employee % companies:02d}"
                employee += 1
        wb.save(path)
        wb.close()
    else:
        completed_at = datetime.now()
        now = completed_at.strftime("%Y-%m-%d %H:%M:%S")
        cache = {"version": 6, "engine_signature": "synthetic", "pdf_backend_signature": "synthetic",
                 "created_at": now, "updated_at": now, "entries": {}, "paths": {}}
        for index in range(30000):
            key = f"{index:064x}_8"
            relative = f"{index:06d}.jpg"
            # Spread entries across roughly three hours, three records per second.
            verified_at = (completed_at - timedelta(seconds=(29999 - index) // 3)).strftime("%Y-%m-%d %H:%M:%S")
            cache["entries"][key] = {"verified_at": verified_at, "ocr_text": "合成缓存资料" * 60,
                                     "material_type": "身份证", "analysis_state": "complete",
                                     "index_scope": "flat_ocr", "source_size": 8, "sample_filename": relative}
            cache["paths"][relative] = {"cache_key": key, "source_size": 8, "source_mtime_ns": 123}
        (source / "cache.json").write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")


def run(source: Path, output: Path, case: str) -> dict:
    if case == "footer":
        from hr_toolkit.tools.personnel_change_merge import merge_personnel_changes
        return merge_personnel_changes(source / "2026年4月新增.xlsx", output,
                                       template_path=source / "summary.xlsx").to_dict()
    if case == "salary":
        from hr_toolkit.tools.salary_split import split_salary_by_company
        return split_salary_by_company(source / "salary.xlsx", output).to_dict()
    from hr_toolkit.tools import material_collector as mc
    seed = json.loads((source / "cache.json").read_text(encoding="utf-8"))
    cache = dict(seed, entries={}, paths={})
    result = {"checkpoint_count": 0, "written_bytes": 0, "trim_seconds": 0.0, "save_seconds": 0.0}
    output.mkdir()
    cache_path = output / "cache.json"
    for index, (key, entry) in enumerate(seed["entries"].items(), 1):
        cache["entries"][key] = entry
        relative = entry["sample_filename"]
        cache["paths"][relative] = seed["paths"][relative]
        if (index <= 1000 and index % 100 == 0) or (index > 1000 and index % 1000 == 0):
            start = time.perf_counter()
            mc._trim_cache_by_age_and_size(cache)
            result["trim_seconds"] += time.perf_counter() - start
            start = time.perf_counter()
            with patch.object(mc, "_beijing_now_str", return_value=seed["updated_at"]):
                if not mc._save_ocr_cache(cache_path, cache):
                    raise RuntimeError("Checkpoint failed")
            result["save_seconds"] += time.perf_counter() - start
            result["written_bytes"] += cache_path.stat().st_size
            result["checkpoint_count"] += 1
    result["entries"] = len(cache["entries"])
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--case", choices=("footer", "salary", "cache"), required=True)
    parser.add_argument("--companies", type=int, choices=(2, 10, 50), default=50)
    parser.add_argument("--label", default="before")
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--profile", action="store_true")
    args = parser.parse_args()
    sys.path.insert(0, str(args.source_root.resolve()))
    # Keep this runner's manifest helper even when testing an older checkout.
    from benchmark_business_pipeline import digest, output_manifest, peak_rss_bytes
    root = args.root.resolve()
    source = root / "inputs"
    if args.prepare:
        prepare(source, args.case, args.companies)
        return
    sources = {str(p.relative_to(source)): digest(p) for p in source.rglob("*") if p.is_file()}
    output = root / "result"
    destination = root / args.label
    if output.exists():
        raise RuntimeError(f"Output already exists: {output}")
    if destination.exists():
        raise RuntimeError(f"Report label already exists: {destination}")
    profile = cProfile.Profile() if args.profile else None
    started = time.perf_counter()
    if profile:
        profile.enable()
    result = run(source, output, args.case)
    if profile:
        profile.disable()
        profile.dump_stats(str(root / f"{args.label}.prof"))
    report = {"case": args.case, "seconds": time.perf_counter() - started,
              "peak_rss_bytes": peak_rss_bytes(), "result": result,
              "sources": sources, "outputs": output_manifest(output),
              "python": sys.version, "platform": platform.platform(),
              "source_root": str(args.source_root.resolve()), "profiled": args.profile}
    assert sources == {str(p.relative_to(source)): digest(p) for p in source.rglob("*") if p.is_file()}, "Source changed"
    destination.mkdir(exist_ok=False)
    output.rename(destination / "result")
    (destination / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"case": args.case, "label": args.label, "seconds": report["seconds"],
                      "peak_rss_bytes": report["peak_rss_bytes"], "output_files": len(report["outputs"])}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
